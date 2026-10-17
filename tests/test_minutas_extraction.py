from copy import deepcopy
from io import BytesIO
import re

from docx import Document
from reportlab.pdfgen import canvas
import pytest

from database.minutas_referencias import EXTRACTION_COLUMNS
from database.store import ROOT
from services.minutas_extraction import (
    DOCX_MIME,
    EXTRACTOR_VERSION,
    PDF_MIME,
    MinutasExtractionError,
    MinutasExtractionService,
    extract_content,
)
from services.minutas_library import MinutasAccessDenied, SUBTIPOS
from services.oficios import MAX_FILE


def _pdf(*pages):
    output = BytesIO()
    document = canvas.Canvas(output)
    for text in pages:
        text_object = document.beginText(72, 760)
        for line in text.splitlines() or [""]:
            text_object.textLine(line)
        document.drawText(text_object)
        document.showPage()
    document.save()
    return output.getvalue()


def _blank_pdf():
    output = BytesIO()
    document = canvas.Canvas(output)
    document.showPage()
    document.save()
    return output.getvalue()


def _docx():
    output = BytesIO()
    document = Document()
    document.add_paragraph("RELATÓRIO")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Irregularidade"
    table.cell(0, 1).text = "Defesa"
    document.add_paragraph("CONCLUSÃO")
    document.save(output)
    return output.getvalue()


def _reference(identifier="ref-1", procuradora="Elvira Samara", subtype="Prefeitura"):
    return {
        "id": identifier,
        "procuradora": procuradora,
        "procuradora_chave": procuradora.casefold(),
        "numero_processo": identifier,
        "subtipo": subtype,
        "auditoria_nome": "auditoria.pdf",
        "auditoria_mime": PDF_MIME,
        "auditoria_tamanho": 100,
        "auditoria_path": f"{identifier}/auditoria.pdf",
        "parecer_nome": "parecer.pdf",
        "parecer_mime": PDF_MIME,
        "parecer_tamanho": 100,
        "parecer_path": f"{identifier}/parecer.pdf",
    }


class FakeRepository:
    def __init__(self):
        self.references = {}
        self.documents = {}
        self.extractions = {}
        self.upserts = 0

    def list_documents(self, reference_id):
        return sorted(
            (
                deepcopy(document)
                for document in self.documents.values()
                if document["referencia_id"] == reference_id
            ),
            key=lambda document: document["posicao"],
        )

    def get_extraction(self, path):
        value = self.extractions.get(path)
        return deepcopy(value) if value else None

    def list_extractions(self, reference_id):
        return [
            deepcopy(value)
            for value in self.extractions.values()
            if value["referencia_id"] == reference_id
        ]

    def list_extractions_by_procuradora(self, procuradora_key):
        return [
            deepcopy(value)
            for value in self.extractions.values()
            if value["procuradora_chave"] == procuradora_key
            and value["status"] == "PROCESSADO"
        ]

    def upsert_extraction(self, record):
        self.upserts += 1
        self.extractions[record["arquivo_path"]] = deepcopy(record)
        return deepcopy(record)


class FakeStorage:
    def __init__(self):
        self.objects = {}
        self.downloads = []

    def download(self, path):
        self.downloads.append(path)
        return self.objects[path]


class _Gestor:
    ativo = True
    pode_admin = True


class FakeLibrary:
    specialty_id = "elvira.pca.rpps"
    principal = _Gestor()
    def __init__(self, repository, storage, authorized=True):
        self.repository = repository
        self.storage = storage
        self.authorized = authorized

    def _authorize(self):
        if not self.authorized:
            raise MinutasAccessDenied("Acesso não autorizado a este módulo.")
        return "admin@mpc.pb.gov.br"

    def get(self, identifier):
        self._authorize()
        return deepcopy(self.repository.references[identifier])


def _install_reference(repository, storage, reference, *, opinion=None):
    repository.references[reference["id"]] = deepcopy(reference)
    storage.objects[reference["auditoria_path"]] = _pdf(
        "RELATÓRIO",
        "ACHADOS DE AUDITORIA",
        "CONCLUSÃO TÉCNICA",
        "Conteúdo técnico preservado integralmente conforme o documento original",
    )
    storage.objects[reference["parecer_path"]] = opinion or _pdf(
        "RELATÓRIO",
        "FUNDAMENTAÇÃO",
        "CONCLUSÃO MINISTERIAL",
        "Entendimento da Procuradora preservado sem qualquer reformulação textual",
    )


def test_extraction_sql_is_additive_private_and_enables_rls():
    sql = (ROOT / "database" / "minutas_extracoes.sql").read_text(encoding="utf-8")
    table = sql.split("CREATE TABLE IF NOT EXISTS", 1)[1].split(
        "CREATE INDEX", 1
    )[0]
    declared = []
    for line in table.splitlines():
        match = re.match(
            r"^    ([a-z0-9_]+)\s+(?:UUID|TEXT|INTEGER|BIGINT)\b", line
        )
        if match:
            declared.append(match.group(1))

    assert declared == EXTRACTION_COLUMNS.split(",")
    assert "REFERENCES mpc_portarias.minutas_referencias(id)" in sql
    assert "ENABLE ROW LEVEL SECURITY" in sql
    assert "FROM PUBLIC, anon, authenticated" in sql
    assert "FORCE ROW LEVEL SECURITY" not in sql
    assert "CREATE POLICY" not in sql


def test_pdf_pages_and_docx_blocks_preserve_origin_and_order():
    pdf = extract_content(
        _pdf("Primeira página com conteúdo suficiente", "Segunda página conclusiva"),
        PDF_MIME,
        "auditoria",
    )
    docx = extract_content(_docx(), DOCX_MIME, "parecer")

    assert [unit["pagina"] for unit in pdf["unidades"]] == [1, 2]
    assert "Primeira página" in pdf["unidades"][0]["texto"]
    assert "Segunda página" in pdf["unidades"][1]["texto"]
    assert [unit["tipo"] for unit in docx["unidades"]] == [
        "paragrafo",
        "tabela",
        "paragrafo",
    ]
    assert "Irregularidade\tDefesa" in docx["texto_integral"]
    assert all("pagina" not in unit for unit in docx["unidades"])
    assert [section["secao"] for section in docx["secoes"]] == [
        "relatorio",
        "conclusao_ministerial",
    ]


@pytest.mark.parametrize(
    "transition, closing",
    [
        ("É o relatório. Passo a opinar.", "É o Parecer."),
        ("E O RELATORIO.   PASSO A\nOPINAR.", "E o parecer."),
    ],
)
def test_real_opinion_markers_identify_sections_without_rewriting_text(
    transition, closing
):
    original = (
        "EMENTA: Prestação de contas anuais. Regularidade com ressalvas.\n\n"
        "O processo foi instruído pela Auditoria e houve manifestação da defesa.\n"
        f"{transition}\n"
        "A análise ministerial considera os documentos e as normas aplicáveis.\n"
        "Ante o exposto, opino pela regularidade com ressalvas das contas.\n"
        f"{closing}"
    )

    extracted = extract_content(_pdf(original), PDF_MIME, "parecer")

    assert extracted["status"] == "PROCESSADO"
    assert [section["secao"] for section in extracted["secoes"]] == [
        "ementa",
        "relatorio",
        "fundamentacao",
        "conclusao_ministerial",
    ]
    assert "EMENTA:" in extracted["secoes"][0]["texto"]
    assert "O processo foi instruído" in extracted["secoes"][1]["texto"]
    assert "passo" in extracted["secoes"][2]["texto"].casefold()
    assert "Ante o exposto" in extracted["secoes"][3]["texto"]
    assert closing in extracted["texto_integral"]
    assert extracted["unidades"][0]["texto"] == extracted["texto_integral"]


def test_blank_scanned_and_corrupt_pdf_need_attention_without_ocr():
    blank = extract_content(_blank_pdf(), PDF_MIME, "auditoria")

    assert blank["status"] == "ATENCAO"
    assert "OCR" in blank["erro"]
    assert blank["unidades"] == [{"pagina": 1, "texto": ""}]
    with pytest.raises(MinutasExtractionError, match="ler"):
        extract_content(b"%PDF-corrompido", PDF_MIME, "auditoria")
    with pytest.raises(MinutasExtractionError, match="10 MB"):
        extract_content(b"x" * (MAX_FILE + 1), PDF_MIME, "auditoria")


def test_validation_only_page_is_warned_and_final_signature_does_not_invalidate():
    extracted = extract_content(
        _pdf(
            "Fundamentação técnica e jurídica suficiente para a análise das contas.",
            "Documento assinado eletronicamente\nTRAMITA\nCódigo verificador 12345",
        ),
        PDF_MIME,
        "parecer",
    )

    assert extracted["status"] == "PROCESSADO"
    assert "página 2" in extracted["erro"]
    assert "TRAMITA" in extracted["texto_integral"]
    assert "assinado eletronicamente" in extracted["unidades"][1]["texto"]

    validation_only = extract_content(
        _pdf("TRAMITA\nPágina 1 de 1\nCódigo verificador 12345"),
        PDF_MIME,
        "parecer",
    )
    assert validation_only["status"] == "ATENCAO"
    assert "Nenhum texto substantivo" in validation_only["erro"]


def test_isolated_page_failure_preserves_other_pages(monkeypatch):
    class Page:
        def __init__(self, text=None, fails=False):
            self.text = text
            self.fails = fails

        def extract_text(self):
            if self.fails:
                raise RuntimeError("page failure")
            return self.text

    class Reader:
        def __init__(self, _source):
            self.pages = [
                Page("Primeira página preservada com conteúdo substantivo suficiente."),
                Page(fails=True),
                Page("Terceira página preservada após a falha isolada da anterior."),
            ]

    monkeypatch.setattr("services.minutas_extraction.PdfReader", Reader)

    extracted = extract_content(b"%PDF-sintetico", PDF_MIME, "auditoria")

    assert extracted["status"] == "ATENCAO"
    assert "página 2" in extracted["erro"]
    assert [unit["pagina"] for unit in extracted["unidades"]] == [1, 2, 3]
    assert extracted["unidades"][1]["texto"] == ""
    assert "Primeira página preservada" in extracted["texto_integral"]
    assert "Terceira página preservada" in extracted["texto_integral"]


def test_corrupt_document_does_not_interrupt_other_documents():
    repository = FakeRepository()
    storage = FakeStorage()
    reference = _reference("corrupt-test")
    _install_reference(repository, storage, reference)
    storage.objects[reference["auditoria_path"]] = b"%PDF-corrompido"
    service = MinutasExtractionService(FakeLibrary(repository, storage))

    result = service.process_reference(reference["id"])

    assert [item["status"] for item in result] == ["ATENCAO", "PROCESSADO"]
    assert result[0]["texto_integral"] == ""
    assert "ler" in result[0]["erro"]


@pytest.mark.parametrize("subtype", SUBTIPOS)
def test_all_pca_subtypes_are_processed_without_changing_originals(subtype):
    repository = FakeRepository()
    storage = FakeStorage()
    reference = _reference(f"ref-{SUBTIPOS.index(subtype)}", subtype=subtype)
    _install_reference(repository, storage, reference)
    before = deepcopy(storage.objects)
    service = MinutasExtractionService(FakeLibrary(repository, storage))

    result = service.process_reference(reference["id"])

    assert len(result) == 2
    assert storage.objects == before
    assert {item["status"] for item in result} == {"PROCESSADO"}


def test_complete_sedh_sequence_hash_reuse_and_forced_reprocessing():
    repository = FakeRepository()
    storage = FakeStorage()
    reference = _reference("TC 02650/25", subtype=SUBTIPOS[-1])
    _install_reference(repository, storage, reference)
    documents = (
        ("defesa", "Relatório de Análise de Defesa", 1, _pdf("ANÁLISE", "DEFESA")),
        ("cota", "Cota Ministerial", 2, _docx()),
        (
            "complemento",
            "Relatório de Complementação de Instrução",
            3,
            _pdf("COMPLEMENTAÇÃO", "CONCLUSÃO"),
        ),
    )
    for identifier, document_type, position, content in documents:
        mime = DOCX_MIME if identifier == "cota" else PDF_MIME
        path = f"{reference['id']}/{identifier}"
        repository.documents[identifier] = {
            "id": identifier,
            "referencia_id": reference["id"],
            "tipo": document_type,
            "posicao": position,
            "nome": identifier + (".docx" if mime == DOCX_MIME else ".pdf"),
            "mime": mime,
            "tamanho": len(content),
            "path": path,
        }
        storage.objects[path] = content
    service = MinutasExtractionService(FakeLibrary(repository, storage))

    first = service.process_reference(reference["id"])
    second = service.process_reference(reference["id"])

    assert [(item["papel"], item["posicao"]) for item in first] == [
        ("auditoria", 0),
        ("complementar", 1),
        ("complementar", 2),
        ("complementar", 3),
        ("parecer", None),
    ]
    assert all(len(item["arquivo_sha256"]) == 64 for item in first)
    assert all(item["extrator_versao"] == EXTRACTOR_VERSION for item in first)
    assert all(item["reutilizado"] is True for item in second)
    assert repository.upserts == 5

    before_other_extractions = {
        path: deepcopy(value)
        for path, value in repository.extractions.items()
        if path != repository.documents["cota"]["path"]
    }
    storage.downloads.clear()
    reprocessed = service.reprocess_document(
        reference["id"], repository.documents["cota"]["path"]
    )

    assert len(reprocessed) == 1
    assert storage.downloads == [repository.documents["cota"]["path"]]
    assert {
        path: value
        for path, value in repository.extractions.items()
        if path != repository.documents["cota"]["path"]
    } == before_other_extractions
    assert repository.upserts == 6

    selected_path = repository.documents["cota"]["path"]
    del storage.objects[selected_path]
    failed = service.reprocess_document(reference["id"], selected_path)

    assert failed[0]["status"] == "ATENCAO"
    assert storage.downloads == [selected_path, selected_path]
    assert {
        path: value
        for path, value in repository.extractions.items()
        if path != selected_path
    } == before_other_extractions
    assert repository.upserts == 7


def test_context_is_isolated_by_prosecutor_and_excludes_target_opinion():
    repository = FakeRepository()
    storage = FakeStorage()
    target = _reference("Taperoá")
    same_office = _reference("Rio Tinto")
    other_office = _reference("Outro", procuradora="Outra Procuradora")
    service = MinutasExtractionService(FakeLibrary(repository, storage))
    for reference in (target, same_office, other_office):
        _install_reference(repository, storage, reference)
        service.process_reference(reference["id"])

    target_opinion = repository.extractions[target["parecer_path"]]
    duplicate_path = same_office["auditoria_path"]
    repository.extractions[duplicate_path].update(
        arquivo_sha256=target_opinion["arquivo_sha256"],
        texto_integral=target_opinion["texto_integral"],
    )
    repository.extractions[same_office["parecer_path"]].update(
        arquivo_sha256="historical-opinion-unique-hash",
        texto_integral="Parecer histórico autônomo e autorizado.",
    )

    sources = service.context_sources(target["id"])

    assert all(item["procuradora"] == target["procuradora"] for item in sources)
    assert not any(
        item["referencia_id"] == target["id"] and item["papel"] == "parecer"
        for item in sources
    )
    assert not any(item["arquivo_path"] == duplicate_path for item in sources)
    assert any(item["referencia_id"] == same_office["id"] for item in sources)


def test_server_side_authorization_precedes_storage_and_persistence():
    repository = FakeRepository()
    storage = FakeStorage()
    reference = _reference()
    _install_reference(repository, storage, reference)
    service = MinutasExtractionService(
        FakeLibrary(repository, storage, authorized=False)
    )

    with pytest.raises(MinutasAccessDenied):
        service.process_reference(reference["id"])

    assert storage.downloads == []
    assert repository.extractions == {}
