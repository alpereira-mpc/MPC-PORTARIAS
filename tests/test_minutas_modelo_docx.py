"""O DOCX gerado usa o modelo vigente do Procurador e da especialidade."""

import base64
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

from services.minutas_docx import (
    MinutasDocxError,
    MinutasDocxService,
    _apply_numbering,
    _numbering_definition,
)
from services.minutas_generation import MinutasGenerationError, MinutasGenerationService
from services.minutas_library import MinutasLibraryService
from services.minutas_generation import _draft_header_lines
from services.minutas_modelo_docx import (
    MENSAGEM_MODELO_AUSENTE,
    MENSAGEM_MODELO_INVALIDO,
    impressao_modelo,
    renderizar,
    validar_pacote,
)
from services.minutas_modelos import MinutasModeloVisualService
from services.minutas_orientacao import aplicar_bloco
from services.minutas_procuradores import exigir_procurador
from services.minutas_specialties import CAMARAS, PREFEITURAS, SHARED_INSTRUCTION_MOTOR
from tests.test_minutas_instrucao import _structure
from tests.test_minutas_modelos import _install
from tests.test_minutas_specialties import GestorAcervo, Storage, TEST_IDENTITY, build_store


PNG_A = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
PNG_B = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+ip1sAAAAASUVORK5CYII="
)
VENENO = "Processo secreto do modelo visual que não pode reaparecer na minuta gerada."


def _fonte(run, nome, tamanho):
    run.font.name = nome
    run.font.size = Pt(tamanho)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), nome)


def _modelo(
    nome_arquivo,
    *,
    fonte,
    esquerda,
    superior,
    logo,
    largura,
    ementa=WD_ALIGN_PARAGRAPH.LEFT,
    assinatura=WD_ALIGN_PARAGRAPH.CENTER,
):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = fonte
    normal.font.size = Pt(12)
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(superior)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(esquerda)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = header.add_run()
    image.add_picture(BytesIO(logo), width=Emu(largura))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"Rodapé exclusivo {fonte}")
    body = document.add_paragraph(VENENO)
    body.paragraph_format.first_line_indent = Cm(1.2)
    _fonte(body.runs[0], fonte, 12)
    ementa_item = document.add_paragraph("Ementa de exemplo que não pode sair.")
    ementa_item.alignment = ementa
    ementa_item.paragraph_format.left_indent = Cm(6.5)
    _fonte(ementa_item.runs[0], fonte, 10)
    title = document.add_paragraph("Irregularidade de exemplo do modelo")
    title.runs[0].bold = True
    title.runs[0].italic = True
    _fonte(title.runs[0], fonte, 12)
    name = document.add_paragraph("NOME DE EXEMPLO DO MODELO")
    name.alignment = assinatura
    name.runs[0].bold = True
    _fonte(name.runs[0], fonte, 14)
    role = document.add_paragraph("Cargo de exemplo do modelo")
    role.alignment = assinatura
    _fonte(role.runs[0], fonte, 11)
    output = BytesIO()
    document.save(output)
    return nome_arquivo, output.getvalue()


def _gabinete(store, storage, specialty, sigla):
    from database.minutas_referencias import MinutasReferenciasStore

    repo = object.__new__(MinutasReferenciasStore)
    repo.store, repo.specialty_id, repo.include_acervo = store, specialty, True
    return MinutasLibraryService(
        repo,
        storage,
        TEST_IDENTITY,
        {"MINUTAS_PILOTO_ENABLED": True, "MINUTAS_PILOTO_EMAILS": [TEST_IDENTITY["email"]]},
        procurador=exigir_procurador(sigla),
        principal=GestorAcervo(),
    )


def _habilitar(store, specialty):
    with store.connection() as connection:
        connection.execute(
            "UPDATE minutas_especialidades SET motor=?, estado=? WHERE id=?",
            (SHARED_INSTRUCTION_MOTOR, "Em testes", specialty),
        )


def _exportar(library, processo="02659/25"):
    key = processo.replace("/", "")
    analyzed = {"fontes_versao": "v", "processo_alvo_id": f"instrucao-{key}", "analise": {}}
    generated = {
        "fontes_versao": "v",
        "processo_alvo_id": f"instrucao-{key}",
        "estrutura": _structure(),
        "metadados_cabecalho": {},
    }
    return MinutasDocxService(library).export_instruction(
        analyzed, generated, processo=processo, exercicio=2025, origem="Prefeitura Municipal de Exemplo",
    )


def _pacote(content, name):
    with ZipFile(BytesIO(content)) as package:
        return package.read(name).decode("utf-8")


def _logotipos(content):
    document = Document(BytesIO(content))
    return [
        relationship.target_part.blob
        for section in document.sections
        for relationship in section.header.part.rels.values()
        if relationship.reltype.endswith("/image")
    ]


@pytest.fixture
def ambiente(tmp_path):
    store = build_store(tmp_path / "minutas.db")
    _install(store)
    storage = Storage()
    _habilitar(store, PREFEITURAS)
    _habilitar(store, CAMARAS)
    return store, storage


def test_pacote_com_macro_ou_caminho_externo_e_recusado():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as package:
        package.writestr("[Content_Types].xml", "<Types></Types>")
        package.writestr("word/document.xml", "<w:document/>")
        package.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(MinutasDocxError, match="inválido ou incompatível"):
        validar_pacote(buffer.getvalue())
    slipped = BytesIO()
    with ZipFile(slipped, "w") as package:
        package.writestr("[Content_Types].xml", "<Types></Types>")
        package.writestr("word/../fora.txt", b"x")
    with pytest.raises(MinutasDocxError, match="inválido ou incompatível"):
        validar_pacote(slipped.getvalue())


def test_bradson_prefeitura_usa_o_modelo_cadastrado_e_nao_o_de_elvira(ambiente):
    store, storage = ambiente
    bradson = _gabinete(store, storage, PREFEITURAS, "BLTC")
    elvira = _gabinete(store, storage, PREFEITURAS, "ESPO")
    modelo_bradson = _modelo(
        "bradson-prefeitura.docx", fonte="Times New Roman", esquerda=3.2, superior=2.4,
        logo=PNG_B, largura=900000, ementa=WD_ALIGN_PARAGRAPH.RIGHT,
        assinatura=WD_ALIGN_PARAGRAPH.CENTER,
    )
    modelo_elvira = _modelo(
        "elvira-prefeitura.docx", fonte="Arial", esquerda=1.8, superior=5.2,
        logo=PNG_A, largura=1400000, ementa=WD_ALIGN_PARAGRAPH.LEFT,
        assinatura=WD_ALIGN_PARAGRAPH.LEFT,
    )
    MinutasModeloVisualService(bradson).save(PREFEITURAS, modelo_bradson)
    MinutasModeloVisualService(elvira).save(PREFEITURAS, modelo_elvira)

    gerado = _exportar(bradson, "02659/25")
    documento = Document(BytesIO(gerado.content))
    texto = "\n".join(item.text for item in documento.paragraphs)
    xml = _pacote(gerado.content, "word/document.xml")
    estilos = _pacote(gerado.content, "word/styles.xml")
    ementa = next(item for item in documento.paragraphs if item.text.startswith("Ementa:"))
    assinatura = next(item for item in documento.paragraphs if "BRADSON" in item.text)

    assert gerado.modelo["procuradora_chave"] == "bltc"
    assert gerado.modelo["especialidade_id"] == PREFEITURAS
    assert gerado.modelo["sha256"] == sha256(modelo_bradson[1]).hexdigest()
    assert documento.sections[0].left_margin.cm == pytest.approx(3.2, abs=0.05)
    assert documento.sections[0].top_margin.cm == pytest.approx(2.4, abs=0.05)
    assert ementa.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert assinatura.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "Times New Roman" in estilos
    assert "Times New Roman" in xml
    assert "Minuta Procuradora" not in estilos
    assert "Minuta Procuradora" not in xml
    assert _logotipos(gerado.content) == [PNG_B]
    assert "BRADSON TIBÉRIO LUNA CAMELO" in texto
    assert "ELVIRA SAMARA" not in texto
    assert VENENO not in texto
    assert "NOME DE EXEMPLO" not in texto
    assert "02659/25" in texto
    assert "Rodapé exclusivo Times New Roman" in "\n".join(
        item.text for item in documento.sections[0].footer.paragraphs
    )
    assert gerado.content[:2] == b"PK"
    Document(BytesIO(gerado.content))

    outra = _exportar(elvira, "02659/25")
    outra_doc = Document(BytesIO(outra.content))
    assert outra.modelo["procuradora_chave"] == "elvira samara"
    assert outra_doc.sections[0].left_margin.cm == pytest.approx(1.8, abs=0.05)
    assert _logotipos(outra.content) == [PNG_A]
    assert "Arial" in _pacote(outra.content, "word/styles.xml")
    assert "BRADSON" not in "\n".join(item.text for item in outra_doc.paragraphs)
    assert "ELVIRA SAMARA PEREIRA DE OLIVEIRA" in "\n".join(
        item.text for item in outra_doc.paragraphs
    )


def test_especialidade_e_versao_nao_contaminam_outro_modelo(ambiente):
    store, storage = ambiente
    prefeitura = _gabinete(store, storage, PREFEITURAS, "BLTC")
    camara = _gabinete(store, storage, CAMARAS, "BLTC")
    primeiro = _modelo(
        "prefeitura-v1.docx", fonte="Times New Roman", esquerda=3.2, superior=2.2,
        logo=PNG_B, largura=900000,
    )
    segundo = _modelo(
        "prefeitura-v2.docx", fonte="Times New Roman", esquerda=3.2, superior=4.4,
        logo=PNG_B, largura=900000,
    )
    camaras = _modelo(
        "camara.docx", fonte="Calibri", esquerda=1.5, superior=2.2,
        logo=PNG_A, largura=700000, assinatura=WD_ALIGN_PARAGRAPH.LEFT,
    )
    servico = MinutasModeloVisualService(prefeitura)
    assert impressao_modelo(prefeitura) == "ausente"
    servico.save(PREFEITURAS, primeiro)
    antes = impressao_modelo(prefeitura)
    gerado = _exportar(prefeitura)
    assert Document(BytesIO(gerado.content)).sections[0].top_margin.cm == pytest.approx(2.2, abs=0.05)
    servico.save(PREFEITURAS, segundo)
    depois = impressao_modelo(prefeitura)
    assert antes != depois
    assert Document(BytesIO(_exportar(prefeitura).content)).sections[0].top_margin.cm == pytest.approx(4.4, abs=0.05)
    servico.save(PREFEITURAS, primeiro)
    assert Document(BytesIO(_exportar(prefeitura).content)).sections[0].top_margin.cm == pytest.approx(2.2, abs=0.05)
    MinutasModeloVisualService(camara).save(CAMARAS, camaras)
    camara_doc = Document(BytesIO(_exportar(camara, "02108/26").content))
    assert camara_doc.sections[0].left_margin.cm == pytest.approx(1.5, abs=0.05)
    assert _logotipos(MinutasDocxService(camara).export_instruction(
        {"fontes_versao": "v", "processo_alvo_id": "instrucao-0210826", "analise": {}},
        {
            "fontes_versao": "v",
            "processo_alvo_id": "instrucao-0210826",
            "estrutura": _structure(),
            "metadados_cabecalho": {},
        },
        processo="02108/26",
        exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    ).content) == [PNG_A]
    assert _logotipos(_exportar(prefeitura).content) == [PNG_B]


def test_ausencia_ou_modelo_invalido_nao_cai_no_padrao_de_elvira(ambiente):
    store, storage = ambiente
    bradson = _gabinete(store, storage, PREFEITURAS, "BLTC")
    chamadas = []
    servico = MinutasGenerationService(bradson, generate=lambda contexto: chamadas.append(contexto))
    with pytest.raises(MinutasDocxError, match="não há um modelo visual válido"):
        _exportar(bradson)
    with pytest.raises(MinutasGenerationError, match="não há um modelo visual válido"):
        servico._generate({})
    assert chamadas == []
    assert MENSAGEM_MODELO_AUSENTE.startswith("Não foi possível gerar o DOCX")

    salvo = MinutasModeloVisualService(bradson).save(
        PREFEITURAS,
        _modelo(
            "valido.docx", fonte="Times New Roman", esquerda=3, superior=2,
            logo=PNG_B, largura=800000,
        ),
    )
    storage.objects[salvo["path"]] = b"PK\x03\x04isto-nao-e-docx"
    with pytest.raises(MinutasDocxError, match="inválido ou incompatível"):
        _exportar(bradson)
    with pytest.raises(MinutasGenerationError, match="inválido ou incompatível"):
        servico._generate({})
    assert chamadas == []
    assert MENSAGEM_MODELO_INVALIDO.startswith("Não foi possível gerar o DOCX")


def test_modelo_valido_nao_dispensa_a_orientacao_do_assessor(ambiente):
    store, storage = ambiente
    bradson = _gabinete(store, storage, PREFEITURAS, "BLTC")
    MinutasModeloVisualService(bradson).save(
        PREFEITURAS,
        _modelo(
            "valido.docx", fonte="Times New Roman", esquerda=3, superior=2,
            logo=PNG_B, largura=800000,
        ),
    )
    chamadas = []
    servico = MinutasGenerationService(
        bradson, generate=lambda contexto: chamadas.append(contexto) or {"ok": True},
    )
    with pytest.raises(MinutasGenerationError, match="Informe a orientação"):
        servico._generate({})
    contexto = {}
    aplicar_bloco(
        contexto,
        "Sustentar a linha argumentativa indicada pelo assessor nesta minuta experimental.",
    )
    assert servico._generate(contexto)["ok"] is True
    assert chamadas[0]["orientacao_juridica_usuario"]["texto"].startswith("Sustentar")


OFICIAL_ASSESSOR = (
    Path(__file__).resolve().parents[1]
    / "tmp"
    / "diagnostico-modelo"
    / "minuta oficial do assessor.docx"
)


def _jc(paragraph):
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is None:
        return ""
    align = properties.find(qn("w:jc"))
    if align is None:
        return ""
    return str(align.get(qn("w:val")) or "")


def _estilo(paragraph):
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is None:
        return ""
    style = properties.find(qn("w:pStyle"))
    if style is None:
        return ""
    return str(style.get(qn("w:val")) or "")


def _num_id(paragraph):
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is None:
        return None
    number = properties.find(qn("w:numPr"))
    if number is None:
        return None
    identifier = number.find(qn("w:numId"))
    if identifier is None:
        return None
    return identifier.get(qn("w:val"))


def _negrito(paragraph):
    for run in paragraph._p.iter(qn("w:r")):
        text = "".join(node.text or "" for node in run.findall(qn("w:t"))).strip()
        if not text:
            continue
        properties = run.find(qn("w:rPr"))
        return properties is not None and properties.find(qn("w:b")) is not None
    return False


def _fonte_run(paragraph):
    for run in paragraph._p.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            continue
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            continue
        return fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi")) or ""
    return ""


def _definir_jc(paragraph, value):
    properties = paragraph._p.get_or_add_pPr()
    align = properties.find(qn("w:jc"))
    if align is None:
        align = OxmlElement("w:jc")
        properties.append(align)
    align.set(qn("w:val"), value)


def _sombrear(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def _largura(cell, dxa):
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(dxa))
    width.set(qn("w:type"), "dxa")


def _largura_atual(cell):
    properties = cell._tc.find(qn("w:tcPr"))
    if properties is None:
        return ""
    width = properties.find(qn("w:tcW"))
    if width is None:
        return ""
    return str(width.get(qn("w:w")) or "")


def _formato_numero(document, number_id):
    numbering = document.part.numbering_part.element
    abstract_id = None
    for number in numbering.findall(qn("w:num")):
        if number.get(qn("w:numId")) == str(number_id):
            reference = number.find(qn("w:abstractNumId"))
            abstract_id = reference.get(qn("w:val")) if reference is not None else None
    if abstract_id is None:
        return ""
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) != abstract_id:
            continue
        level = abstract.find(qn("w:lvl"))
        if level is None:
            return ""
        fmt = level.find(qn("w:numFmt"))
        return str(fmt.get(qn("w:val")) or "") if fmt is not None else ""
    return ""


def _mapa_numeracao(document):
    numbering = document.part.numbering_part.element
    return {
        number.get(qn("w:numId")): _formato_numero(document, number.get(qn("w:numId")))
        for number in numbering.findall(qn("w:num"))
    }


def _texto_tabela(table):
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def _texto_completo(document):
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(_texto_tabela(table) for table in document.tables)
    return "\n".join(parts)


def _modelo_regional(
    nome_arquivo,
    *,
    fonte_relatorio,
    fonte_fundamento,
    esquerda,
    superior,
    logo,
    largura,
    fill,
    kind,
    marker,
    alinhamento_data,
    alinhamento_assinatura,
    largura_ementa,
):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(superior)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(esquerda)
    section.right_margin = Cm(2)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = header.add_run()
    image.add_picture(BytesIO(logo), width=Emu(largura))
    footer = section.footer.paragraphs[0]
    footer.add_run(f"Rodapé exclusivo {fonte_relatorio}")
    relatorio = document.styles.add_style("EstiloRelatorioGabinete", WD_STYLE_TYPE.PARAGRAPH)
    fundamento = document.styles.add_style("EstiloFundamentoGabinete", WD_STYLE_TYPE.PARAGRAPH)
    relatorio.font.name = fonte_relatorio
    relatorio.font.size = Pt(12)
    fundamento.font.name = fonte_fundamento
    fundamento.font.size = Pt(11)
    for line in (
        "Processo TC nº 00000/00",
        "Natureza: texto antigo do modelo",
        "Unidade Jurisdicionada: Município Antigo",
        "Exercício: 1999",
    ):
        item = document.add_paragraph(line)
        item.runs[0].bold = True
        _definir_jc(item, "both")
    document.add_paragraph("")
    ementa = document.add_table(rows=1, cols=2)
    ementa.cell(0, 0).text = ""
    _largura(ementa.cell(0, 1), largura_ementa)
    paragrafo = ementa.cell(0, 1).paragraphs[0]
    paragrafo.text = ""
    paragrafo.add_run("EMENTA:").bold = True
    paragrafo.add_run(" ")
    paragrafo.add_run(VENENO)
    document.add_paragraph("")
    faixa = document.add_table(rows=1, cols=1)
    faixa.cell(0, 0).text = "PARECER Nº "
    _sombrear(faixa.cell(0, 0), fill)
    document.add_paragraph("")
    corpo = document.add_paragraph(
        "Texto antigo do relatório deste modelo, longo o bastante para representar o corpo.",
        style=relatorio,
    )
    _fonte(corpo.runs[0], fonte_relatorio, 12)
    analise = document.add_paragraph(
        "Argumentação antiga do modelo sobre outro processo, outro gestor e outros valores.",
        style=fundamento,
    )
    _fonte(analise.runs[0], fonte_fundamento, 11)
    prova = document.add_table(rows=2, cols=2)
    prova.cell(0, 0).text = "VALOR_RESIDUAL_MODELO"
    prova.cell(1, 0).text = "349430"
    imagem = document.add_paragraph()
    imagem.add_run().add_picture(BytesIO(PNG_A), width=Emu(120000))
    document.add_paragraph(
        "ANTE O EXPOSTO, pugna-se pelo dispositivo do processo que serviu de modelo.",
        style=relatorio,
    )
    number_id = _numbering_definition(
        document, kind=kind, marker=marker, left_twips=720, hanging_twips=360,
    )
    for text in (
        "Pedido antigo que não pode ser copiado.",
        "Outro pedido exclusivo do processo usado como modelo.",
    ):
        item = document.add_paragraph(text, style=relatorio)
        _fonte(item.runs[0], fonte_relatorio, 12)
        _apply_numbering(item, number_id)
    document.add_paragraph("")
    document.add_paragraph("")
    data = document.add_paragraph("João Pessoa, 1 de janeiro de 2020.")
    _definir_jc(data, alinhamento_data)
    document.add_paragraph("")
    nome = document.add_paragraph("NOME DE EXEMPLO DO MODELO")
    nome.runs[0].bold = True
    _definir_jc(nome, alinhamento_assinatura)
    cargo = document.add_paragraph("Cargo de exemplo do modelo")
    _definir_jc(cargo, alinhamento_assinatura)
    _definir_jc(document.add_paragraph("kacf"), "start")
    output = BytesIO()
    document.save(output)
    return nome_arquivo, output.getvalue()


def _perfil_render():
    return {
        "nome": "BRADSON TIBÉRIO LUNA CAMELO",
        "cargo": "Procurador do Ministério Público de Contas do Estado da Paraíba",
        "localidade": "João Pessoa",
    }


def test_tabela_nao_reconhecida_nao_e_apagada_em_silencio():
    document = Document()
    document.add_paragraph("Planilha antiga")
    table = document.add_table(rows=4, cols=3)
    table.cell(0, 0).text = "Valores do processo modelo"
    output = BytesIO()
    document.save(output)
    with pytest.raises(MinutasDocxError, match="inválido ou incompatível"):
        renderizar(
            output.getvalue(),
            {},
            _structure(),
            _perfil_render(),
            ["Processo TC Nº 02659/25"],
        )


def test_gabinetes_preservam_tabelas_estilos_e_dispositivo_proprios(ambiente):
    store, storage = ambiente
    bradson = _gabinete(store, storage, PREFEITURAS, "BLTC")
    elvira = _gabinete(store, storage, PREFEITURAS, "ESPO")
    modelo_bradson = _modelo_regional(
        "bradson-estrutura.docx",
        fonte_relatorio="Times New Roman",
        fonte_fundamento="Calibri",
        esquerda=3.4,
        superior=2.6,
        logo=PNG_B,
        largura=900000,
        fill="D8D8D8",
        kind="lowerLetter",
        marker="%1)",
        alinhamento_data="end",
        alinhamento_assinatura="center",
        largura_ementa=5000,
    )
    modelo_elvira = _modelo_regional(
        "elvira-estrutura.docx",
        fonte_relatorio="Georgia",
        fonte_fundamento="Arial",
        esquerda=1.9,
        superior=4.8,
        logo=PNG_A,
        largura=1400000,
        fill="1F4E79",
        kind="decimal",
        marker="%1.",
        alinhamento_data="right",
        alinhamento_assinatura="start",
        largura_ementa=3200,
    )
    MinutasModeloVisualService(bradson).save(PREFEITURAS, modelo_bradson)
    MinutasModeloVisualService(elvira).save(PREFEITURAS, modelo_elvira)

    def conferir(library, modelo, nome, cargo, fonte_relatorio, fonte_fundamento, fill, formato, jc_data, jc_cargo, largura):
        gerado = Document(BytesIO(_exportar(library).content))
        texto = _texto_completo(gerado)
        relatorio = next(
            item for item in gerado.paragraphs if item.text.startswith("Relatorio da instrucao")
        )
        fundamento = next(
            item for item in gerado.paragraphs if item.text.startswith("Fundamentacao geral")
        )
        titulo = next(
            item for item in gerado.paragraphs if item.text == "Pagamento sem cobertura"
        )
        pedido = next(
            item for item in gerado.paragraphs
            if item.text.startswith("Conferir a instrucao")
        )
        data = next(
            item for item in gerado.paragraphs
            if any(
                month in item.text.casefold()
                for month in ("janeiro", "fevereiro", "março", "abril", "maio", "junho",
                              "julho", "agosto", "setembro", "outubro", "novembro", "dezembro")
            )
        )
        assinatura = next(item for item in gerado.paragraphs if nome in item.text)
        cargo_item = next(item for item in gerado.paragraphs if cargo in item.text)
        ementa = next(table for table in gerado.tables if "EMENTA" in table.cell(0, 1).text.upper())
        faixa = next(table for table in gerado.tables if "PARECER" in _texto_tabela(table).upper())
        identificacao = next(item for item in gerado.paragraphs if item.text.startswith("Processo TC"))

        assert len(gerado.tables) == 2
        assert not ementa.cell(0, 0).text.strip()
        assert "Ementa de teste da instrucao." in ementa.cell(0, 1).text
        assert "EMENTA:" in ementa.cell(0, 1).text
        assert not any(item.text.startswith("Ementa:") for item in gerado.paragraphs)
        assert _largura_atual(ementa.cell(0, 1)) == str(largura)
        assert fill in "".join(
            str(node.get(qn("w:fill")) or "")
            for node in faixa._tbl.iter(qn("w:shd"))
        ).upper()
        assert not any(item.text.strip().casefold() == "parecer nº" for item in gerado.paragraphs)
        assert _negrito(identificacao)
        assert not _negrito(relatorio)
        assert not _negrito(fundamento)
        assert _negrito(titulo)
        assert _estilo(relatorio) != _estilo(fundamento)
        assert _fonte_run(relatorio) == fonte_relatorio
        assert _fonte_run(fundamento) == fonte_fundamento
        assert _num_id(pedido)
        assert _formato_numero(gerado, _num_id(pedido)) == formato
        assert _mapa_numeracao(gerado) == _mapa_numeracao(Document(BytesIO(modelo[1])))
        assert _jc(data) == jc_data
        assert _jc(assinatura) == jc_cargo
        assert _jc(cargo_item) == jc_cargo
        assert _negrito(assinatura)
        assert not _negrito(cargo_item)
        assert len(list(gerado.element.body.iter(qn("w:drawing")))) == 0
        assert _logotipos(_exportar(library).content) == [logo_bytes(library)]
        assert VENENO not in texto
        assert "kacf" not in texto
        assert "VALOR_RESIDUAL_MODELO" not in texto
        assert "00000/00" not in texto
        assert "NOME DE EXEMPLO" not in texto
        assert "Pedido antigo" not in texto
        assert "02659/25" in texto
        assert nome in texto

    def logo_bytes(library):
        return PNG_B if library is bradson else PNG_A

    conferir(
        bradson, modelo_bradson, "BRADSON TIBÉRIO LUNA CAMELO",
        "Procurador do Ministério Público de Contas do Estado da Paraíba",
        "Times New Roman", "Calibri", "D8D8D8", "lowerLetter", "end", "center", 5000,
    )
    conferir(
        elvira, modelo_elvira, "ELVIRA SAMARA PEREIRA DE OLIVEIRA",
        "Procuradora-Geral do Ministério Público de Contas do Estado da Paraíba",
        "Georgia", "Arial", "1F4E79", "decimal", "right", "start", 3200,
    )
    texto_bradson = _texto_completo(Document(BytesIO(_exportar(bradson).content)))
    texto_elvira = _texto_completo(Document(BytesIO(_exportar(elvira).content)))
    assert "ELVIRA" not in texto_bradson
    assert "BRADSON" not in texto_elvira
    assert "Minuta Procuradora" not in _pacote(_exportar(elvira).content, "word/document.xml")


@pytest.mark.skipif(not OFICIAL_ASSESSOR.is_file(), reason="amostra local do modelo do assessor ausente")
def test_minuta_oficial_do_assessor_preserva_regioes_visuais():
    estrutura = _structure()
    estrutura["pedidos"] = [
        "Primeiro pedido novo da minuta.",
        "Segundo pedido novo da minuta.",
        "Terceiro pedido novo da minuta.",
        "Quarto pedido novo da minuta.",
    ]
    alvo = {
        "numero_processo": "02659/25",
        "origem": "Prefeitura Municipal de Piloto",
        "responsavel": "Gestor de Teste",
        "exercicio": "2025",
        "data_parecer": "2026-10-20",
    }
    gerado = Document(BytesIO(renderizar(
        OFICIAL_ASSESSOR.read_bytes(),
        alvo,
        estrutura,
        _perfil_render(),
        _draft_header_lines(alvo),
    )))
    texto = _texto_completo(gerado)
    ementa = next(table for table in gerado.tables if "EMENTA" in _texto_tabela(table).upper())
    faixa = next(table for table in gerado.tables if "PARECER" in _texto_tabela(table).upper())
    relatorio = next(item for item in gerado.paragraphs if "Relatorio da instrucao" in item.text)
    fundamento = next(item for item in gerado.paragraphs if "Fundamentacao geral" in item.text)
    titulo = next(item for item in gerado.paragraphs if item.text == "Pagamento sem cobertura")
    pedidos = [
        item for item in gerado.paragraphs
        if item.text.endswith("pedido novo da minuta.")
    ]
    identificacao = next(item for item in gerado.paragraphs if item.text.startswith("Processo TC"))
    data = next(item for item in gerado.paragraphs if "20 de outubro de 2026" in item.text)
    nome = next(item for item in gerado.paragraphs if "BRADSON TIBÉRIO LUNA CAMELO" in item.text)
    cargo = next(item for item in gerado.paragraphs if "Procurador do Ministério Público de Contas" in item.text)
    origem = Document(OFICIAL_ASSESSOR)

    assert len(gerado.tables) == 2
    assert len(origem.tables) == 2
    assert not ementa.cell(0, 0).text.strip()
    assert ementa.cell(0, 1).text.startswith("EMENTA:")
    assert "Ementa de teste da instrucao." in ementa.cell(0, 1).text
    assert not any("Ementa de teste da instrucao." in item.text for item in gerado.paragraphs)
    assert _largura_atual(ementa.cell(0, 1)) == _largura_atual(origem.tables[0].cell(0, 1))
    assert "D8D8D8" in "".join(
        str(node.get(qn("w:fill")) or "") for node in faixa._tbl.iter(qn("w:shd"))
    ).upper()
    assert "PARECER" in _texto_tabela(faixa).upper()
    assert not any(item.text.strip().casefold() in {"parecer nº", "parecer n°"} for item in gerado.paragraphs)
    assert _negrito(identificacao)
    assert "02659/25" in identificacao.text
    assert any(
        item.text == "Unidade Jurisdicionada: Prefeitura Municipal de Piloto"
        for item in gerado.paragraphs
    )
    assert any(item.text.startswith("Exercício:") and item.text.endswith("2025") for item in gerado.paragraphs)
    assert not any(item.text.startswith("Origem:") for item in gerado.paragraphs)
    assert not _negrito(relatorio)
    assert not _negrito(fundamento)
    assert _negrito(titulo)
    assert _estilo(relatorio) and _estilo(relatorio) != _estilo(fundamento)
    assert len(pedidos) == 4
    assert len({_num_id(item) for item in pedidos}) == 1
    assert _formato_numero(gerado, _num_id(pedidos[0])) == "lowerLetter"
    assert _mapa_numeracao(gerado) == _mapa_numeracao(origem)
    assert _jc(data) == "end"
    assert _jc(nome) == "center"
    assert _jc(cargo) == "center"
    assert _negrito(nome)
    assert not _negrito(cargo)
    assert nome._p.find(qn("w:pPr")).find(qn("w:keepNext")) is not None
    assert "kacf" not in texto
    assert "02559" not in texto
    assert "FREI MARTINHO" not in texto.upper()
    assert "Pedido antigo" not in texto
    assert len(list(gerado.element.body.iter(qn("w:drawing")))) == 0
    assert len(list(origem.element.body.iter(qn("w:drawing")))) > 0
    assert gerado.sections[0].left_margin == origem.sections[0].left_margin
    assert gerado.sections[0].top_margin == origem.sections[0].top_margin
