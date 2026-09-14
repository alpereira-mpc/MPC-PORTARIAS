from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import hashlib
import tempfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, Twips

from database.memorandos import MemorandosStore
from database.store import Store
from document_generator.memorandos import generate
from services.memorandos import STATUS_FINALIZADO, cabinet_text, digest, fingerprint, validate
from services.memorandos import (
    SERVER_SEARCH_LIMIT,
    filter_servers,
    prefix_search,
    read_server_xlsx,
    search_key,
    signature_role,
)
from tests.test_postgresql import pg_store, pg_url  # noqa: F401


def person(name, gender, registration, cargo="Analista", lotacao="Secretaria da 1ª Câmara"):
    return {
        "nome": name,
        "genero": gender,
        "matricula": registration,
        "cargo": cargo,
        "lotacao": lotacao,
    }


def record(**changes):
    today = date.today()
    payload = {
        "tipo": "SUBSTITUICAO",
        "data_inicio": today.isoformat(),
        "data_fim": (today + timedelta(days=1)).isoformat(),
        "natureza_funcao": "Cargo comissionado",
        "motivo": "Licença especial",
        "motivo_texto": "",
        "gabinete_snapshot": "lotada no gabinete da Procuradoria-Geral",
        "signatario_id": 1,
        "signatario_nome": "Elvira Samara Pereira de Oliveira",
        "signatario_cargo": "Procuradora-Geral",
        "etapas": [
            {
                "substituido": person("Ana", "Feminino", "1", "A"),
                "substituto": person("Beto", "Masculino", "2", "B"),
            }
        ],
    }
    payload.update(changes)
    return payload


def test_server_import_upsert_and_invalid_rows(store):
    service = MemorandosStore(store)
    rows = [
        {"nome": "Ana", "matricula": "100", "cargo": "A", "setor": "1CAM"},
        {"nome": "Sem matrícula", "matricula": "", "cargo": "B", "setor": "X"},
    ]
    first = service.import_servers(
        rows, actor_email="admin@test", filename="x.xlsx", content_hash="a", administrator=True
    )
    assert first["incluidos"] == 1 and first["inconsistentes"] == 1
    service.import_servers(
        [{"nome": "Ana Atualizada", "matricula": "100", "cargo": "Novo", "setor": "2CAM"}],
        actor_email="admin@test",
        filename="x.xlsx",
        content_hash="b",
        administrator=True,
    )
    found = service.servers("100")
    assert len(found) == 1 and found[0]["cargo"] == "Novo"


def test_server_search_is_strict_name_prefix_or_numeric_registration():
    people = [
        {"id": 1, "nome": "José da Silva", "matricula": "4455", "setor": "Protocolo"},
        {"id": 2, "nome": "Silvana Costa", "matricula": "100", "setor": "Gabinete da Procuradoria"},
        {"id": 3, "nome": "Ana Clara", "matricula": "200", "setor": "Secretaria"},
        {"id": 4, "nome": "Carlos Souza", "matricula": "300", "setor": "Arquivo"},
    ]
    assert [row["id"] for row in filter_servers(people, "JOSE")] == [1]
    assert [row["id"] for row in filter_servers(people, "jose")] == [1]
    assert [row["id"] for row in filter_servers(people, "silva")] == [2]
    assert [row["id"] for row in filter_servers(people, "jose da")] == [1]
    assert filter_servers(people, "lara") == []
    assert [row["id"] for row in filter_servers(people, "4455")] == [1]
    assert filter_servers(people, "gabinete") == []
    assert filter_servers(people, "   ") == []
    crowded = [{"id": i, "nome": f"Ana {i:02d}", "matricula": str(1000 + i), "setor": "X"} for i in range(25)]
    assert len(filter_servers(crowded, "ana")) == SERVER_SEARCH_LIMIT
    hits, truncated = prefix_search(crowded, "ana")
    assert truncated and len(hits) == SERVER_SEARCH_LIMIT


def test_niltamir_progressive_prefix_and_registration():
    people = [
        {"id": 1, "nome": "Ana Costa", "matricula": "1", "setor": "SECRETARIA"},
        {"id": 2, "nome": "Niltamir Galdino Guedes", "matricula": "3702391", "setor": "PROGE"},
        {"id": 3, "nome": "Nelson Silva", "matricula": "99", "setor": "PROGE"},
        {"id": 4, "nome": "Bruno Nilo", "matricula": "2", "setor": "X"},
        {"id": 2, "nome": "Niltamir Galdino Guedes", "matricula": "3702391", "setor": "PROGE"},
    ]
    by_n = filter_servers(people, "N")
    assert [row["id"] for row in by_n] == [3, 2]
    assert [row["id"] for row in filter_servers(people, "NI")] == [2]
    assert [row["id"] for row in filter_servers(people, "NIL")] == [2]
    assert [row["id"] for row in filter_servers(people, "NILT")] == [2]
    assert [row["id"] for row in filter_servers(people, "NILTA")] == [2]
    assert [row["id"] for row in filter_servers(people, "nilta")] == [2]
    assert filter_servers(people, "zzz") == []
    assert [row["id"] for row in filter_servers(people, "37023")] == [2]
    assert [row["id"] for row in filter_servers(people, "3702391")] == [2]
    assert filter_servers(people, "PROGE") == []
    from services.memorandos_ui import _label

    assert _label({"nome": "Niltamir Galdino Guedes", "matricula_original": "3702391", "setor": "PROGE"}) == (
        "Niltamir Galdino Guedes — 3702391 — PROGE"
    )
    assert search_key("  NILT  ") == "nilt"


def test_server_administration_requires_administrator(store):
    service = MemorandosStore(store)
    with pytest.raises(ValueError, match="não autorizado"):
        service.import_servers(
            [], actor_email="user@test", filename="x.xlsx", content_hash="a", administrator=False
        )
    with pytest.raises(ValueError, match="não autorizado"):
        service.update_server(1, {"nome": "X"}, actor_email="user@test", administrator=False)
    imported = service.import_servers(
        [{"nome": "Ana", "matricula": "100", "cargo": "A", "setor": "X"}],
        actor_email="admin@test",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    assert imported["incluidos"] == 1
    server_id = service.servers("100")[0]["id"]
    service.update_server(
        server_id,
        {"nome": "Ana", "cargo": "B", "setor": "X", "genero": "Feminino", "ativo": True},
        actor_email="admin@test",
        administrator=True,
    )
    assert service.servers("100")[0]["cargo"] == "B"
    with pytest.raises(ValueError, match="não autorizado"):
        service.update_server(
            server_id,
            {"nome": "Ana", "cargo": "Hack", "setor": "X", "genero": "Feminino", "ativo": False},
            actor_email="user@test",
            administrator=False,
        )
    assert service.servers("100")[0]["cargo"] == "B"


def test_synthetic_xlsx_reports_duplicate_zero_and_blank():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Nome", "Matrícula", "Cargo", "Setor"])
    sheet.append(["Ana", "100", "A", "1CAM"])
    sheet.append(["Ana repetida", "100", "B", "2CAM"])
    sheet.append(["Zero", "0", "C", "X"])
    sheet.append(["Vazia", "", "D", "Y"])
    sheet.append(["Mesmo nome", "200", "E", "Z"])
    out = BytesIO()
    book.save(out)
    rows = read_server_xlsx(out.getvalue())
    service = MemorandosStore(Store(path=Path(tempfile.mkdtemp()) / "memo.db"))
    report = service.preview_import(rows)
    assert report["novos"] == 2 and len(report["duplicidades"]) == 1
    assert len(report["matriculas_zero"]) == 1 and len(report["sem_matricula"]) == 1


def test_draft_snapshot_preview_final_file_and_official_number(store):
    service = MemorandosStore(store)
    payload = record()
    identifier = service.save_draft(payload, actor_email="author@test")
    draft = service.get(identifier)
    assert draft["status"] == "RASCUNHO" and draft["situacao"] == "EM ANDAMENTO"
    pdf_bytes = b"%PDF- synthetic"
    digest = hashlib.sha256(pdf_bytes).hexdigest()
    service.finalize(
        identifier,
        preview_hash=fingerprint(payload),
        pdf_bytes=pdf_bytes,
        filename="memo.pdf",
    )
    assert service.get(identifier)["status"] == STATUS_FINALIZADO
    assert service.file(identifier) == ("memo.pdf", pdf_bytes)
    assert hashlib.sha256(service.file(identifier)[1]).hexdigest() == digest
    service.set_official_number(identifier, "123/2026", "author@test")
    assert service.get(identifier)["numero_oficial"] == "123/2026"
    assert service.file(identifier)[1] == pdf_bytes
    with pytest.raises(ValueError):
        service.save_draft(payload, actor_email="author@test", identifier=identifier)


def test_preview_fingerprint_is_invalidated_after_change(store):
    payload = record()
    first = fingerprint(payload)
    payload["motivo"] = "Férias regulamentares"
    assert fingerprint(payload) != first
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    with pytest.raises(ValueError, match="prévia"):
        MemorandosStore(store).finalize(
            identifier, preview_hash=first, pdf_bytes=b"%PDF- x", filename="x.pdf"
        )


def test_self_substitution_and_broken_chain_are_rejected():
    payload = record()
    payload["etapas"][0]["substituto"] = dict(payload["etapas"][0]["substituido"])
    with pytest.raises(ValueError, match="cadeia|repetida|autossubstituição"):
        validate(payload)
    payload = record()
    payload["etapas"].append(
        {
            "substituido": person("Outra", "Feminino", "9"),
            "substituto": person("Caio", "Masculino", "3"),
        }
    )
    with pytest.raises(ValueError, match="quebrada"):
        validate(payload)


def test_cascade_cycle_is_rejected(store):
    payload = record()
    payload["etapas"].append(
        {
            "substituido": payload["etapas"][0]["substituto"],
            "substituto": payload["etapas"][0]["substituido"],
        }
    )
    with pytest.raises(ValueError, match="cadeia"):
        MemorandosStore(store).save_draft(payload, actor_email="author@test")


def test_two_and_three_stage_cascades_are_valid(store):
    payload = record()
    payload["etapas"].append(
        {
            "substituido": dict(payload["etapas"][0]["substituto"]),
            "substituto": person("Caio", "Masculino", "3", "C"),
        }
    )
    validate(payload)
    payload["etapas"].append(
        {
            "substituido": dict(payload["etapas"][1]["substituto"]),
            "substituto": person("Dora", "Feminino", "4", "D"),
        }
    )
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    assert MemorandosStore(store).get(identifier)["etapas"][2]["substituido"]["nome"] == "Caio"


def test_general_prosecution_cabinet_rule():
    elvira = {"nome": "Elvira Samara Pereira de Oliveira", "genero": "feminino"}
    assert cabinet_text(elvira, "Feminino") == "lotada no gabinete da Procuradoria-Geral"
    assert cabinet_text(elvira, "Masculino") == "lotado no gabinete da Procuradoria-Geral"
    other = {"nome": "Luciano Andrade Farias", "genero": "masculino"}
    assert cabinet_text(other, "Feminino") == "lotada no gabinete do Procurador Luciano Andrade Farias"


def test_signature_role_uses_catalog_gender_not_the_name():
    elvira = {"nome": "Elvira Samara Pereira de Oliveira", "genero": "feminino", "funcao": "Procurador-Geral"}
    assert signature_role(elvira) == "Procuradora-Geral"
    man = {"nome": "Alguém", "genero": "masculino", "funcao": "Procurador-Geral"}
    assert signature_role(man) == "Procurador-Geral"


def test_editable_snapshot_does_not_change_master(store):
    service = MemorandosStore(store)
    service.import_servers(
        [{"nome": "Ana", "matricula": "100", "cargo": "Efetivo", "setor": "1CAM"}],
        actor_email="admin@test",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    payload = record()
    payload["etapas"][0]["substituido"]["cargo"] = "Comissionado"
    MemorandosStore(store).save_draft(payload, actor_email="author@test")
    assert service.servers("100")[0]["cargo"] == "Efetivo"


def test_draft_can_be_deleted_final_cannot(store):
    service = MemorandosStore(store)
    identifier = service.save_draft(record(), actor_email="author@test")
    service.delete_draft(identifier)
    with pytest.raises(ValueError, match="não encontrado"):
        service.get(identifier)
    identifier = service.save_draft(record(), actor_email="author@test")
    service.finalize(
        identifier, preview_hash=fingerprint(record()), pdf_bytes=b"%PDF- x", filename="x.pdf"
    )
    with pytest.raises(ValueError, match="rascunho"):
        service.delete_draft(identifier)


def test_list_filters_and_pagination(store):
    service = MemorandosStore(store)
    first = service.save_draft(record(), actor_email="author@test")
    second = record()
    second["etapas"][0]["substituido"] = person("Carla", "Feminino", "8")
    service.save_draft(second, actor_email="author@test")
    rows = service.list(search="Carla", limit=1, offset=0)
    assert len(rows) == 1 and "Carla" in rows[0]["cadeia"]
    assert service.get(first)["etapas"][0]["substituido"]["nome"] == "Ana"


def test_document_contains_institutional_fields():
    payload = record(
        gabinete_snapshot="lotada no gabinete da Procuradoria-Geral",
        etapas=[
            {
                "substituido": person("Niltamir Galdino Guedes", "Masculino", "10", "Chefe"),
                "substituto": person("Ana Claudia da Costa Ferreira", "Feminino", "20", "Assessora", "Secretaria da 2ª Câmara"),
            }
        ],
    )
    content = generate(payload)
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    section = doc.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.05
    assert abs(section.left_margin.cm - 3.0) < 0.05
    assert abs(section.right_margin.cm - 2.0) < 0.05
    title = next(p for p in doc.paragraphs if p.text == "MEMORANDO")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.paragraph_format.left_indent == Twips(-567)
    assert title.runs[0].bold and title.runs[0].font.size == Pt(14)
    assert title.runs[0].font.name == "Times New Roman"
    dest = next(p for p in doc.paragraphs if "Excelentíssimo" in p.text)
    assert int(dest.paragraph_format.left_indent or 0) == 0
    assert title.paragraph_format.space_before == Pt(18)
    assert title.paragraph_format.space_after == Pt(18)
    body = next(p for p in doc.paragraphs if p.text.startswith("Ao cumprimentá-lo"))
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(body.paragraph_format.first_line_indent.cm - 1.25) < 0.02
    assert body.runs[0].font.size == Pt(13)
    assert body.runs[0].font.name == "Times New Roman"
    subject = next(p for p in doc.paragraphs if p.text.startswith("Assunto:"))
    assert not subject.runs[0].bold
    assert subject.runs[0].text == "Assunto:"
    assert subject.runs[1].bold and "Substituição" in subject.runs[1].text
    closing = next(p for p in doc.paragraphs if p.text.startswith("Com os meus melhores"))
    assert closing.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
    assert (closing.paragraph_format.first_line_indent or Pt(0)).pt == 0
    signature = next(p for p in doc.paragraphs if p.text.startswith("ELVIRA SAMARA"))
    assert signature.alignment == WD_ALIGN_PARAGRAPH.CENTER and signature.runs[0].bold
    role = next(p for p in doc.paragraphs if "Procuradora-Geral do Ministério" in p.text)
    assert role.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert not role.runs[0].bold
    assert role.runs[0].font.size == Pt(13)
    bold = "".join(run.text for run in body.runs if run.bold)
    plain = "".join(run.text for run in body.runs if not run.bold)
    assert "Niltamir Galdino Guedes" in bold
    assert "matrícula nº 10" in bold
    assert "Ana Claudia da Costa Ferreira" in bold
    assert "matrícula nº 20" in bold
    from services.memorandos import short_date
    assert short_date(payload["data_inicio"]) in bold
    assert short_date(payload["data_fim"]) in bold
    assert "Chefe" in plain and "gabinete da Procuradoria-Geral" in plain
    for item in (
        "MEMORANDO",
        "Presidente do Tribunal de Contas",
        "Senhor Presidente",
        "Niltamir Galdino Guedes",
        "Ana Claudia da Costa Ferreira",
        "gabinete da Procuradoria-Geral",
        "ELVIRA SAMARA",
        "Procuradora-Geral",
        "Assunto:",
    ):
        assert item in text
    assert "gabinete da Procuradora Elvira" not in text
    header_xml = doc.sections[0].header._element.xml
    assert "a:blip" not in header_xml and "v:imagedata" not in header_xml


def test_memorandos_embed_the_same_logo_as_portarias():
    from zipfile import ZipFile
    from database.store import ROOT as PROJECT_ROOT
    from document_generator.memorandos import PORTARIAS_LOGO, PORTARIAS_LOGO_HEIGHT, PORTARIAS_LOGO_WIDTH

    portarias_jpeg = ZipFile(PROJECT_ROOT / "templates" / "simples.docx").read("word/media/image1.jpeg")
    assert PORTARIAS_LOGO.is_file()
    assert PORTARIAS_LOGO.read_bytes() == portarias_jpeg
    content = generate(record())
    with ZipFile(BytesIO(content)) as package:
        media = [name for name in package.namelist() if name.startswith("word/media/")]
        assert len(media) == 1
        assert package.read(media[0]) == portarias_jpeg
        document = package.read("word/document.xml").decode("utf-8")
        headers = [name for name in package.namelist() if "header" in name and name.endswith(".xml")]
        for name in headers:
            xml = package.read(name)
            assert b"a:blip" not in xml and b"v:imagedata" not in xml
    assert str(PORTARIAS_LOGO_WIDTH.emu) in document
    assert str(PORTARIAS_LOGO_HEIGHT.emu) in document
    doc = Document(BytesIO(content))
    logo_paragraph = doc.paragraphs[0]
    assert logo_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert logo_paragraph.paragraph_format.left_indent == Twips(-567)
    assert int(logo_paragraph.paragraph_format.right_indent or 0) == 0
    assert PORTARIAS_LOGO_WIDTH == Pt(163.5)
    assert PORTARIAS_LOGO_HEIGHT == Pt(105.75)


def test_memorando_title_pdf_center_is_within_three_points():
    from document_generator.memorandos import pdf
    from document_generator.pdf import PdfUnavailable
    from pypdf import PdfReader

    try:
        content = pdf(record())
    except (PdfUnavailable, RuntimeError) as exc:
        pytest.skip(str(exc))
    page = PdfReader(BytesIO(content)).pages[0]
    page_center = float(page.mediabox.width) / 2
    found = []

    def visitor(text, cm, tm, font_dict, font_size):
        if (text or "").strip() == "MEMORANDO":
            found.append((float(tm[4]), float(font_size or 14)))

    page.extract_text(visitor_text=visitor)
    assert found, "não foi possível localizar MEMORANDO no PDF"
    start, size = found[0]
    # Measured glyph width of MEMORANDO in Times 14 pt on the previous render.
    width = 97.93 * (size / 14.0)
    center = start + width / 2
    assert abs(center - page_center) <= 3


def test_cascade_document_mentions_each_stage():
    payload = record(
        etapas=[
            {
                "substituido": person("Kátia Cilene Brandão Antunes", "Feminino", "11"),
                "substituto": person("Maria da Luz de Lima", "Feminino", "12"),
            },
            {
                "substituido": person("Maria da Luz de Lima", "Feminino", "12"),
                "substituto": person("Ana Claudia da Costa Ferreira", "Feminino", "13"),
            },
        ]
    )
    text = "\n".join(p.text for p in Document(BytesIO(generate(payload))).paragraphs)
    assert "Em decorrência da substituição acima" in text
    assert "Kátia Cilene Brandão Antunes" in text
    assert "Maria da Luz de Lima" in text
    assert "Ana Claudia da Costa Ferreira" in text
    cascade = next(p for p in Document(BytesIO(generate(payload))).paragraphs if p.text.startswith("Em decorrência"))
    bold = "".join(run.text for run in cascade.runs if run.bold)
    assert "Maria da Luz de Lima" in bold
    assert "Ana Claudia da Costa Ferreira" in bold
    assert "matrícula nº 12" in bold
    assert "matrícula nº 13" in bold


def test_real_docx_pdf_roundtrip_and_stored_bytes(store):
    from document_generator.memorandos import pdf
    from document_generator.pdf import PdfUnavailable
    from pypdf import PdfReader

    payload = record(
        etapas=[
            {
                "substituido": person("Niltamir Galdino Guedes", "Masculino", "10"),
                "substituto": person("Ana Claudia da Costa Ferreira", "Feminino", "20"),
            }
        ]
    )
    try:
        content = pdf(payload)
    except (PdfUnavailable, RuntimeError) as exc:
        pytest.skip(str(exc))
    assert content.startswith(b"%PDF-")
    extracted = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    assert "MEMORANDO" in extracted
    assert "Niltamir" in extracted
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    MemorandosStore(store).finalize(
        identifier, preview_hash=fingerprint(payload), pdf_bytes=content, filename="memo.pdf"
    )
    name, stored = MemorandosStore(store).file(identifier)
    assert name == "memo.pdf" and stored == content
    assert hashlib.sha256(stored).hexdigest() == hashlib.sha256(content).hexdigest()
    MemorandosStore(store).set_official_number(identifier, "99/2026", "author@test")
    assert MemorandosStore(store).file(identifier)[1] == content


def test_new_form_starts_without_selected_participants():
    import inspect
    from services import memorandos_ui as ui

    pick = inspect.getsource(ui._pick)
    chain = inspect.getsource(ui._chain)
    source = pick + chain + inspect.getsource(ui._editor)
    assert ui.LABEL_REPLACEMENT == "Substituto(a)"
    assert "Primeiro substituto" not in source
    assert "Novo substituto" not in chain
    assert "index=None" in pick
    assert 'filter_mode="prefix"' in pick
    assert "st.selectbox" in pick
    assert "st.text_input" not in pick
    assert "st.radio" not in pick
    assert "prefix_search" not in pick
    assert "options[0]" not in pick
    assert "people[0]" not in source
    assert "by_id" in pick
    assert pick.count("_pick(") >= 1
    assert chain.count("_pick(") == 3
    assert ui.PLACEHOLDER_AWAY == "Selecione o servidor afastado"
    assert ui.PLACEHOLDER_REPLACEMENT == "Selecione o substituto(a)"
    assert "disabled=not complete" in chain
    assert chain.index('expander("Etapa 1"') < chain.index("memorando_away")
    assert chain.index(f'**{{LABEL_AWAY}}**') < chain.index(f'**{{LABEL_REPLACEMENT}}**')


def test_new_memorandum_ui_has_empty_participant_selects(store, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from tests.access_testing import enable_login
    from database.store import ROOT

    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    enable_login(monkeypatch, store)
    monkeypatch.setattr("database.store.Store", lambda: store)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    app.button(key="open_memorandos").click().run()
    app.radio(key="memorandos_nav").set_value("Novo Memorando").run()
    assert not app.exception
    away = next(x for x in app.selectbox if x.label == "Servidor afastado")
    replacement = next(x for x in app.selectbox if x.label == "Substituto(a)")
    assert away.value is None
    assert replacement.value is None
    assert away.options and replacement.options
    assert not any(x.label == "Primeiro substituto" for x in app.selectbox)
    assert not any(x.label == "Cargo/função documental" for x in app.text_input)
    assert not any(x.label == "Pré-visualizar PDF" for x in app.button)
    assert not any(x.label == "Gerar prévia" for x in app.button)
    assert not any(x.label == "Finalizar memorando" for x in app.button)
    cascade = next(x for x in app.button if x.label == "Adicionar substituição em cascata")
    assert cascade.disabled


def test_niltamir_picker_keeps_native_select_and_selection_by_id(store, monkeypatch):
    import streamlit as st
    from services.memorandos_ui import NAV_KEY, NAV_NEW, _label

    assert tuple(int(part) for part in st.__version__.split(".")[:2]) >= (1, 56)
    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana Costa", "matricula": "1", "cargo": "A", "setor": "SECRETARIA"},
            {"nome": "Niltamir Galdino Guedes", "matricula": "3702391", "cargo": "Chefe", "setor": "PROGE"},
            {"nome": "Nelson Silva", "matricula": "99", "cargo": "B", "setor": "PROGE"},
            {"nome": "Bruno Nilo", "matricula": "2", "cargo": "C", "setor": "X"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="niltamir",
        administrator=True,
    )
    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    people = MemorandosStore(store).servers()
    niltamir = next(row for row in people if row["nome"].startswith("Niltamir"))
    nelson = next(row for row in people if row["nome"].startswith("Nelson"))
    bruno = next(row for row in people if row["nome"].startswith("Bruno"))
    ana = next(row for row in people if row["nome"].startswith("Ana"))
    label = _label(niltamir)
    assert label == "Niltamir Galdino Guedes — 3702391 — PROGE"
    away = next(x for x in app.selectbox if x.label == "Servidor afastado")
    replacement = next(x for x in app.selectbox if x.label == "Substituto(a)")
    assert away.value is None and replacement.value is None
    assert label in away.options
    assert _label(nelson) in away.options
    assert _label(bruno) in away.options
    away.set_value(niltamir["id"]).run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value == niltamir["id"]
    assert next(x for x in app.selectbox if x.label == "Substituto(a)").value is None
    next(x for x in app.selectbox if x.label == "Servidor afastado").set_value(None).run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value is None
    next(x for x in app.selectbox if x.label == "Servidor afastado").set_value(niltamir["id"]).run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value == niltamir["id"]
    next(x for x in app.selectbox if x.label == "Substituto(a)").set_value(ana["id"]).run()
    assert next(x for x in app.selectbox if x.label == "Substituto(a)").value == ana["id"]
    next(x for x in app.button if x.label == "Adicionar substituição em cascata").click().run()
    cascade = [
        box for box in app.selectbox
        if box.label == "Substituto(a)" and box.key == "memorando_cascade_1"
    ]
    assert cascade and cascade[0].value is None
    cascade[0].set_value(nelson["id"]).run()
    assert next(x for x in app.selectbox if x.key == "memorando_cascade_1").value == nelson["id"]


def test_system_preview_can_be_finalized_without_any_docx_upload(store):
    from services.memorandos import SOURCE_SYSTEM, preview_is_valid

    payload = record()
    preview = {
        "source": SOURCE_SYSTEM,
        "fingerprint": fingerprint(payload),
        "pdf": b"%PDF- generated",
        "pdf_name": "auto.pdf",
        "docx": b"PK-generated",
        "docx_name": "auto.docx",
    }
    assert preview_is_valid(preview, payload)
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    MemorandosStore(store).finalize(
        identifier,
        preview_hash=fingerprint(payload),
        pdf_bytes=preview["pdf"],
        filename="auto.pdf",
        docx_bytes=preview["docx"],
        docx_filename="auto.docx",
        uploaded=False,
    )
    assert MemorandosStore(store).file(identifier)[1] == b"%PDF- generated"


def test_system_finalize_does_not_require_docx_bytes(store):
    payload = record()
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    MemorandosStore(store).finalize(
        identifier,
        preview_hash=fingerprint(payload),
        pdf_bytes=b"%PDF- only-system",
        filename="only.pdf",
        uploaded=False,
    )
    assert MemorandosStore(store).file(identifier) == ("only.pdf", b"%PDF- only-system")


def test_upload_is_optional_and_does_not_block_system_preview():
    import inspect
    from services import memorandos_ui as ui
    from services.memorandos import preview_is_valid, SOURCE_SYSTEM

    payload = record()
    preview = {"source": SOURCE_SYSTEM, "fingerprint": fingerprint(payload), "pdf": b"%PDF- x"}
    assert preview_is_valid(preview, payload)
    assert "upload_bytes" not in inspect.getsource(ui._finalize_active)
    assert "file_uploader" not in inspect.getsource(ui._finalize_active)
    documents = inspect.getsource(ui._documents)
    assert "Opcional:" in documents
    assert "memo_finalize_system" in documents


def test_system_preview_docx_is_the_pdf_source(monkeypatch):
    from document_generator.memorandos import preview_documents

    seen = []

    def fake_convert(content, engine="auto"):
        seen.append(content)
        return b"%PDF-1.4 from-docx", "test"

    monkeypatch.setattr("document_generator.pdf.convert", fake_convert)
    docx, pdf_bytes = preview_documents(record())
    assert docx[:2] == b"PK"
    assert seen == [docx]
    assert pdf_bytes == b"%PDF-1.4 from-docx"


def test_uploaded_docx_is_validated_and_converted(monkeypatch):
    from document_generator.memorandos import generate
    from services.memorandos import MIME_DOCX, preview_is_valid, upload_fingerprint, validate_docx

    payload = record()
    docx = generate(payload)
    name, mime = validate_docx("memorando.docx", docx)
    assert name.endswith(".docx") and mime == MIME_DOCX
    seen = []

    def fake_convert(content, engine="auto"):
        seen.append(content)
        return b"%PDF-1.4 uploaded", "test"

    monkeypatch.setattr("document_generator.pdf.convert", fake_convert)
    from document_generator.pdf import convert

    pdf_bytes, _ = convert(docx)
    assert seen == [docx]
    preview = {
        "source": "DOCX_ENVIADO",
        "fingerprint": upload_fingerprint(payload, docx, pdf_bytes),
        "docx": docx,
        "pdf": pdf_bytes,
    }
    assert preview_is_valid(preview, payload)
    with pytest.raises(ValueError):
        validate_docx("memorando.docm", docx)
    with pytest.raises(ValueError):
        validate_docx("memorando.docx", b"")
    with pytest.raises(ValueError):
        validate_docx("memorando.docx", b"PK\x03\x04not-a-document")


def test_swapping_uploaded_docx_invalidates_previous_pdf():
    from document_generator.memorandos import generate
    from services.memorandos import preview_is_valid, upload_fingerprint

    payload = record()
    first = generate(payload)
    second = generate({**payload, "motivo": "Férias regulamentares"})
    assert first != second
    pdf_bytes = b"%PDF-1.4 first"
    preview = {
        "source": "DOCX_ENVIADO",
        "fingerprint": upload_fingerprint(payload, first, pdf_bytes),
        "docx": first,
        "pdf": pdf_bytes,
    }
    assert preview_is_valid(preview, payload)
    preview["docx"] = second
    assert not preview_is_valid(preview, payload)
    assert digest(first) != digest(second)


def test_form_change_invalidates_system_and_uploaded_preview():
    from services.memorandos import preview_is_valid, upload_fingerprint

    payload = record()
    system = {"source": "GERADA_PELO_SISTEMA", "fingerprint": fingerprint(payload), "pdf": b"%PDF- x", "docx": b"PK"}
    assert preview_is_valid(system, payload)
    changed = {**payload, "motivo": "Férias regulamentares"}
    assert not preview_is_valid(system, changed)
    uploaded = {
        "source": "DOCX_ENVIADO",
        "fingerprint": upload_fingerprint(payload, b"PK-docx", b"%PDF- y"),
        "docx": b"PK-docx",
        "pdf": b"%PDF- y",
    }
    assert preview_is_valid(uploaded, payload)
    assert not preview_is_valid(uploaded, changed)


def test_finalize_uses_latest_preview_and_keeps_pdf_immutable(store):
    from document_generator.memorandos import generate
    from services.memorandos import MIME_DOCX, upload_fingerprint

    service = MemorandosStore(store)
    payload = record()
    identifier = service.save_draft(payload, actor_email="author@test")
    stale = b"%PDF- stale"
    active = b"%PDF- active"
    docx = generate(payload)
    service.finalize(
        identifier,
        preview_hash=upload_fingerprint(payload, docx, active),
        pdf_bytes=active,
        filename="final.pdf",
        docx_bytes=docx,
        docx_filename="final.docx",
        uploaded=True,
    )
    assert service.file(identifier)[1] == active
    assert service.file(identifier)[1] != stale
    assert service.file(identifier, MIME_DOCX)[1] == docx
    service.set_official_number(identifier, "1/2026", "author@test")
    assert service.file(identifier)[1] == active
    assert service.file(identifier, MIME_DOCX)[1] == docx


def test_legacy_history_without_docx_still_downloads_pdf(store):
    from services.memorandos import MIME_DOCX

    service = MemorandosStore(store)
    payload = record()
    identifier = service.save_draft(payload, actor_email="author@test")
    service.finalize(
        identifier,
        preview_hash=fingerprint(payload),
        pdf_bytes=b"%PDF- only",
        filename="only.pdf",
    )
    assert service.file(identifier) == ("only.pdf", b"%PDF- only")
    assert service.file(identifier, MIME_DOCX) is None


def test_system_finalize_stores_docx_and_pdf(store):
    from document_generator.memorandos import generate
    from services.memorandos import MIME_DOCX

    payload = record()
    docx = generate(payload)
    pdf_bytes = b"%PDF- system"
    identifier = MemorandosStore(store).save_draft(payload, actor_email="author@test")
    MemorandosStore(store).finalize(
        identifier,
        preview_hash=fingerprint(payload),
        pdf_bytes=pdf_bytes,
        filename="sys.pdf",
        docx_bytes=docx,
        docx_filename="sys.docx",
    )
    assert MemorandosStore(store).file(identifier)[1] == pdf_bytes
    assert MemorandosStore(store).file(identifier, MIME_DOCX)[1] == docx


def test_open_memorandos_from_home_sets_overview():
    import inspect
    from portal import open_memorandos
    from services.memorandos_ui import NAV_KEY, NAV_NEW, NAV_OVERVIEW, _reset_new_form

    source = inspect.getsource(open_memorandos) + inspect.getsource(_reset_new_form)
    assert "Visão Geral" in inspect.getsource(open_memorandos)
    assert "NAV_KEY" in inspect.getsource(_reset_new_form)
    assert "memorandos_nav" in inspect.getsource(open_memorandos)
    assert NAV_NEW not in inspect.getsource(open_memorandos)


def test_memorandos_navigation_stays_in_sync(store, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from tests.access_testing import enable_login
    from database.store import ROOT
    from services.memorandos_ui import NAV_KEY, NAV_HISTORY, NAV_NEW, NAV_OVERVIEW

    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    enable_login(monkeypatch, store)
    monkeypatch.setattr("database.store.Store", lambda: store)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    app.button(key="open_memorandos").click().run()
    nav = app.radio(key=NAV_KEY)
    assert nav.value == NAV_OVERVIEW
    assert app.session_state[NAV_KEY] == NAV_OVERVIEW
    assert "Em andamento" not in nav.options
    assert NAV_HISTORY in nav.options
    assert not any(x.label == "Servidor afastado" for x in app.selectbox)
    assert any(m.label == "EM ANDAMENTO" for m in app.metric)
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    assert app.radio(key=NAV_KEY).value == NAV_NEW
    assert any(x.label == "Servidor afastado" for x in app.selectbox)
    assert not any(m.label == "EM ANDAMENTO" for m in app.metric)
    app.run()
    assert app.radio(key=NAV_KEY).value == NAV_NEW
    assert any(x.label == "Servidor afastado" for x in app.selectbox)
    app.radio(key=NAV_KEY).set_value(NAV_OVERVIEW).run()
    assert app.radio(key=NAV_KEY).value == NAV_OVERVIEW
    assert app.session_state[NAV_KEY] == NAV_OVERVIEW
    assert not any(x.label == "Servidor afastado" for x in app.selectbox)
    assert any(m.label == "EM ANDAMENTO" for m in app.metric)
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    app.sidebar.radio(key="portal_module").set_value("Início").run()
    app.button(key="open_memorandos").click().run()
    assert app.radio(key=NAV_KEY).value == NAV_OVERVIEW
    assert not any(x.label == "Servidor afastado" for x in app.selectbox)


def _login_memorandos_app(store, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from tests.access_testing import enable_login
    from database.store import ROOT

    enable_login(monkeypatch, store)
    monkeypatch.setattr("database.store.Store", lambda: store)
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()


def test_xlsx_import_does_not_select_first_server(store, monkeypatch):
    from services.memorandos_ui import NAV_KEY, NAV_NEW, NAV_OVERVIEW

    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    assert not any(x.label == "Servidor afastado" for x in app.selectbox)
    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    app.radio(key=NAV_KEY).set_value(NAV_OVERVIEW).run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    away = next(x for x in app.selectbox if x.label == "Servidor afastado")
    replacement = next(x for x in app.selectbox if x.label == "Substituto(a)")
    first_id = MemorandosStore(store).servers()[0]["id"]
    assert away.value is None
    assert replacement.value is None
    assert away.value != first_id


def test_rerun_preserves_manual_participant_choices(store, monkeypatch):
    from services.memorandos_ui import NAV_KEY, NAV_NEW

    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    away = next(x for x in app.selectbox if x.label == "Servidor afastado")
    chosen = MemorandosStore(store).servers()[0]["id"]
    away.set_value(chosen).run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value == chosen
    assert next(x for x in app.selectbox if x.label == "Substituto(a)").value is None
    app.run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value == chosen
    assert next(x for x in app.selectbox if x.label == "Substituto(a)").value is None


def test_entering_new_memorandum_clears_previous_choices(store, monkeypatch):
    from services.memorandos_ui import NAV_KEY, NAV_NEW, NAV_OVERVIEW

    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    away = next(x for x in app.selectbox if x.label == "Servidor afastado")
    away.set_value(MemorandosStore(store).servers()[0]["id"]).run()
    app.radio(key=NAV_KEY).set_value(NAV_OVERVIEW).run()
    app.radio(key=NAV_KEY).set_value(NAV_NEW).run()
    assert next(x for x in app.selectbox if x.label == "Servidor afastado").value is None
    assert next(x for x in app.selectbox if x.label == "Substituto(a)").value is None


def test_active_preview_exists_without_uploaded_docx():
    from services.memorandos import SOURCE_SYSTEM, SOURCE_UPLOAD, preview_is_valid

    payload = record()
    system = {
        "source": SOURCE_SYSTEM,
        "fingerprint": fingerprint(payload),
        "pdf": b"%PDF- ok",
        "docx": b"PK-ok",
    }
    assert preview_is_valid(system, payload)
    empty_upload = {"source": SOURCE_UPLOAD, "fingerprint": "x", "pdf": b"%PDF- x", "docx": None}
    assert not preview_is_valid(empty_upload, payload)


def test_admin_sees_server_base_tab_common_user_does_not(store, monkeypatch):
    from services.memorandos_ui import NAV_BASE, NAV_KEY, NAV_HISTORY, NAV_NEW, NAV_OVERVIEW, _nav_pages
    from services.access import Principal
    from tests.access_testing import seed_access

    admin_pages = _nav_pages(
        Principal(
            id=1, nome="Admin", email="a@t", perfil="ADMINISTRADOR", ativo=True,
            pode_portarias=True, pode_agenda=True, pode_oficios=True, pode_admin=True,
            gabinetes=(), pode_memorandos=True,
        )
    )
    user_pages = _nav_pages(
        Principal(
            id=2, nome="User", email="u@t", perfil="USUARIO", ativo=True,
            pode_portarias=False, pode_agenda=False, pode_oficios=False, pode_admin=False,
            gabinetes=(), pode_memorandos=True,
        )
    )
    assert admin_pages == [NAV_OVERVIEW, NAV_NEW, NAV_HISTORY, NAV_BASE]
    assert user_pages == [NAV_OVERVIEW, NAV_NEW, NAV_HISTORY]
    assert "Em andamento" not in admin_pages
    assert "Em andamento" not in user_pages
    assert NAV_BASE not in user_pages

    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    assert not app.exception
    nav = app.radio(key=NAV_KEY)
    assert NAV_BASE in nav.options
    nav.set_value(NAV_BASE).run()
    assert not app.exception
    assert app.radio(key=NAV_KEY).value == NAV_BASE
    uploaders = app.get("file_uploader")
    assert any(getattr(x, "label", None) == "Nova planilha XLSX" for x in uploaders)

    seed_access(
        store,
        email="memo.user@test.local",
        nome="Usuário Memorandos",
        perfil="USUARIO",
        pode_admin=False,
        pode_portarias=False,
        pode_agenda=False,
        pode_oficios=False,
        pode_memorandos=True,
    )
    monkeypatch.setattr(
        "services.access.oidc_identity",
        lambda: {"email": "memo.user@test.local", "name": "Usuário Memorandos", "email_verified": True},
    )
    from streamlit.testing.v1 import AppTest
    from database.store import ROOT

    common = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    common.button(key="open_memorandos").click().run()
    assert not common.exception
    common_nav = common.radio(key=NAV_KEY)
    assert common_nav.value == NAV_OVERVIEW
    assert NAV_BASE not in common_nav.options
    assert NAV_NEW in common_nav.options
    assert not any(x.label == "Ler e validar planilha" for x in common.button)
    assert not any(getattr(x, "label", None) == "Nova planilha XLSX" for x in common.get("file_uploader"))
    common.sidebar.radio(key="portal_module").set_value("Início").run()
    common.session_state[NAV_KEY] = NAV_BASE
    common.sidebar.radio(key="portal_module").set_value("Memorandos").run()
    assert not common.exception
    assert common.radio(key=NAV_KEY).value == NAV_OVERVIEW
    assert NAV_BASE not in common.radio(key=NAV_KEY).options
    assert any(m.label == "EM ANDAMENTO" for m in common.metric)
    assert not any(x.label == "Ler e validar planilha" for x in common.button)


def test_base_ui_does_not_hardcode_administrative_backend_flag():
    import inspect
    from services import memorandos_ui as ui

    source = inspect.getsource(ui._base) + inspect.getsource(ui.render)
    assert "administrator=True" not in source
    assert "administrator=principal.administrator" in inspect.getsource(ui._base)
    assert "if not principal.administrator" in inspect.getsource(ui._base)


def test_server_correction_opens_empty_and_closes_after_save(store, monkeypatch):
    from services.memorandos_ui import NAV_KEY, NAV_BASE

    MemorandosStore(store).import_servers(
        [
            {"nome": "Ana", "matricula": "1", "cargo": "A", "setor": "X"},
            {"nome": "Beto", "matricula": "2", "cargo": "B", "setor": "Y"},
        ],
        actor_email="admin@test.local",
        filename="x.xlsx",
        content_hash="a",
        administrator=True,
    )
    app = _login_memorandos_app(store, monkeypatch)
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_BASE).run()
    assert any(x.label == "Corrigir servidor" for x in app.button)
    assert not any(x.label == "Servidor" for x in app.selectbox)
    assert not any(x.label == "Nome" for x in app.text_input)
    assert not any(x.label == "Salvar correção" for x in app.button)
    next(x for x in app.button if x.label == "Corrigir servidor").click().run()
    picker = next(x for x in app.selectbox if x.label == "Servidor")
    assert picker.value is None
    first_id = MemorandosStore(store).all_servers(include_inactive=True)[0]["id"]
    assert picker.value != first_id
    assert not any(x.label == "Nome" for x in app.text_input)
    picker.set_value(first_id).run()
    assert any(x.label == "Nome" for x in app.text_input)
    assert any(x.label == "Cargo" for x in app.text_input)
    assert any(x.label == "Setor" for x in app.text_input)
    assert any(x.label == "Gênero" for x in app.selectbox)
    assert any(x.label == "Ativo" for x in app.checkbox)
    next(x for x in app.text_input if x.label == "Nome").set_value("Ana Atualizada").run()
    next(x for x in app.button if x.label == "Salvar correção").click().run()
    assert not app.exception
    saved = MemorandosStore(store).all_servers(include_inactive=True)
    assert any(row["nome"] == "Ana Atualizada" for row in saved)
    assert any(x.label == "Corrigir servidor" for x in app.button)
    assert not any(x.label == "Servidor" for x in app.selectbox)
    assert not any(x.label == "Nome" for x in app.text_input)
    assert any("Cadastro atualizado com sucesso" in (x.value or "") for x in app.success)


def test_legacy_em_andamento_session_falls_back_to_history(monkeypatch):
    from services.access import Principal
    from services import memorandos_ui

    state = {memorandos_ui.NAV_KEY: memorandos_ui.NAV_ONGOING_LEGACY}
    monkeypatch.setattr(memorandos_ui.st, "session_state", state)
    captured = {}

    def fake_radio(label, pages, **kwargs):
        captured["pages"] = list(pages)
        captured["value"] = state.get(memorandos_ui.NAV_KEY)
        return state[memorandos_ui.NAV_KEY]

    monkeypatch.setattr(memorandos_ui.st, "radio", fake_radio)
    principal = Principal(
        id=1,
        nome="Admin",
        email="a@t",
        perfil="ADMINISTRADOR",
        ativo=True,
        pode_portarias=True,
        pode_agenda=True,
        pode_oficios=True,
        pode_admin=True,
        gabinetes=(),
        pode_memorandos=True,
    )
    pages = memorandos_ui._nav_pages(principal)
    page = memorandos_ui._current_page(pages)
    assert page == memorandos_ui.NAV_HISTORY
    assert state[memorandos_ui.NAV_KEY] == memorandos_ui.NAV_HISTORY
    assert memorandos_ui.NAV_ONGOING_LEGACY not in captured["pages"]
    assert memorandos_ui.NAV_HISTORY in captured["pages"]


def test_history_filters_situacao_in_sql(store):
    service = MemorandosStore(store)
    today = date.today()
    ongoing = service.save_draft(
        record(data_inicio=today.isoformat(), data_fim=(today + timedelta(days=2)).isoformat()),
        actor_email="author@test",
    )
    scheduled = service.save_draft(
        record(data_inicio=(today + timedelta(days=5)).isoformat(), data_fim=(today + timedelta(days=8)).isoformat()),
        actor_email="author@test",
    )
    ended = service.save_draft(
        record(data_inicio=(today - timedelta(days=10)).isoformat(), data_fim=(today - timedelta(days=2)).isoformat()),
        actor_email="author@test",
    )
    rows = service.list(situacao="EM ANDAMENTO", limit=50)
    assert {r["id"] for r in rows} >= {ongoing}
    assert scheduled not in {r["id"] for r in rows}
    rows = service.list(situacao="AGENDADA", limit=50)
    assert scheduled in {r["id"] for r in rows}
    rows = service.list(situacao="ENCERRADA", limit=50)
    assert ended in {r["id"] for r in rows}


def _finalize_memo(service, payload=None, *, official=None, extra_docx=False):
    payload = payload or record()
    identifier = service.save_draft(payload, actor_email="author@test")
    kwargs = dict(
        preview_hash=fingerprint(payload),
        pdf_bytes=b"%PDF- synthetic",
        filename="memo.pdf",
    )
    if extra_docx:
        kwargs["docx_bytes"] = b"PK-docx"
        kwargs["docx_filename"] = "memo.docx"
    service.finalize(identifier, **kwargs)
    if official:
        service.set_official_number(identifier, official, "author@test")
    return identifier


def test_hard_delete_requires_admin_and_confirmation(store):
    service = MemorandosStore(store)
    identifier = _finalize_memo(service)
    other = _finalize_memo(service, record(etapas=[{
        "substituido": person("Carla", "Feminino", "8"),
        "substituto": person("Dora", "Feminino", "9"),
    }]))
    with pytest.raises(ValueError, match="não autorizado"):
        service.delete_finalized(identifier, administrator=False, confirmation="EXCLUIR")
    assert service.get(identifier)["id"] == identifier
    with pytest.raises(ValueError, match="EXCLUIR"):
        service.delete_finalized(identifier, administrator=True, confirmation="excluir")
    with pytest.raises(ValueError, match="EXCLUIR"):
        service.delete_finalized(identifier, administrator=True, confirmation="")
    assert service.get(identifier)["id"] == identifier
    snapshot = service.delete_finalized(identifier, administrator=True, confirmation="EXCLUIR")
    assert snapshot["id"] == identifier
    with pytest.raises(ValueError, match="não encontrado"):
        service.get(identifier)
    assert service.get(other)["id"] == other
    with store.connection(read_only=True) as c:
        assert c.execute("SELECT 1 FROM memorandos WHERE id=?", (identifier,)).fetchone() is None
        assert c.execute("SELECT 1 FROM memorandos_substituicao WHERE memorando_id=?", (identifier,)).fetchone() is None
        assert c.execute("SELECT 1 FROM memorandos_substituicao_etapas WHERE memorando_id=?", (identifier,)).fetchone() is None
        assert c.execute("SELECT 1 FROM memorandos_arquivos WHERE memorando_id=?", (identifier,)).fetchone() is None
        assert c.execute("SELECT 1 FROM memorandos WHERE id=?", (other,)).fetchone()
        assert c.execute("SELECT COUNT(*) FROM servidores").fetchone()[0] >= 0


def test_hard_delete_with_official_number_and_files(store):
    from services.memorandos import MIME_DOCX

    service = MemorandosStore(store)
    identifier = _finalize_memo(service, official="99/2026", extra_docx=True)
    assert service.file(identifier)[1].startswith(b"%PDF-")
    assert service.file(identifier, MIME_DOCX)[1].startswith(b"PK")
    service.delete_finalized(identifier, administrator=True, confirmation="EXCLUIR")
    with pytest.raises(ValueError, match="não encontrado"):
        service.get(identifier)
    assert service.file(identifier) is None
    assert service.file(identifier, MIME_DOCX) is None


def test_hard_delete_rolls_back_on_failure(store):
    from contextlib import contextmanager

    service = MemorandosStore(store)
    identifier = _finalize_memo(service, extra_docx=True)
    original = store.connection

    @contextmanager
    def wrapped(*, read_only=False, isolation=None, statement_timeout=None):
        with original(
            read_only=read_only, isolation=isolation, statement_timeout=statement_timeout
        ) as inner:
            class Proxy:
                def execute(self, sql, parameters=()):
                    text = " ".join(str(sql).split()).upper()
                    if text.startswith("DELETE FROM MEMORANDOS WHERE") and "ARQUIVOS" not in text:
                        raise RuntimeError("falha forçada")
                    return inner.execute(sql, parameters)

                def __getattr__(self, name):
                    return getattr(inner, name)

            yield Proxy()

    store.connection = wrapped
    with pytest.raises(RuntimeError, match="falha forçada"):
        service.delete_finalized(identifier, administrator=True, confirmation="EXCLUIR")
    store.connection = original
    assert service.get(identifier)["status"] == STATUS_FINALIZADO
    assert service.file(identifier)[1].startswith(b"%PDF-")


def test_hard_delete_audit_and_alerts(store):
    from database.audit import AuditStore
    from services.access import resolve_principal
    from services.alerts import collect_alerts, get_alert_summary
    from services.audit import registrar_evento
    from services.pending import collect_pending
    from tests.access_testing import TEST_IDENTITY

    service = MemorandosStore(store)
    today = date.today()
    payload = record(data_inicio=today.isoformat(), data_fim=(today + timedelta(days=3)).isoformat())
    identifier = _finalize_memo(service, payload)
    principal = resolve_principal(store, TEST_IDENTITY)
    items, _, _ = collect_pending(store, principal, today=today)
    assert any(i.source_id == identifier for i in items)
    snapshot = service.delete_finalized(identifier, administrator=True, confirmation="EXCLUIR")
    registrar_evento(
        store,
        evento="MEMORANDO_EXCLUIDO_DEFINITIVAMENTE",
        modulo="memorandos",
        acao="EXCLUIR",
        principal=principal,
        entidade_tipo="memorando",
        entidade_id=identifier,
        detalhes={
            "gabinete": snapshot["gabinete"][:80],
            "status": snapshot["status"],
            "numero_oficial": snapshot["numero_oficial"] or None,
        },
    )
    items, _, _ = collect_pending(store, principal, today=today)
    assert all(i.source_id != identifier for i in items)
    alerts, _, _ = collect_alerts(store, principal)
    assert all(i.source_id != identifier for i in alerts)
    summary = get_alert_summary(store, principal)
    assert all(i.source_id != identifier for i in summary["top"])
    events = AuditStore(store).list_events({"evento": "MEMORANDO_EXCLUIDO_DEFINITIVAMENTE"})
    assert events
    assert events[0]["usuario_email"] == TEST_IDENTITY["email"]
    assert events[0]["entidade_id"] == identifier
    with store.connection(read_only=True) as c:
        assert c.execute("SELECT COUNT(*) FROM servidores").fetchone()[0] >= 0


def test_hard_delete_postgres_contract(pg_store):
    service = MemorandosStore(pg_store)
    identifier = _finalize_memo(service, extra_docx=True)
    other = _finalize_memo(
        service,
        record(etapas=[{
            "substituido": person("Eva", "Feminino", "11"),
            "substituto": person("Fábio", "Masculino", "12"),
        }]),
    )
    service.delete_finalized(identifier, administrator=True, confirmation="EXCLUIR")
    with pytest.raises(ValueError, match="não encontrado"):
        service.get(identifier)
    assert service.get(other)["id"] == other


def test_hard_delete_ui_is_admin_only():
    import inspect
    from services import memorandos_ui

    details = inspect.getsource(memorandos_ui._details)
    controls = inspect.getsource(memorandos_ui._hard_delete_controls)
    listing = inspect.getsource(memorandos_ui._listing)
    assert "principal.administrator" in details
    assert "EXCLUIR" in controls
    assert "Confirmar exclusão definitiva" in controls
    assert "situacao" in listing
    assert "Em andamento" not in inspect.getsource(memorandos_ui._nav_pages)
    assert "NAV_HISTORY" in inspect.getsource(memorandos_ui.render)
    assert "list(limit=200)" not in inspect.getsource(memorandos_ui.render)
    assert "situacao_counts" in inspect.getsource(memorandos_ui.render)


def test_admin_sees_hard_delete_common_user_does_not(store, monkeypatch):
    from tests.access_testing import enable_login, seed_access
    from streamlit.testing.v1 import AppTest
    from database.store import ROOT
    from services.memorandos_ui import NAV_KEY, NAV_HISTORY

    identifier = _finalize_memo(MemorandosStore(store))
    enable_login(monkeypatch, store)
    monkeypatch.setattr("database.store.Store", lambda: store)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    app.button(key="open_memorandos").click().run()
    app.radio(key=NAV_KEY).set_value(NAV_HISTORY).run()
    assert not app.exception
    assert any(b.label == "Excluir definitivamente" for b in app.button)
    assert identifier

    seed_access(
        store,
        email="comum@test.local",
        perfil="USUARIO",
        pode_portarias=False,
        pode_agenda=False,
        pode_oficios=False,
        pode_admin=False,
        pode_memorandos=True,
    )
    monkeypatch.setattr(
        "services.access.oidc_identity",
        lambda: {"email": "comum@test.local", "name": "Comum", "email_verified": True},
    )
    common = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    common.button(key="open_memorandos").click().run()
    common.radio(key=NAV_KEY).set_value(NAV_HISTORY).run()
    assert not common.exception
    assert not any(b.label == "Excluir definitivamente" for b in common.button)

