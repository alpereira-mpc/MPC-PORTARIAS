"""Modelo visual por especialidade: formatação na minuta, sem conteúdo e sem RAG."""

import base64
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from services.minutas_docx import ELVIRA_TEMPLATE, MinutasDocxService, read_visual_model
from services.minutas_library import MinutasValidationError
from services.minutas_modelos import MinutasModeloVisualService, active_model_bytes
from services.minutas_specialties import CATALOG, RPPS
from tests.test_minutas_docx import Library as DocxLibrary, _payload, _summary, _target
from tests.test_minutas_specialties import Storage, build_store, library


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
POISON = (
    "Processo TC 02209/26. Instituto de Previdência de Bananeiras. "
    "Opino pela irregularidade das contas. João Pessoa."
)


def _install(store, *, models=True):
    sql = (Path(__file__).resolve().parents[1] / "database/minutas_especialidades.sql").read_text(encoding="utf-8")
    if models:
        sql += "\n" + (Path(__file__).resolve().parents[1] / "database/minutas_modelos_visuais.sql").read_text(encoding="utf-8")
    sql = re.sub(r"--[^\n]*", "", sql).replace("mpc_portarias.", "").replace("::text", "")
    with store.connection() as connection:
        for statement in sql.split(";"):
            statement = statement.strip()
            if statement.startswith(("CREATE TABLE", "CREATE INDEX", "CREATE UNIQUE INDEX", "CREATE OR REPLACE VIEW")):
                connection.execute(statement.replace("CREATE OR REPLACE VIEW", "CREATE VIEW"))
        for profile in CATALOG:
            connection.execute(
                "INSERT INTO minutas_especialidades(id,procuradora_chave,grupo,nome,perfil,estado,motor) "
                "VALUES(?,?,?,?,?,?,?)",
                tuple(profile[key] for key in ("id", "procuradora_chave", "grupo", "nome", "perfil", "estado", "motor")),
            )


def _count(store, table):
    with store.connection() as connection:
        return connection.execute(f"SELECT COUNT(*) AS n FROM {table}").fetchone()["n"]


def _model(path_name, *, top=5, logo=2.1, poison=POISON):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.4)
    run = section.header.paragraphs[0].add_run()
    run.add_picture(BytesIO(PNG), width=Inches(logo))
    body = document.add_paragraph(poison)
    body.paragraph_format.first_line_indent = Cm(2)
    body.paragraph_format.space_after = Pt(10)
    body.runs[0].font.size = Pt(13)
    ementa = document.add_paragraph("Ementa secreta do modelo sobre o jurisdicionado.")
    ementa.paragraph_format.left_indent = Cm(8)
    ementa.paragraph_format.space_after = Pt(18)
    ementa.runs[0].font.size = Pt(10)
    heading = document.add_paragraph("Irregularidade inventada no modelo")
    heading.runs[0].bold = True
    heading.runs[0].italic = True
    heading.runs[0].font.size = Pt(12)
    name = document.add_paragraph("FULANO DE TAL DO MODELO")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(30)
    name.runs[0].bold = True
    name.runs[0].italic = True
    name.runs[0].font.size = Pt(16)
    role = document.add_paragraph("Cargo inventado no modelo")
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role.runs[0].font.size = Pt(10)
    output = BytesIO()
    document.save(output)
    return path_name, output.getvalue()


@pytest.fixture
def prepared(tmp_path):
    store = build_store(tmp_path / "minutas.db")
    _install(store)
    storage = Storage()
    return store, storage, library(store, storage)


def test_modelo_visual_nao_cria_processo_nem_extracao(prepared):
    store, storage, scoped = prepared
    before = (_count(store, "minutas_referencias"), _count(store, "minutas_extracoes"),
              _count(store, "minutas_acervo_documentos"))
    saved = MinutasModeloVisualService(scoped).save(RPPS, _model("modelo visual.docx"))
    assert saved["estado"] == "CADASTRADO"
    assert saved["path"].startswith("modelos/")
    assert _count(store, "minutas_referencias") == before[0]
    assert _count(store, "minutas_extracoes") == before[1]
    assert _count(store, "minutas_acervo_documentos") == before[2]
    with store.connection() as connection:
        texts = [row["texto_integral"] for row in connection.execute(
            "SELECT texto_integral FROM minutas_extracoes"
        ).fetchall()]
    assert POISON not in "\n".join(texts)


def test_versionamento_preserva_anterior_e_um_ativo(prepared):
    store, storage, scoped = prepared
    service = MinutasModeloVisualService(scoped)
    first = service.save(RPPS, _model("primeiro.docx", top=5))
    second = service.save(RPPS, _model("segundo.docx", top=4.2, logo=1.8))
    assert first["path"] != second["path"]
    assert first["path"] in storage.objects and second["path"] in storage.objects
    versions = service.versions(RPPS)
    assert [item["versao"] for item in versions] == [1, 2]
    assert [int(item["ativo"]) for item in versions] == [0, 1]
    assert service.active(RPPS)["nome"] == "segundo.docx"
    again = service.save(RPPS, _model("segundo.docx", top=4.2, logo=1.8))
    assert again["estado"] == "ATUAL" and len(service.versions(RPPS)) == 2
    restored = service.save(RPPS, _model("primeiro.docx", top=5))
    assert restored["estado"] == "REATIVADO" and restored["versao"] == 1
    assert int(service.active(RPPS)["versao"]) == 1
    assert len(storage.objects) == 2


def test_selecao_por_especialidade_e_fallback(prepared):
    store, storage, scoped = prepared
    service = MinutasModeloVisualService(scoped)
    service.save("elvira.pca.camaras", _model("camara.docx", top=4))
    assert service.active(RPPS) is None
    assert active_model_bytes(scoped) is None
    service.save(RPPS, _model("rpps.docx", top=5))
    assert service.active("elvira.pca.camaras")["nome"] == "camara.docx"
    assert sha256(active_model_bytes(scoped)).digest() == sha256(
        storage.objects[service.active(RPPS)["path"]]
    ).digest()
    camaras = library(store, storage, "elvira.pca.camaras")
    assert sha256(active_model_bytes(camaras)).digest() != sha256(
        storage.objects[service.active(RPPS)["path"]]
    ).digest()


def test_exportacao_instrucao_camara_prioriza_modelo_da_especialidade(prepared):
    from services.minutas_specialties import CAMARAS, SHARED_INSTRUCTION_MOTOR
    from tests.test_minutas_instrucao import _structure

    store, storage, scoped = prepared
    with store.connection() as connection:
        connection.execute(
            "UPDATE minutas_especialidades SET motor=?, estado=? WHERE id=?",
            (SHARED_INSTRUCTION_MOTOR, "Em testes", CAMARAS),
        )
    camaras = library(store, storage, CAMARAS)
    analyzed = {"fontes_versao": "v", "processo_alvo_id": "instrucao-0210826"}
    generated = {
        "fontes_versao": "v",
        "processo_alvo_id": "instrucao-0210826",
        "estrutura": _structure(),
    }
    fallback = MinutasDocxService(camaras).export_instruction(
        analyzed, generated, processo="02108/26", exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    )
    fallback_doc = Document(BytesIO(fallback.content))
    assert fallback_doc.sections[0].left_margin.cm == pytest.approx(2.75, abs=0.05)
    MinutasModeloVisualService(scoped).save(RPPS, _model("rpps.docx", top=5))
    MinutasModeloVisualService(camaras).save(CAMARAS, _model("camara.docx", top=4))
    exported = MinutasDocxService(camaras).export_instruction(
        analyzed, generated, processo="02108/26", exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    )
    document = Document(BytesIO(exported.content))
    assert document.sections[0].left_margin.cm == pytest.approx(3, abs=0.05)
    assert document.styles["Minuta Ementa"].paragraph_format.left_indent.cm == pytest.approx(8, abs=0.05)
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Câmara Municipal de Exemplo" in body
    assert POISON not in body and "Bananeiras" not in body
    assert "RPPS" not in body


def test_sem_tabela_o_cadastro_nao_grava_e_a_exportacao_segue(tmp_path):
    store = build_store(tmp_path / "minutas.db")
    _install(store, models=False)
    storage = Storage()
    scoped = library(store, storage)
    assert MinutasModeloVisualService(scoped).enabled() is False
    with pytest.raises(MinutasValidationError, match="modelo institucional"):
        MinutasModeloVisualService(scoped).save(RPPS, _model("ausente.docx"))
    assert storage.objects == {}
    exported = _export_with(scoped)
    section = Document(BytesIO(exported.content)).sections[0]
    assert section.top_margin.cm == pytest.approx(6.35, abs=0.02)


def test_docx_aplica_formatacao_e_descarta_o_texto_do_modelo(prepared):
    store, storage, scoped = prepared
    content = _model("modelo - bananeiras.docx", top=5, logo=2.1)[1]
    spec = read_visual_model(content)
    logo = spec.pop("logo")
    dumped = json.dumps(spec, ensure_ascii=False).casefold()
    assert isinstance(logo, bytes) and logo.startswith(b"\x89PNG")
    assert "bananeir" not in dumped and "02209" not in dumped and "fulano" not in dumped
    assert spec["top_cm"] == pytest.approx(5, abs=0.05)
    MinutasModeloVisualService(scoped).save(RPPS, ("modelo - bananeiras.docx", content))
    exported = _export_with(scoped)
    document = Document(BytesIO(exported.content))
    section = document.sections[0]
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert section.top_margin.cm >= 5
    assert section.left_margin.cm == pytest.approx(3, abs=0.05)
    assert section.header_distance.cm == pytest.approx(0.8, abs=0.05)
    assert document.styles["Minuta Ementa"].paragraph_format.left_indent.cm == pytest.approx(8, abs=0.05)
    assert document.styles["Minuta Corpo"].font.size.pt == pytest.approx(13, abs=0.05)
    extent = document.sections[0].header.paragraphs[0]._p.xpath(".//wp:extent")[0]
    assert int(extent.get("cx")) / 914400 == pytest.approx(2.1, abs=0.05)
    assert "ELVIRA SAMARA PEREIRA DE OLIVEIRA" in body
    assert "FULANO" not in body and "Bananeiras" not in body and "02209/26" not in body
    assert "irregularidade das contas" not in body.casefold()
    with BytesIO(exported.content) as raw:
        xml = raw.getvalue().decode("latin1").casefold()
    assert "bananeir" not in xml and "fulano" not in xml and "02209" not in xml


def _export_with(scoped):
    target = _target()
    analyzed, generated = _payload(target)
    exporter_library = DocxLibrary(target)
    exporter_library.repository = scoped.repository
    exporter_library.storage = scoped.storage
    exporter_library.specialty_id = scoped.specialty_id
    return MinutasDocxService(exporter_library, source_summary=_summary).export(
        target["id"], analyzed, generated
    )
