from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.minutas_docx import (
    DOCX_MIME,
    ELVIRA_TEMPLATE,
    IDENTIFICATION_AFTER_PT,
    EMENTA_AFTER_PT,
    LOGO_AFTER_PT,
    LOGO_FROM_TOP_CM,
    MinutasDocxError,
    MinutasDocxService,
)
from services.minutas_generation import (
    DEVICE_INTRO_HARMONY,
    DEVICE_INTRO_NEUTRAL,
    MinutasGenerationService,
    RAG_RECOVERY_VERSION,
    REPORT_TRANSITION,
)
from services.minutas_library import MinutasAccessDenied


VERSION = "fontes-atuais"


def _target(
    identifier="target-taperoa",
    process="TC 01653/26",
    key="0165326",
    **changes,
):
    target = {
        "id": identifier,
        "especialidade_id": "elvira.pca.rpps",
        "classificacao_estado": "CLASSIFICADO",
        "procuradora": "Elvira Samara",
        "procuradora_chave": "elvira samara",
        "numero_processo": process,
        "processo_chave": key,
        "exercicio": 2026,
        "subtipo": "Instituto/Autarquia",
    }
    target.update(changes)
    return target


class Library:
    specialty_id = "elvira.pca.rpps"
    def __init__(self, target=None, *, authorized=True):
        self.target = target or _target()
        self.authorized = authorized

    def _authorize(self):
        if not self.authorized:
            raise MinutasAccessDenied("Acesso não autorizado a este módulo.")
        return "admin@test.local"

    def get(self, identifier):
        self._authorize()
        if str(identifier) != str(self.target["id"]):
            raise MinutasDocxError("Referência não encontrada.")
        return self.target


def _summary(_identifier):
    return {"fontes_versao": VERSION, "issues": []}


def _structure(*, long=False):
    body = (
        "A instrução demonstrou os fatos documentados e preservou a distinção "
        "entre defesa, Auditoria e avaliação ministerial."
    )
    if long:
        body = "\n\n".join(f"Parágrafo jurídico sintético {index}. {body}" for index in range(45))
    return {
        "identificacao_processual": "Processo TC 01653/26",
        "ementa": "PRESTAÇÃO DE CONTAS ANUAIS. ANÁLISE DOCUMENTAL.",
        "relatorio": "Relatório processual com acentuação e valor de R$ 1.234,56.",
        "fundamentacao_geral": body,
        "topicos_irregularidades": [
            {
                "titulo": "18.1 Receita intraorçamentária",
                "irregularidades_ids": ["18.1"],
                "avaliacoes_irregularidades": [
                    {
                        "irregularidade_id": "18.1",
                        "tratamento_ministerial_defesa": "MANTIDA_REJEICAO",
                        "fundamentos_divergencia": [],
                        "estado_saneamento": "ALEGADO_NAO_COMPROVADO",
                        "fontes_saneamento": [],
                    }
                ],
                "texto": (
                    "A fundamentação examina o art. 37 da Constituição e os "
                    "documentos efetivamente indicados.\n\n"
                    "“Trecho documental sintético para conferência.”"
                ),
            }
        ],
        "analise_sancionatoria": {
            "conclusao": "MULTA_NAO_CABIVEL",
            "irregularidades_ids": ["18.1"],
            "fundamentos_legais": [],
            "responsaveis": [],
            "dano_ao_erario": "NAO_COMPROVADO_OU_NAO_QUANTIFICADO",
            "texto": "Não há suporte suficiente para propor multa neste caso.",
        },
        "conclusao_pedidos": (
            "Ante o exposto, o Ministério Público de Contas, em atuação no "
            "âmbito do Tribunal de Contas, manifesta-se pelo julgamento das "
            "contas com ressalvas.\n"
            "1. Julgue-se o feito conforme a fundamentação."
        ),
        "pedidos": [
            "Julgue-se o feito conforme a fundamentação.",
            (
                "Expeçam-se recomendações:\n"
                "a) mantenha os controles previdenciários;\n"
                "b) documente as providências adotadas."
            ),
        ],
        "encerramento": "É o Parecer.",
    }


def _payload(target=None, *, sufficient=True, version=VERSION, warnings=()):
    target = target or _target()
    identity = {
        "processo_alvo_id": str(target["id"]),
        "processo_alvo_chave": "0165326",
        "procuradora_chave": target["procuradora_chave"],
    }
    analyzed = {
        **identity,
        "suficiente_para_geracao": sufficient,
        "fontes_versao": version,
        "rag_recovery_version": RAG_RECOVERY_VERSION,
    }
    generated = {
        **identity,
        "fontes_versao": version,
        "estrutura": _structure(),
        "advertencias_qualidade": list(warnings),
    }
    return analyzed, generated


def _export(*, target=None, structure=None):
    target = target or _target()
    analyzed, generated = _payload(target)
    if structure is not None:
        generated["estrutura"] = structure
    return MinutasDocxService(
        Library(target), source_summary=_summary
    ).export(target["id"], analyzed, generated)


def _image_blobs(document):
    return [
        relationship.target_part.blob
        for section in document.sections
        for relationship in section.header.part.rels.values()
        if relationship.reltype.endswith("/image")
    ]


def test_docx_is_editable_a4_and_matches_institutional_geometry():
    exported = _export()
    document = Document(BytesIO(exported.content))
    section = document.sections[0]

    assert exported.name == "MINUTA_PARECER_TC_01653_26.docx"
    assert exported.mime == DOCX_MIME
    assert exported.content[:2] == b"PK"
    assert section.page_width.cm == pytest.approx(21.001, abs=0.02)
    assert section.page_height.cm == pytest.approx(29.7, abs=0.02)
    assert section.top_margin.cm == pytest.approx(6.35, abs=0.02)
    assert section.bottom_margin.cm == pytest.approx(2.0, abs=0.02)
    assert section.left_margin.cm == pytest.approx(2.75, abs=0.02)
    assert section.right_margin.cm == pytest.approx(1.75, abs=0.02)
    assert document.styles["Normal"].font.name == "Arial"
    assert document.styles["Normal"].font.size.pt == pytest.approx(12)


def test_docx_reproduces_logo_header_and_spacing_without_empty_body_paragraphs():
    exported = _export()
    document = Document(BytesIO(exported.content))
    section = document.sections[0]
    header = section.header
    drawing_paragraph = next(
        paragraph for paragraph in header.paragraphs if paragraph._p.xpath(".//w:drawing")
    )
    extent = drawing_paragraph._p.xpath(".//wp:extent")[0]
    body_paragraphs = document.paragraphs
    ementa = next(
        paragraph for paragraph in body_paragraphs if paragraph.text.startswith("Ementa:")
    )
    responsible = next(
        paragraph
        for paragraph in body_paragraphs
        if paragraph.text.startswith("Responsável:")
    )

    assert section.different_first_page_header_footer is False
    assert section.header_distance.cm == pytest.approx(0.5, abs=0.02)
    assert len(_image_blobs(document)) == 1
    assert drawing_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert drawing_paragraph.paragraph_format.left_indent.cm == pytest.approx(0)
    assert drawing_paragraph.paragraph_format.right_indent.cm == pytest.approx(0)
    assert int(extent.get("cx")) / 914400 == pytest.approx(2.62, abs=0.01)
    assert int(extent.get("cy")) / 914400 == pytest.approx(1.84, abs=0.01)
    logo_from_top_cm = (
        section.header_distance.cm
        + drawing_paragraph.paragraph_format.space_before.pt * 2.54 / 72
    )
    assert drawing_paragraph.paragraph_format.space_before.pt == pytest.approx(
        (LOGO_FROM_TOP_CM - 0.5) * 72 / 2.54, abs=0.4
    )
    assert drawing_paragraph.paragraph_format.space_after.pt == pytest.approx(
        LOGO_AFTER_PT
    )
    assert responsible.paragraph_format.space_after.pt == pytest.approx(
        IDENTIFICATION_AFTER_PT
    )
    assert ementa.style.paragraph_format.space_after.pt == pytest.approx(
        EMENTA_AFTER_PT
    )
    assert logo_from_top_cm == pytest.approx(LOGO_FROM_TOP_CM, abs=0.05)
    logo_bottom_cm = logo_from_top_cm + int(extent.get("cy")) / 914400 * 2.54
    assert section.top_margin.cm - logo_bottom_cm >= 0.8
    assert not any(not paragraph.text for paragraph in body_paragraphs)


def test_docx_uses_original_logo_and_contains_only_target_content():
    exported = _export()
    output = Document(BytesIO(exported.content))
    reference = Document(ELVIRA_TEMPLATE)
    output_images = _image_blobs(output)
    reference_images = _image_blobs(reference)
    body_text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    footer_text = "\n".join(
        paragraph.text
        for section in output.sections
        for paragraph in section.footer.paragraphs
    )

    assert len(output_images) == 1
    assert sha256(output_images[0]).digest() == sha256(reference_images[0]).digest()
    assert "Processo TC Nº 01653/26" in body_text
    assert "Parecer Nº" in body_text
    assert "MINUTA-TC-01653/26" not in body_text
    assert "Origem: Não informada" in body_text
    assert "Responsável: Não informado" in body_text
    assert "ELVIRA SAMARA PEREIRA DE OLIVEIRA" in body_text
    assert "Procuradora-Geral do Ministério Público de Contas" in body_text
    assert "Análise sancionatória" in body_text
    assert "Não há suporte suficiente para propor multa" in body_text
    assert "MINUTA EXPERIMENTAL — SUJEITA À REVISÃO HUMANA" in footer_text
    forbidden = ("bananeir", "02209/26", " ycn", "joão pessoa,")
    assert not any(item in body_text.casefold() for item in forbidden)
    with ZipFile(BytesIO(exported.content)) as package:
        assert not any(name.endswith("vbaProject.bin") for name in package.namelist())
        xml_text = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        ).casefold()
    assert "bananeir" not in xml_text
    assert ">ycn<" not in xml_text
    assert "w:fldsimple" not in xml_text
    assert "logomarca do ministério público de contas da paraíba" in xml_text


def test_docx_preserves_formatting_and_groups_recommendations_in_one_item():
    exported = _export()
    document = Document(BytesIO(exported.content))
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text]
    all_text = "\n".join(paragraph.text for paragraph in paragraphs)

    ementa = next(paragraph for paragraph in paragraphs if paragraph.text.startswith("Ementa:"))
    heading = next(paragraph for paragraph in paragraphs if paragraph.text.startswith("18.1"))
    transcription = next(
        paragraph for paragraph in paragraphs if paragraph.text.startswith("“Trecho")
    )
    request = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text == "Julgue-se o feito conforme a fundamentação."
    )
    recommendation = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text.startswith("Pela expedição de recomenda")
    )
    alineas = [
        paragraph
        for paragraph in paragraphs
        if paragraph.style.name == "Minuta Alínea"
    ]

    assert ementa.style.name == "Minuta Ementa"
    assert ementa.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert ementa.style.paragraph_format.left_indent.cm == pytest.approx(
        6.752, abs=0.02
    )
    assert heading.style.name == "Minuta Irregularidade"
    assert heading.runs[0].bold is True and heading.runs[0].italic is True
    assert transcription.style.name == "Minuta Transcrição"
    assert request._p.pPr.numPr is not None
    assert recommendation._p.pPr.numPr is not None
    assert all(run.bold is not True for run in request.runs)
    assert all(run.bold is not True for run in recommendation.runs)
    assert "mantenha os controles previdenciários" in "\n".join(
        paragraph.text for paragraph in [recommendation, *alineas]
    )
    assert "documente as providências adotadas" in "\n".join(
        paragraph.text for paragraph in [recommendation, *alineas]
    )
    assert len(alineas) == 2
    assert all(paragraph._p.pPr.numPr is not None for paragraph in alineas)
    assert all_text.count("Julgue-se o feito conforme a fundamentação.") == 1
    assert "R$ 1.234,56" in all_text
    assert "art. 37" in all_text
    with ZipFile(BytesIO(exported.content)) as package:
        numbering = package.read("word/numbering.xml").decode("utf-8")
    assert 'w:val="decimal"' in numbering
    assert 'w:val="lowerLetter"' in numbering
    assert 'w:val="%1)"' in numbering


def test_docx_keeps_already_consolidated_recommendation_idempotent():
    target = _target()
    normalized, _warnings = MinutasGenerationService._normalize_structure(
        _structure(), target
    )

    document = Document(BytesIO(_export(target=target, structure=normalized).content))
    recommendations = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Pela expedição de recomenda")
    ]

    assert len(recommendations) == 1
    assert recommendations[0].count("Pela expedição de recomenda") == 1
    assert "no sentido de Recomendar" not in recommendations[0]


def test_docx_prefers_structured_paragraphs_and_keeps_legacy_compatibility():
    structure = _structure()
    structure["relatorio"] = "Texto legado que não deve prevalecer."
    structure["relatorio_paragrafos"] = [
        "Apresentação do processo.",
        "Histórico da instrução.",
        "É o relatório. Passa-se a opinar.",  # variação da IA; o backend padroniza
        REPORT_TRANSITION,
    ]
    structure["fundamentacao_geral_paragrafos"] = [
        "Premissas jurídicas gerais.",
        "Limites da prova documental.",
    ]
    structure["topicos_irregularidades"][0]["paragrafos"] = [
        "A Auditoria manteve a irregularidade.",
        "A defesa apresentou sua alegação.",
        "A análise ministerial confrontou os elementos.",
    ]

    document = Document(BytesIO(_export(structure=structure).content))
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert "Texto legado que não deve prevalecer." not in texts
    for expected in (
        "Apresentação do processo.",
        "Histórico da instrução.",
        "A Auditoria manteve a irregularidade.",
        "A defesa apresentou sua alegação.",
        "A análise ministerial confrontou os elementos.",
    ):
        assert texts.count(expected) == 1
    assert texts.count(REPORT_TRANSITION) == 1
    assert "Passa-se a opinar" not in "\n".join(texts)
    transition = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == REPORT_TRANSITION
    )
    assert transition.style.name == "Minuta Transição"
    assert transition.runs
    assert all(run.bold is True for run in transition.runs if run.text.strip())

    legacy = Document(BytesIO(_export(structure=_structure()).content))
    assert any(
        paragraph.text.startswith("Relatório processual")
        for paragraph in legacy.paragraphs
    )
    assert [paragraph.text for paragraph in legacy.paragraphs].count(
        REPORT_TRANSITION
    ) == 1


def test_docx_includes_authorized_origin_and_responsible_when_available():
    target = _target(parecer_numero="999/2025")
    analyzed, generated = _payload(target)
    generated["metadados_cabecalho"] = {
        "origem": "Instituto de Previdência do Município de Taperoá",
        "responsavel": "André Batista de Queiroz",
    }

    document = Document(
        BytesIO(
            MinutasDocxService(Library(target), source_summary=_summary).export(
                target["id"], analyzed, generated
            ).content
        )
    )
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert texts[:5] == [
        "Parecer Nº",
        "Processo TC Nº 01653/26",
        "Origem: Instituto de Previdência do Município de Taperoá",
        "Natureza: Prestação de Contas Anuais - exercício 2026",
        "Responsável: André Batista de Queiroz",
    ]
    assert not any("MINUTA-TC-01653/26" in text for text in texts)
    assert not any("999/2025" in text for text in texts)


def test_docx_keeps_only_footer_warning_and_preserves_justified_negative_request():
    structure = _structure()
    structure["pedidos"] = [
        "Afaste-se a multa, diante da controvérsia expressamente fundamentada."
    ]
    structure["encerramento"] = (
        "É o Parecer. Documento experimental sem caráter oficial e sujeito à revisão."
    )

    document = Document(BytesIO(_export(structure=structure).content))
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    footer_text = "\n".join(
        paragraph.text
        for section in document.sections
        for paragraph in section.footer.paragraphs
    )

    assert "Afaste-se a multa" in body_text
    assert body_text.count("É o Parecer.") == 1
    assert "experimental" not in body_text.casefold()
    assert footer_text.count("MINUTA EXPERIMENTAL — SUJEITA À REVISÃO HUMANA") == 1


@pytest.mark.parametrize("long", (False, True))
def test_docx_supports_short_and_long_drafts(long):
    exported = _export(structure=_structure(long=long))
    document = Document(BytesIO(exported.content))

    assert document.paragraphs
    assert document.paragraphs[-1].text.startswith("Procuradora-Geral")
    if long:
        assert len(document.paragraphs) > 55
        assert len(document.sections) == 1
        assert document.sections[0].different_first_page_header_footer is False
        assert len(_image_blobs(document)) == 1
        headings = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.style.name == "Minuta Irregularidade"
        ]
        assert headings
        assert all(
            paragraph.style.paragraph_format.keep_with_next is True
            for paragraph in headings
        )


def test_export_revalidates_authorization_target_and_sources_but_accepts_warnings():
    target = _target()
    analyzed, generated = _payload(target)

    with pytest.raises(MinutasAccessDenied):
        MinutasDocxService(
            Library(target, authorized=False), source_summary=_summary
        ).export(target["id"], analyzed, generated)

    wrong = dict(generated, processo_alvo_id="outro-processo")
    with pytest.raises(MinutasDocxError, match="não pertence"):
        MinutasDocxService(Library(target), source_summary=_summary).export(
            target["id"], analyzed, wrong
        )

    stale = dict(generated, fontes_versao="fontes-antigas")
    with pytest.raises(MinutasDocxError, match="fontes foram alteradas"):
        MinutasDocxService(Library(target), source_summary=_summary).export(
            target["id"], analyzed, stale
        )

    blocked = dict(analyzed, suficiente_para_geracao=False)
    blocked["pendencias_conferencia"] = [
        {
            "irregularidade": "18.3",
            "categoria": "Fato documentado",
            "documento": "auditoria.pdf",
            "pagina": 4,
            "natureza": "trecho não localizado na página indicada",
            "bloqueante": True,
        }
    ]
    exported = MinutasDocxService(
        Library(target), source_summary=_summary
    ).export(target["id"], blocked, generated)
    body = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(exported.content)).paragraphs
    )
    assert "INFORMAÇÕES AUXILIARES PARA REVISÃO" not in body
    assert "18.3" not in body
    assert "auditoria.pdf, p. 4" not in body
    assert "não integram a fundamentação oficial" not in body


def test_docx_omits_auxiliary_review_section_without_warnings():
    exported = _export()
    body = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(exported.content)).paragraphs
    )

    assert "INFORMAÇÕES AUXILIARES PARA REVISÃO" not in body


def test_export_requires_the_visual_template(tmp_path):
    target = _target()
    analyzed, generated = _payload(target)

    with pytest.raises(MinutasDocxError, match="modelo visual institucional"):
        MinutasDocxService(
            Library(target),
            source_summary=_summary,
            template_path=tmp_path / "ausente.docx",
        ).export(target["id"], analyzed, generated)


def test_docx_device_intro_avoids_duplication_and_adapts_to_audit_position():
    harmony = _structure()
    harmony["conclusao_pedidos"] = (
        "Ante o exposto, em harmonia com o entendimento do Órgão Técnico, "
        "o Ministério Público de Contas opina pelo (a):\n"
        "Julgue-se o feito conforme a fundamentação."
    )
    document = Document(BytesIO(_export(structure=harmony).content))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert texts.count(DEVICE_INTRO_HARMONY) == 1
    assert DEVICE_INTRO_NEUTRAL not in texts
    assert not any(
        "manifesta-se pelo julgamento" in text.casefold() for text in texts
    )

    divergence = _structure()
    divergence["topicos_irregularidades"][0]["avaliacoes_irregularidades"][0][
        "tratamento_ministerial_defesa"
    ] = "PARCIALMENTE_ACOLHIDA"
    divergence["conclusao_pedidos"] = (
        "Ante o exposto, em harmonia com o entendimento do Órgão Técnico, "
        "o Ministério Público de Contas opina pelo (a):"
    )
    divergent = Document(BytesIO(_export(structure=divergence).content))
    divergent_texts = [
        paragraph.text for paragraph in divergent.paragraphs if paragraph.text
    ]
    assert divergent_texts.count(DEVICE_INTRO_NEUTRAL) == 1
    assert DEVICE_INTRO_HARMONY not in divergent_texts


def test_docx_closing_spacing_and_request_runs_are_not_fully_bold():
    structure = _structure()
    structure["conclusao_pedidos"] = DEVICE_INTRO_HARMONY
    structure["pedidos"] = [
        "Regularidade com ressalvas das contas do responsável.",
        (
            "Pela expedição de recomendações à gestão, no sentido de que:\n"
            "a) mantenha os controles previdenciários;\n"
            "b) documente as providências adotadas."
        ),
        "Aplicação de multa proporcional à falha formal persistente.",
    ]
    document = Document(BytesIO(_export(structure=structure).content))
    closing = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "É o Parecer."
    )
    requests = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Minuta Pedido"
    ]
    alineas = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Minuta Alínea"
    ]

    assert closing.style.name == "Minuta Encerramento"
    assert closing.style.paragraph_format.space_before.pt == pytest.approx(24)
    assert closing.runs
    assert all(run.bold is not True for run in closing.runs)
    assert closing.style.paragraph_format.first_line_indent.cm == pytest.approx(
        3, abs=0.02
    )
    assert len(requests) == 3
    assert len(alineas) == 2
    assert all(
        run.bold is not True
        for paragraph in requests
        for run in paragraph.runs
    )


def test_docx_supports_short_and_long_ementa_with_optional_header_fields():
    short = _structure()
    short["ementa"] = "PCA."
    long = _structure()
    long["ementa"] = " ".join(["PRESTAÇÃO DE CONTAS ANUAIS."] * 40)
    target = _target()
    analyzed, generated = _payload(target)
    generated["metadados_cabecalho"] = {"origem": "", "responsavel": ""}
    generated["estrutura"] = short
    short_doc = Document(
        BytesIO(
            MinutasDocxService(Library(target), source_summary=_summary)
            .export(target["id"], analyzed, generated)
            .content
        )
    )
    long_doc = Document(BytesIO(_export(structure=long).content))
    assert any(paragraph.text.startswith("Ementa: PCA.") for paragraph in short_doc.paragraphs)
    assert any(
        paragraph.text.startswith("Ementa:") and len(paragraph.text) > 200
        for paragraph in long_doc.paragraphs
    )
    assert any("Origem: Não informada" in paragraph.text for paragraph in short_doc.paragraphs)
