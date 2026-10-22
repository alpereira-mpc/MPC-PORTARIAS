"""Identidade e estrutura jurídica: fontes atuais, PDF real e DOCX real."""

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document

from services.minutas_consistencia import integrar_sancoes
from services.minutas_docx import MinutasDocxError, MinutasDocxService
from services.minutas_generation import (
    MinutasGenerationError, MinutasGenerationService, _resolved_header_metadata,
    _target_with_header_metadata,
)
from services.minutas_instrucao import montar_contexto, receber_pdf
from tests.test_minutas_docx import _export, _structure, _target
from tests.test_minutas_instrucao import _pdf


def fonte(text):
    return {"fonte_id": "auditoria-atual", "conteudo": [{"pagina": 1, "texto": text}]}


@pytest.mark.parametrize("role", [
    "Prefeito Municipal", "Presidente da Câmara Municipal", "Diretor-Presidente",
    "Secretária Municipal", "Secretário Estadual", "Ordenador de Despesas", "Gestor de RPPS",
])
def test_nome_apos_cargo_e_antes_do_cargo(role):
    for text in (f"Sr. Sebastião Pinto Dantas, {role}.", f"{role} Sebastião Pinto Dantas."):
        metadata = _resolved_header_metadata({"exercicio": 2025}, [fonte(text)])
        assert metadata["responsavel"] == "Sebastião Pinto Dantas"


def test_pdf_extraido_relatorio_normalizacao_e_docx_sem_metadado_manual():
    text = ("Processo TC 02659/25. Contas do exercício de 2025.\n"
            "Sr. Sebastião Pinto Dantas, Prefeito Municipal.\n"
            "O auditor Pedro da Silva examinou os documentos.")
    pdf = receber_pdf("auditoria.pdf", _pdf(text), "relatorio_inicial")
    context = montar_contexto([pdf], processo="02659/25", exercicio=2025,
                             origem="Prefeitura Municipal de Frei Martinho")["contexto"]
    target = _target(process="02659/25", key="0265925", exercicio=2025, responsavel="")
    metadata = _resolved_header_metadata(target, context["documentos_processo_alvo"])
    assert metadata["responsavel"] == "Sebastião Pinto Dantas"
    target = _target_with_header_metadata(target, metadata)
    structure = _structure()
    structure["relatorio"] = "Contas de 2025 do Sr. Sebastião Pinto Dantas, Prefeito Municipal."
    normalized, _ = MinutasGenerationService._normalize_structure(structure, target, context)
    assert not normalized["pendencias_identificacao"]
    assert normalized["identidade_responsaveis"]["evidencias"][0]["pagina"] == 1
    output = Document(BytesIO(_export(target=target, structure=normalized).content))
    rendered = "\n".join(p.text for p in output.paragraphs)
    assert "Responsável: Sebastião Pinto Dantas" in rendered
    assert "02659/25" in rendered


def test_multiplos_e_sucessivos_preservam_periodos_sem_escolha_por_ia():
    text = ("Responsável: Ana de Sousa (01/01/2025 a 30/06/2025)\n"
            "Responsável: João da Silva (01/07/2025 a 31/12/2025)")
    metadata = _resolved_header_metadata({"exercicio": 2025}, [fonte(text)])
    assert "Ana de Sousa (01/01/2025 a 30/06/2025)" in metadata["responsavel"]
    assert "João da Silva (01/07/2025 a 31/12/2025)" in metadata["responsavel"]
    target = _target_with_header_metadata({"exercicio": 2025}, metadata)
    assert target["responsavel"] == metadata["responsavel"]
    collective = _resolved_header_metadata({}, [fonte("Responsáveis: Ana de Sousa; João da Silva")])
    assert collective["responsavel"] == "Ana de Sousa; João da Silva"


def test_conflito_nao_e_desempatado_por_ia_ou_metadado():
    sources = [fonte("Responsável: Ana de Sousa"), fonte("Responsável: João da Silva")]
    for target in ({}, {"responsavel": "Ana de Sousa"}):
        metadata = _resolved_header_metadata(target, sources, {"resumo": "Responsável: Ana de Sousa"})
        assert "responsavel" not in metadata
        assert metadata["pendencias_identificacao"]
        assert not _target_with_header_metadata(target, metadata).get("responsavel")


def test_exercicio_gestor_atual_e_papeis_processuais_nao_substituem_responsavel():
    sources = [fonte("João da Silva, Prefeito Municipal no exercício de 2024.\n"
                     "Atual Prefeito Pedro de Sousa.\n"
                     "Procurador: Paulo da Costa.\n"
                     "Advogado: José da Silva.\n"
                     "Ana de Sousa, Prefeita Municipal no exercício de 2025.")]
    assert _resolved_header_metadata({"exercicio": 2025}, sources)["responsavel"] == "Ana de Sousa"


def test_sem_fonte_nao_aproveita_nome_do_corpo_ou_do_rag():
    structure = _structure()
    structure["relatorio"] = "José da Costa, Prefeito Municipal."
    target = _target(responsavel="Não informado")
    metadata = _resolved_header_metadata(target, [], {"resumo": structure["relatorio"]})
    assert not metadata.get("responsavel")
    with pytest.raises(MinutasDocxError, match="responsável"):
        _export(target=target, structure=structure)


def test_metadado_confirmado_sem_nome_no_corpo_e_divergencia_explicita():
    assert _export(target=_target(responsavel="Maria de Sousa")).content
    structure = _structure()
    structure["relatorio"] = "Ana da Costa, Prefeita Municipal."
    with pytest.raises(MinutasDocxError, match="corpo identifica"):
        _export(target=_target(responsavel="Maria de Sousa"), structure=structure)


@pytest.mark.parametrize("title", ["Análise sancionatória", "III. Dosimetria da sanção", "Aplicação de penalidades"])
def test_capitulo_generico_bloqueado_sem_apagar_fundamentacao(title):
    structure = _structure()
    structure["topicos_irregularidades"][0]["titulo"] = title
    original = deepcopy(structure)
    with pytest.raises(MinutasGenerationError, match="Capítulo sancionatório"):
        MinutasGenerationService._normalize_structure(structure, _target())
    assert structure == original


@pytest.mark.parametrize("conclusion", ["MULTA_CABIVEL", "MULTA_NAO_CABIVEL"])
def test_fundamentacao_e_pedido_legitimo_preservados_sem_capitulo(conclusion):
    structure = _structure()
    structure["analise_sancionatoria"].update(conclusao=conclusion, texto="A conduta documentada no item 18.1 fundamenta a consequência proposta.")
    if conclusion == "MULTA_CABIVEL":
        structure["pedidos"].append("Aplicar multa pela conduta fundamentada no item 18.1.")
    normalized = integrar_sancoes(structure)
    assert normalized == integrar_sancoes(normalized)
    assert structure["analise_sancionatoria"]["texto"] in normalized["topicos_irregularidades"][0]["texto"]
    assert normalized["pedidos"] == structure["pedidos"]
    doc = Document(BytesIO(_export(structure=normalized).content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Análise sancionatória" not in text
    assert structure["analise_sancionatoria"]["texto"] in text


def test_docx_confere_campo_em_tabela_e_rejeita_nome_so_no_corpo():
    doc = Document()
    doc.add_paragraph("Relatório: Maria de Sousa, Prefeita Municipal.")
    out = BytesIO(); doc.save(out)
    with pytest.raises(MinutasDocxError, match="não preservou"):
        MinutasDocxService._conferir_identificacao_docx(out.getvalue(), ["Responsável: Maria de Sousa"])
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Responsável:"
    row.cells[1].text = "Maria de Sousa"
    out = BytesIO(); doc.save(out)
    MinutasDocxService._conferir_identificacao_docx(out.getvalue(), ["Responsável: Maria de Sousa"])
