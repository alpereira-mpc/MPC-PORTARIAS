"""Perfis visuais de PCA de Prefeituras, sem depender dos DOCX de referência."""

import base64
import re
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Inches

from services.minutas_docx import (
    MinutasDocxError,
    MinutasDocxService,
    _apply_numbering,
    _build_docx,
    _numbering_definition,
    _signature_date,
)
from services.minutas_generation import _draft_header_lines
from services.minutas_modelo_docx import _eh_faixa_paragrafo, renderizar, resolver_para_exportacao
from services.minutas_pca_formato import PERFIS, _e_data, complementar_minuta_pca
from services.minutas_procuradores import (
    assinatura_bradson,
    assinatura_neutra,
    exigir_procurador,
    listar_ativos,
    visual_bradson,
    visual_neutro,
)
from services.minutas_specialties import CAMARAS, PREFEITURAS, RPPS


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
MODULO = Path(__file__).resolve().parents[1] / "services" / "minutas_pca_formato.py"
EMENTA_CURTA = "Contas regulares com ressalva formal."
EMENTA_LONGA = "A prestação de contas exige conferência documental do exercício. " * 30
CITACAO = "Art. 16 da Lei de Responsabilidade Fiscal condiciona a criação de despesa."
DESTAQUE = (
    "O gestor descumpriu o dever de licitar. "
    "“A contratação direta sem pesquisa de preços viola a regra aplicável.” "
    "A falha permanece."
)
TITULO = "Ausência de pesquisa de preços"
CORPO = "O apontamento segue sem documento de pesquisa."
CONCLUSAO = "Ante o exposto, opina-se pela irregularidade das contas."
TRANSICAO = "É o Relatório. Passo a opinar."


def _estrutura(ementa):
    return {
        "ementa": ementa,
        "relatorio": "Relato objetivo dos atos de instrução.\n\n" + TRANSICAO,
        "relatorio_paragrafos": ["Relato objetivo dos atos de instrução.", TRANSICAO],
        "fundamentacao_geral": f"A análise permanece no limite da instrução.\n\n{CITACAO}\n\n{DESTAQUE}",
        "fundamentacao_geral_paragrafos": [
            "A análise permanece no limite da instrução.",
            CITACAO,
            DESTAQUE,
        ],
        "topicos_irregularidades": [{
            "titulo": TITULO,
            "texto": CORPO,
            "paragrafos": [CORPO],
        }],
        "analise_sancionatoria": {"texto": "A multa depende de conferência humana."},
        "conclusao_pedidos": CONCLUSAO,
        "pedidos": [
            "Julgue-se irregular a matéria.\na) registre a ciência do gestor;\nb) acompanhe o cumprimento."
        ],
        "encerramento": "É o parecer.",
    }


def _visual(sigla):
    if sigla == "BLTC":
        visual = visual_bradson()
    elif sigla == "ESPO":
        visual = {}
    else:
        visual = visual_neutro()
    visual["logo"] = PNG
    return visual


def _perfil():
    return {
        "nome": "NOME DE TESTE",
        "cargo": "Cargo de teste",
        "localidade": "João Pessoa",
        "template": None,
    }


def _alvo():
    return {"numero_processo": "00000/00", "data_parecer": "2026-10-21"}


def _biblioteca(specialty, sigla):
    biblioteca = type("Biblioteca", (), {})()
    biblioteca.specialty_id = specialty
    biblioteca.procurador = exigir_procurador(sigla)
    biblioteca.repository = None
    biblioteca.storage = None
    return biblioteca


def _servico(specialty, sigla):
    servico = MinutasDocxService.__new__(MinutasDocxService)
    servico.library = _biblioteca(specialty, sigla)
    return servico


def _marcar(paragraph):
    referencia = OxmlElement("w:footnoteReference")
    referencia.set(qn("w:id"), "1")
    paragraph.runs[0]._r.append(referencia)
    paragraph.paragraph_format.page_break_before = True


def _base(sigla, ementa):
    profile = _perfil()
    content = _build_docx(_alvo(), _estrutura(ementa), profile, [
        "Processo TC Nº 00000/00",
        "Origem: Unidade de teste",
    ], _visual(sigla))
    document = Document(BytesIO(content))
    tabela = document.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "FAIXA"
    tabela.cell(0, 1).text = "CINZA"
    meio = next(item for item in document.paragraphs if item.text == CORPO)
    _marcar(meio)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _aplicar(sigla, ementa, *, origem_modelo=False):
    bruto = _base(sigla, ementa)
    data = "" if origem_modelo or PERFIS[sigla].get("preservar_geometria") else (
        "João Pessoa, 21 de outubro de 2026."
    )
    return bruto, complementar_minuta_pca(
        bruto,
        sigla=sigla,
        structure=_estrutura(ementa),
        origem_modelo=origem_modelo,
        data_assinatura=data,
    )


def _paragrafo(document, texto):
    return next(item for item in document.paragraphs if item.text == texto or item.text.startswith(texto))


def _fonte(run):
    propriedades = run._r.find(qn("w:rPr"))
    fontes = propriedades.find(qn("w:rFonts")) if propriedades is not None else None
    if fontes is None:
        return ""
    return fontes.get(qn("w:ascii")) or ""


def _recuo(paragraph, nome):
    propriedades = paragraph._p.find(qn("w:pPr"))
    recuo = propriedades.find(qn("w:ind")) if propriedades is not None else None
    if recuo is None or not recuo.get(qn(nome)):
        return None
    return int(recuo.get(qn(nome))) / 567


def _conferir_ementa(paragraph, spec):
    rotulo = spec.get("rotulo")
    if rotulo:
        assert paragraph.runs[0].text.lower().startswith("ementa")
        assert paragraph.runs[0].bold is rotulo["bold"]
        assert paragraph.runs[0].italic is rotulo["italic"]
        miolo = paragraph.runs[-1]
    else:
        miolo = paragraph.runs[0]
        assert miolo.italic is spec["italic"]
        assert miolo.bold is spec["bold"]
    assert miolo.font.size.pt == pytest.approx(spec["size"])
    if rotulo:
        assert miolo.bold is spec["bold"]
        assert miolo.italic is spec["italic"]


def _numerados(document):
    total = 0
    for paragraph in document.paragraphs:
        propriedades = paragraph._p.find(qn("w:pPr"))
        if propriedades is not None and propriedades.find(qn("w:numPr")) is not None:
            total += 1
    return total


def _quebras(document):
    return sum(
        1 for node in document.element.iter(qn("w:pageBreakBefore"))
    )


def _notas(document):
    return sum(1 for node in document.element.iter(qn("w:footnoteReference")))


@pytest.mark.parametrize("sigla", ["ESPO", "BLTC", "LAF", "MASN", "IBMF", "SBBQ", "MTTF"])
@pytest.mark.parametrize("ementa", [EMENTA_CURTA, EMENTA_LONGA])
def test_perfil_formata_blocos_e_preserva_a_pagina(sigla, ementa):
    bruto, gerado = _aplicar(sigla, ementa)
    antes, depois = Document(BytesIO(bruto)), Document(BytesIO(gerado))
    assert antes.sections[0]._sectPr.xml == depois.sections[0]._sectPr.xml
    assert _numerados(antes) == _numerados(depois) >= 2
    assert _quebras(antes) == _quebras(depois) == 1
    assert _notas(antes) == _notas(depois) == 1
    assert len(depois.tables) == 1
    assert depois.tables[0].cell(0, 0).text == "FAIXA"
    assert "MINUTA EXPERIMENTAL — SUJEITA À REVISÃO HUMANA" in depois.sections[0].footer.paragraphs[0].text
    assert depois.sections[0].header.paragraphs[0]._p.find(".//" + qn("w:drawing")) is not None

    ementa_item = _paragrafo(depois, "Ementa:")
    titulo = _paragrafo(depois, TITULO)
    citacao = _paragrafo(depois, CITACAO)
    corpo = _paragrafo(depois, CORPO)
    conclusao = _paragrafo(depois, CONCLUSAO)
    transicao = _paragrafo(depois, TRANSICAO)
    identificacao = _paragrafo(depois, "Processo TC")
    assert ementa_item.style.name == "Minuta Ementa"
    assert titulo.style.name == "Minuta Irregularidade"
    assert identificacao.style.name == "Minuta Identificação"
    assert identificacao.runs[0].bold is True
    assert _recuo(identificacao, "w:firstLine") is None
    assert transicao.runs[0].bold is True
    assert "NOME DE TESTE" in "\n".join(item.text for item in depois.paragraphs)
    assert "Cargo de teste" in "\n".join(item.text for item in depois.paragraphs)
    destaque = _paragrafo(depois, "O gestor descumpriu")
    assert any(run.italic and "pesquisa de preços" in run.text for run in destaque.runs)
    assert any(not run.italic and "dever de licitar" in run.text for run in destaque.runs)
    assert destaque.text == DESTAQUE

    perfil = PERFIS[sigla]
    if perfil.get("preservar_geometria"):
        assert _fonte(corpo.runs[0]) == "Arial"
        assert _recuo(ementa_item, "w:left") is None
        assert _recuo(corpo, "w:firstLine") is None
        assert titulo.runs[0].bold is True and titulo.runs[0].italic is True
        if sigla == "ESPO":
            assert citacao.runs[0].italic is True
            assert ementa_item.style.paragraph_format.left_indent.cm == pytest.approx(6.752, abs=0.05)
        else:
            assert citacao.runs[0].italic is not True
            assert ementa_item.runs[0].bold is not True
            assert ementa_item.style.paragraph_format.left_indent.cm == pytest.approx(4.75, abs=0.05)
            assert "PAGE" in depois.sections[0].footer.paragraphs[0]._p.xml
            assert depois.paragraphs[-1].text != "João Pessoa, 21 de outubro de 2026."
        assert not any(
            item.text.startswith("João Pessoa,")
            for item in depois.paragraphs
            if item.style.name != "Minuta Cargo"
        )
    else:
        corpo_spec = perfil["corpo"]
        ementa_spec = perfil["ementa"]
        citacao_spec = perfil["citacao"]
        assert _fonte(corpo.runs[0]) == corpo_spec.get("font", "Arial")
        assert _recuo(corpo, "w:firstLine") == pytest.approx(corpo_spec["first_cm"], abs=0.05)
        assert corpo.runs[0].font.size.pt == pytest.approx(corpo_spec["size"])
        assert corpo.runs[0].bold is corpo_spec["bold"]
        assert corpo.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert _recuo(ementa_item, "w:left") == pytest.approx(ementa_spec["left_cm"], abs=0.05)
        _conferir_ementa(ementa_item, ementa_spec)
        assert titulo.runs[0].bold is True
        assert titulo.runs[0].italic is perfil["titulo"]["italic"]
        assert bool(citacao.runs[0].italic) is bool(citacao_spec.get("italic"))
        assert citacao.runs[0].font.size.pt == pytest.approx(citacao_spec["size"])
        if "left_cm" in citacao_spec:
            assert _recuo(citacao, "w:left") == pytest.approx(citacao_spec["left_cm"], abs=0.05)
        assert bool(conclusao.runs[0].bold) is bool(perfil["conclusao"]["bold"])
        data = _paragrafo(depois, "João Pessoa,")
        assert data.alignment == {
            "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
        }[perfil["data"]["align"]]
        textos = [item.text for item in depois.paragraphs]
        assert textos.index(data.text) < textos.index("NOME DE TESTE")
        if "font" in perfil["data"]:
            assert _fonte(data.runs[0]) == perfil["data"]["font"]
    if sigla == "IBMF":
        assert corpo.runs[0].bold is not True
        assert _fonte(corpo.runs[0]) == "Calibri"
        assert corpo.runs[0].font.size.pt == pytest.approx(13)
        assert _recuo(ementa_item, "w:left") == pytest.approx(5, abs=0.05)
        assert ementa_item.runs[-1].font.size.pt == pytest.approx(11.5)
        assert ementa_item.runs[-1].bold is not True
        assert citacao.runs[0].italic is not True
    if sigla == "MTTF":
        assert ementa_item.runs[0].text.lower().startswith("ementa")
        assert ementa_item.runs[0].italic is True
        assert ementa_item.runs[-1].italic is not True
        assert len(ementa_item.runs[-1].text) > 20
        assert citacao.runs[0].italic is not True


def test_modelo_visual_permanece_base_e_o_perfil_so_ajusta_blocos():
    document = Document()
    secao = document.sections[0]
    secao.top_margin = Cm(2.2)
    secao.left_margin = Cm(3.2)
    secao.header.paragraphs[0].add_run().add_picture(BytesIO(PNG), width=Inches(1))
    secao.footer.paragraphs[0].add_run("Rodapé estável do modelo")
    tabela = document.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "FAIXA"
    tabela.cell(0, 1).text = "IDENTIFICACAO"
    document.add_paragraph("Processo TC Nº 00000/00")
    ementa = document.add_paragraph("Ementa: " + EMENTA_CURTA)
    corpos = [
        document.add_paragraph("Relato objetivo dos atos de instrução."),
        document.add_paragraph(TRANSICAO),
        document.add_paragraph(CITACAO),
        document.add_paragraph(DESTAQUE),
        document.add_paragraph(TITULO),
        document.add_paragraph(CORPO),
        document.add_paragraph("Análise sancionatória"),
        document.add_paragraph("A multa depende de conferência humana."),
        document.add_paragraph(CONCLUSAO),
    ]
    for paragraph in corpos:
        paragraph.paragraph_format.first_line_indent = Cm(1.1)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(12)
            run.font.name = "Calibri"
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    fecho = document.add_paragraph("É o Parecer.")
    fecho.alignment = WD_ALIGN_PARAGRAPH.CENTER
    data = document.add_paragraph("João Pessoa, 21 de outubro de 2026.")
    data.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    nome = document.add_paragraph("NOME DE TESTE")
    nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nome.runs[0].bold = True
    pedido = document.add_paragraph("Julgue-se irregular a matéria.")
    _apply_numbering(pedido, _numbering_definition(
        document, kind="decimal", marker="%1.", left_twips=720, hanging_twips=360,
    ))
    bruto = BytesIO()
    document.save(bruto)
    original = bruto.getvalue()
    gerado = complementar_minuta_pca(
        original, sigla="LAF", structure=_estrutura(EMENTA_CURTA),
        origem_modelo=True, data_assinatura="João Pessoa, 21 de outubro de 2026.",
    )
    depois = Document(BytesIO(gerado))
    assert depois.sections[0].top_margin.cm == pytest.approx(2.2, abs=0.05)
    assert depois.sections[0].left_margin.cm == pytest.approx(3.2, abs=0.05)
    assert depois.sections[0].footer.paragraphs[0].text == "Rodapé estável do modelo"
    assert depois.tables[0].cell(0, 0).text == "FAIXA"
    assert _recuo(_paragrafo(depois, "Relato objetivo"), "w:firstLine") == pytest.approx(2.54, abs=0.05)
    assert _fonte(_paragrafo(depois, "Relato objetivo").runs[0]) == "Calibri"
    assert _paragrafo(depois, "Relato objetivo").runs[0].font.size.pt == pytest.approx(13)
    assert _paragrafo(depois, "Relato objetivo").runs[0].bold is not True
    assert _recuo(_paragrafo(depois, "Ementa:"), "w:left") == pytest.approx(6.35, abs=0.05)
    assert _paragrafo(depois, TITULO).runs[0].bold is True
    assert _recuo(_paragrafo(depois, TITULO), "w:firstLine") in (None, 0) or _recuo(
        _paragrafo(depois, TITULO), "w:firstLine"
    ) == pytest.approx(0, abs=0.05)
    assert _paragrafo(depois, CITACAO).runs[0].italic is True
    assert _recuo(_paragrafo(depois, CITACAO), "w:left") == pytest.approx(2.54, abs=0.05)
    assert _paragrafo(depois, CITACAO).runs[0].font.size.pt == pytest.approx(11)
    assert _paragrafo(depois, "João Pessoa,").alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _paragrafo(depois, "É o Parecer.").alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    pedido = _paragrafo(depois, "Julgue-se irregular")
    assert pedido._p.find(qn("w:pPr")).find(qn("w:numPr")) is not None
    assert pedido.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _fonte(_paragrafo(depois, CORPO).runs[0]) == "Calibri"

    preservada = complementar_minuta_pca(
        original, sigla="ESPO", structure=_estrutura(EMENTA_CURTA), origem_modelo=True,
    )
    elvira = Document(BytesIO(preservada))
    assert _recuo(_paragrafo(elvira, "Ementa:"), "w:left") is None
    assert _paragrafo(elvira, CITACAO).runs[0].italic is True
    assert _recuo(_paragrafo(elvira, "Relato objetivo"), "w:firstLine") == pytest.approx(1.1, abs=0.05)
    assert elvira.sections[0].left_margin.cm == pytest.approx(3.2, abs=0.05)


def test_ementa_descartada_pelo_modelo_estrutural_volta_com_o_perfil():
    document = Document()
    document.add_paragraph("Processo TC Nº 00000/00")
    tabela = document.add_table(rows=1, cols=1)
    tabela.cell(0, 0).text = "PARECER"
    document.add_paragraph("Relato objetivo dos atos de instrução.")
    document.add_paragraph("NOME DE TESTE")
    buffer = BytesIO()
    document.save(buffer)
    gerado = Document(BytesIO(complementar_minuta_pca(
        buffer.getvalue(), sigla="LAF", structure=_estrutura(EMENTA_CURTA), origem_modelo=True,
    )))
    ementa = _paragrafo(gerado, "Ementa:")
    assert ementa.runs[0].italic is True
    assert _fonte(ementa.runs[0]) == "Calibri"
    assert ementa.runs[0].font.size.pt == pytest.approx(11)
    assert _recuo(ementa, "w:left") == pytest.approx(6.35, abs=0.05)
    filhos = [child for child in gerado.element.body if child.tag != qn("w:sectPr")]
    assert filhos.index(ementa._p) < filhos.index(gerado.tables[0]._tbl)
    assert gerado.tables[0].cell(0, 0).text == "PARECER"

    bradson = Document()
    faixa = bradson.add_table(rows=1, cols=2)
    faixa.cell(0, 0).text = "EMENTA:"
    faixa.cell(0, 1).text = "Texto já residente na tabela."
    bradson.add_paragraph("Relato objetivo dos atos de instrução.")
    bruto = BytesIO()
    bradson.save(bruto)
    preservada = Document(BytesIO(complementar_minuta_pca(
        bruto.getvalue(), sigla="BLTC", structure=_estrutura(EMENTA_CURTA), origem_modelo=True,
    )))
    assert sum(item.text.lower().count("ementa") for item in preservada.paragraphs) == 0
    assert "Texto já residente na tabela." in preservada.tables[0].cell(0, 1).text


def test_modelo_aplica_o_perfil_mesmo_com_recuo_e_negrito_do_prototipo():
    document = Document()
    ementa = document.add_paragraph("Ementa: " + EMENTA_CURTA)
    ementa.paragraph_format.left_indent = Cm(6.5)
    ementa.runs[0].italic = True
    ementa.runs[0].font.size = Pt(12)
    corpo = document.add_paragraph("Relato objetivo dos atos de instrução.")
    corpo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corpo.runs[0].bold = True
    corpo.runs[0].font.size = Pt(14)
    corpo.runs[0].font.name = "Times New Roman"
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), "CCCCCC")
    corpo._p.get_or_add_pPr().append(sombra)
    nome = document.add_paragraph("NOME DE TESTE")
    nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    faixa = OxmlElement("w:shd")
    faixa.set(qn("w:val"), "clear")
    faixa.set(qn("w:fill"), "CCCCCC")
    nome._p.get_or_add_pPr().append(faixa)
    tabela = document.add_table(rows=1, cols=1)
    celula = tabela.cell(0, 0).paragraphs[0]
    celula.text = "FAIXA"
    banda = OxmlElement("w:shd")
    banda.set(qn("w:val"), "clear")
    banda.set(qn("w:fill"), "D9D9D9")
    celula._p.get_or_add_pPr().append(banda)
    document.add_paragraph("Cargo de teste")
    buffer = BytesIO()
    document.save(buffer)
    original = buffer.getvalue()

    marcilio = Document(BytesIO(complementar_minuta_pca(
        original, sigla="MTTF", structure=_estrutura(EMENTA_CURTA), origem_modelo=True,
        data_assinatura="João Pessoa, 21 de outubro de 2026.",
    )))
    ementa_m = _paragrafo(marcilio, "Ementa:")
    assert _recuo(ementa_m, "w:left") == pytest.approx(8, abs=0.05)
    assert ementa_m.runs[0].italic is True
    assert ementa_m.runs[-1].italic is not True
    assert ementa_m.runs[-1].font.size.pt == pytest.approx(10)
    relato = _paragrafo(marcilio, "Relato objetivo")
    assert relato.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert relato.runs[0].bold is not True
    assert relato.runs[0].font.size.pt == pytest.approx(12)
    assert _fonte(relato.runs[0]) == "Calibri"
    assert _paragrafo(marcilio, "João Pessoa,").alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert _paragrafo(marcilio, "NOME DE TESTE").alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert relato._p.find(qn("w:pPr")).find(qn("w:shd")) is None
    assert _paragrafo(marcilio, "NOME DE TESTE")._p.find(qn("w:pPr")).find(qn("w:shd")) is None
    assert marcilio.tables[0].cell(0, 0).paragraphs[0]._p.find(qn("w:pPr")).find(qn("w:shd")) is not None

    isabella = Document(BytesIO(complementar_minuta_pca(
        original, sigla="IBMF", structure=_estrutura(EMENTA_CURTA), origem_modelo=True,
    )))
    ementa_i = _paragrafo(isabella, "Ementa:")
    assert _recuo(ementa_i, "w:left") == pytest.approx(5, abs=0.05)
    assert ementa_i.runs[0].bold is not True
    assert ementa_i.runs[0].font.size.pt == pytest.approx(11.5)
    assert _fonte(ementa_i.runs[0]) == "Calibri"
    corpo_i = _paragrafo(isabella, "Relato objetivo")
    assert corpo_i.runs[0].bold is not True
    assert corpo_i.runs[0].font.size.pt == pytest.approx(13)
    assert _recuo(corpo_i, "w:firstLine") == pytest.approx(2.5, abs=0.05)


def test_sem_modelo_mantem_o_fallback_existente_e_isola_especialidades():
    class Biblioteca:
        specialty_id = PREFEITURAS
        repository = None
        storage = None

    luciano = Biblioteca()
    luciano.procurador = exigir_procurador("LAF")
    with pytest.raises(MinutasDocxError, match="modelo visual"):
        resolver_para_exportacao(luciano)
    elvira = Biblioteca()
    elvira.procurador = exigir_procurador("ESPO")
    assert resolver_para_exportacao(elvira) is None

    bruto, _gerado = _aplicar("LAF", EMENTA_CURTA)
    camara = _servico(CAMARAS, "LAF")._complementar_pca(
        bruto, _estrutura(EMENTA_CURTA), False, _alvo(), _perfil(),
    )
    rpps = _servico(RPPS, "ESPO")._complementar_pca(
        bruto, _estrutura(EMENTA_CURTA), False, _alvo(), _perfil(),
    )
    assert camara == bruto
    assert rpps == bruto
    contextual = _servico("laf.pca.prefeituras", "LAF")._complementar_pca(
        bruto, _estrutura(EMENTA_CURTA), False, _alvo(), _perfil(),
    )
    assert contextual != bruto
    assert Document(BytesIO(contextual)).sections[0]._sectPr.xml == Document(BytesIO(bruto)).sections[0]._sectPr.xml


def test_falha_do_perfil_devolve_o_docx_original(monkeypatch):
    bruto, _gerado = _aplicar("SBBQ", EMENTA_CURTA)

    def falha(*_args, **_kwargs):
        raise RuntimeError("perfil indisponível")

    monkeypatch.setattr("services.minutas_pca_formato._complementar", falha)
    assert complementar_minuta_pca(
        bruto, sigla="SBBQ", structure=_estrutura(EMENTA_CURTA), origem_modelo=False,
    ) == bruto
    assert complementar_minuta_pca(
        b"nao-e-docx", sigla="SBBQ", structure=_estrutura(EMENTA_CURTA), origem_modelo=False,
    ) == b"nao-e-docx"
    assert complementar_minuta_pca(
        bruto, sigla="INEXISTENTE", structure=_estrutura(EMENTA_CURTA), origem_modelo=False,
    ) == bruto


def test_perfis_cobrem_os_procuradores_ativos_sem_ler_as_referencias():
    assert {pessoa.sigla for pessoa in listar_ativos()} == set(PERFIS)
    fonte = MODULO.read_text(encoding="utf-8")
    assert "REFERENCIAS_LOCAIS" not in fonte
    assert "PCA_PREFEITURAS" not in fonte


def _sombrear_paragrafo(paragraph, fill):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(sombra)


def _alvo_atual():
    return {
        "numero_processo": "00000/00",
        "origem": "Prefeitura Municipal de Teste Alfa",
        "responsavel": "Responsável de Teste Beta",
        "exercicio": "2099",
        "data_parecer": "2026-10-21",
    }


def _perfil_real(sigla):
    procurador = exigir_procurador(sigla)
    if sigla == "BLTC":
        return assinatura_bradson(procurador)
    perfil = assinatura_neutra(procurador)
    perfil["localidade"] = "João Pessoa"
    return perfil


def _pacote_textual(blob):
    """Texto visível do corpo, do cabeçalho e do rodapé, na ordem dos nós."""
    document = Document(BytesIO(blob))
    partes = ["".join(node.text or "" for node in document.element.iter(qn("w:t")))]
    for section in document.sections:
        for part in (section.header, section.footer):
            partes.append("".join(node.text or "" for node in part._element.iter(qn("w:t"))))
    return "\n".join(partes)


def _preencher(blob, sigla, ementa):
    alvo = _alvo_atual()
    perfil = _perfil_real(sigla)
    data = _signature_date(alvo, perfil)
    renderizado = renderizar(blob, alvo, _estrutura(ementa), perfil, _draft_header_lines(alvo))
    gerado = complementar_minuta_pca(
        renderizado,
        sigla=sigla,
        structure=_estrutura(ementa),
        origem_modelo=True,
        data_assinatura=data,
        identificacao=alvo,
    )
    return gerado, perfil, data


def test_citacao_conserva_numero_e_o_cabecalho_recebe_o_processo_atual():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "PROCESSO TC N.º 02529/25"
    document.add_paragraph("Processo TC Nº 02529/25")
    document.add_paragraph("Origem: Prefeitura Municipal de Joca Claudino")
    document.add_paragraph("Gestor: Rinaldo Cipriano de Sousa")
    document.add_paragraph("Exercício: 2024")
    document.add_paragraph("Ementa: texto de exemplo do modelo.")
    document.add_paragraph(
        "A jurisprudência menciona o processo 02529/25 apenas como citação "
        "e esse número precisa permanecer neste parágrafo de fundamentação."
    )
    buffer = BytesIO()
    document.save(buffer)
    gerado = Document(BytesIO(complementar_minuta_pca(
        buffer.getvalue(),
        sigla="LAF",
        structure=_estrutura(EMENTA_CURTA),
        origem_modelo=True,
        data_assinatura="João Pessoa, 21 de outubro de 2026.",
        identificacao=_alvo_atual(),
    )))
    cabecalho = gerado.sections[0].header.paragraphs[0].text
    assert "02529/25" not in cabecalho
    assert "00000/00" in cabecalho
    assert any("Teste Alfa" in item.text for item in gerado.paragraphs)
    assert any("Teste Beta" in item.text for item in gerado.paragraphs)
    assert any(item.text.startswith("Exercício:") and "2099" in item.text for item in gerado.paragraphs)
    citacao = _paragrafo(gerado, "A jurisprudência menciona")
    assert "02529/25" in citacao.text
    assert "Joca Claudino" not in "\n".join(item.text for item in gerado.paragraphs[:6])


def test_data_de_elvira_nao_herda_o_paragrafo_que_so_menciona_o_mes():
    document = Document()
    corpo = document.add_paragraph("Texto antigo do relatório, longo o bastante para representar o corpo.")
    corpo.paragraph_format.first_line_indent = Cm(3)
    corpo.runs[0].font.name = "Arial"
    corpo.runs[0].font.size = Pt(12)
    armadilha = document.add_paragraph("Em outubro de 2024 houve registro do fato neste parágrafo.")
    armadilha.runs[0].italic = True
    armadilha.runs[0].font.size = Pt(11)
    armadilha.paragraph_format.left_indent = Cm(4)
    data = document.add_paragraph("João Pessoa, 21 de outubro de 2026.")
    data.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    data.paragraph_format.first_line_indent = Cm(3)
    data.runs[0].font.name = "Arial"
    data.runs[0].font.size = Pt(12)
    data.runs[0].italic = False
    nome = document.add_paragraph("ELVIRA SAMARA PEREIRA DE OLIVEIRA")
    nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nome.runs[0].bold = True
    nome.runs[0].italic = True
    document.add_paragraph("Procuradora-Geral")
    buffer = BytesIO()
    document.save(buffer)
    alvo = _alvo_atual()
    perfil = _perfil_real("ESPO")
    renderizado = renderizar(
        buffer.getvalue(), alvo, _estrutura(EMENTA_CURTA), perfil, _draft_header_lines(alvo),
    )
    minuta = Document(BytesIO(renderizado))
    data_render = _paragrafo(minuta, "João Pessoa,")
    assert data_render.runs[0].italic is not True
    assert data_render.runs[0].font.size.pt == pytest.approx(12)
    gerado = Document(BytesIO(complementar_minuta_pca(
        renderizado,
        sigla="ESPO",
        structure=_estrutura(EMENTA_CURTA),
        origem_modelo=True,
        data_assinatura=_signature_date(alvo, perfil),
        identificacao=alvo,
    )))
    data_final = _paragrafo(gerado, "João Pessoa,")
    assert data_final.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert data_final.runs[0].italic is not True
    assert data_final.runs[0].font.size.pt == pytest.approx(12)
    assert _fonte(data_final.runs[0]) == "Arial"
    assert _recuo(data_final, "w:firstLine") == pytest.approx(3, abs=0.08)
    assert _recuo(_paragrafo(gerado, "Relato objetivo"), "w:firstLine") == pytest.approx(3, abs=0.08)
    assert sum(1 for item in gerado.paragraphs if _e_data(item.text)) == 1


def test_faixa_sombreada_permanece_e_nao_pinta_o_corpo():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "PROCESSO TC Nº 02316/25"
    document.add_paragraph("Natureza: PCA - Prestação de Contas Anuais")
    document.add_paragraph("Jurisdicionado: Prefeitura Municipal de Caiçara")
    document.add_paragraph("Exercício: 2024")
    faixa = document.add_paragraph("PARECER")
    _sombrear_paragrafo(faixa, "CCCCCC")
    titulo = document.add_paragraph("Irregularidade antiga do modelo")
    titulo.runs[0].bold = True
    titulo.runs[0].font.name = "Arial Narrow"
    titulo.runs[0].font.size = Pt(14)
    _sombrear_paragrafo(titulo, "D9D9D9")
    corpo = document.add_paragraph(
        "Texto antigo do relatório deste modelo, longo o bastante para ser o protótipo do corpo narrativo."
    )
    corpo.runs[0].font.name = "Calibri"
    corpo.runs[0].font.size = Pt(12)
    data = document.add_paragraph("João Pessoa,26 de ago. de 2026.")
    data.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    data.runs[0].font.name = "Calibri"
    nome = document.add_paragraph("Nome Antigo do Modelo")
    nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nome.runs[0].bold = False
    nome.runs[0].font.name = "Calibri"
    nome.runs[0].font.size = Pt(12)
    cargo = document.add_paragraph("Cargo antigo do modelo")
    cargo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cargo.runs[0].font.name = "Calibri"
    document.add_paragraph("ecad")
    buffer = BytesIO()
    document.save(buffer)
    gerado, perfil, _data = _preencher(buffer.getvalue(), "SBBQ", EMENTA_LONGA)
    minuta = Document(BytesIO(gerado))
    faixa_nova = next(item for item in minuta.paragraphs if _eh_faixa_paragrafo(item._p))
    assert faixa_nova._p.find(qn("w:pPr")).find(qn("w:shd")) is not None
    assert re.search(r"\d", faixa_nova.text) is None
    relato = _paragrafo(minuta, "Relato objetivo")
    assert relato._p.find(qn("w:pPr")).find(qn("w:shd")) is None
    titulo_novo = _paragrafo(minuta, TITULO)
    assert titulo_novo._p.find(qn("w:pPr")).find(qn("w:shd")) is not None
    assert titulo_novo.runs[0].bold is True
    nome_novo = _paragrafo(minuta, perfil["nome"])
    assert nome_novo.runs[0].bold is not True
    assert _fonte(nome_novo.runs[0]) != "Arial Narrow"
    assert "ecad" not in _pacote_textual(gerado)
    assert "02316/25" not in _pacote_textual(gerado)
    assert "Caiçara" not in _pacote_textual(gerado)
    assert "00000/00" in minuta.sections[0].header.paragraphs[0].text
    assert sum(1 for item in minuta.paragraphs if _e_data(item.text)) == 1
    assert "ago." not in "\n".join(item.text for item in minuta.paragraphs)

    repetido = complementar_minuta_pca(
        gerado,
        sigla="SBBQ",
        structure=_estrutura(EMENTA_LONGA),
        origem_modelo=True,
        data_assinatura=_data,
        identificacao=_alvo_atual(),
    )
    outra = Document(BytesIO(repetido))
    assert [item.text for item in outra.paragraphs] == [item.text for item in minuta.paragraphs]
    assert sum(1 for item in outra.paragraphs if item.text.strip().casefold().startswith("ementa")) == 1
    assert sum(1 for item in outra.paragraphs if _e_data(item.text)) == 1


def test_servico_repassa_a_identificacao_do_processo_atual():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "PROCESSO TC N.º 02529/25"
    document.add_paragraph("Relato objetivo dos atos de instrução.")
    buffer = BytesIO()
    document.save(buffer)
    gerado = _servico(PREFEITURAS, "LAF")._complementar_pca(
        buffer.getvalue(),
        _estrutura(EMENTA_CURTA),
        True,
        _alvo_atual(),
        _perfil_real("LAF"),
    )
    minuta = Document(BytesIO(gerado))
    assert "00000/00" in minuta.sections[0].header.paragraphs[0].text
    assert "02529/25" not in minuta.sections[0].header.paragraphs[0].text
    reaberto = BytesIO()
    minuta.save(reaberto)
    assert Document(BytesIO(reaberto.getvalue())).paragraphs


def _referencias_locais():
    pasta = Path(__file__).resolve().parents[1] / "REFERENCIAS_LOCAIS" / "PCA_PREFEITURAS"
    chaves = (
        ("bradson", "BLTC"),
        ("luciano", "LAF"),
        ("manoel", "MASN"),
        ("elvira", "ESPO"),
        ("isabella", "IBMF"),
        ("sheyla", "SBBQ"),
        ("marc", "MTTF"),
    )
    if not pasta.is_dir():
        return []
    achados = []
    for chave, sigla in chaves:
        for caminho in sorted(pasta.glob("*.docx")):
            if chave in caminho.name.casefold():
                achados.append((str(caminho), sigla))
                break
    return achados


_REFERENCIAS = _referencias_locais() or [(None, None)]


def _venenos(caminho):
    origem = Document(caminho)
    zonas = []
    for section in origem.sections:
        for part in (section.header, section.footer):
            zonas.extend(item.text for item in part.paragraphs)
    zonas.extend(item.text for item in origem.paragraphs[:12])
    texto = "\n".join(zonas)
    numeros = {
        re.sub(r"\s+", "", item)
        for item in re.findall(r"\d{2,}\s*/\s*\d{2,}", texto)
    }
    numeros.discard("00000/00")
    novo = "prefeitura municipal de teste alfa responsável de teste beta"
    lugares = set()
    for match in re.finditer(
        r"(?:Jurisdicionad[oa]s?|Origem|Entidade|Gestor|Responsável)\s*:\s*(.+)",
        texto,
        re.IGNORECASE,
    ):
        valor = " ".join(match.group(1).split()).strip(" .")
        if valor and valor.casefold() not in novo:
            lugares.add(valor)
        if " de " in valor.casefold():
            token = valor.split(" de ")[-1].strip(" .")
            if len(token) > 4 and token.casefold() not in novo:
                lugares.add(token)
    iniciais = {
        item.text.strip()
        for item in origem.paragraphs[-4:]
        if re.fullmatch(r"[a-z]{2,5}", item.text.strip())
    }
    return numeros, lugares, iniciais


def _sombra(paragraph):
    properties = paragraph._p.find(qn("w:pPr"))
    if properties is None:
        return None
    return properties.find(qn("w:shd"))


@pytest.mark.parametrize(
    ("caminho", "sigla"),
    _REFERENCIAS,
    ids=[sigla or "ausente" for _caminho, sigla in _REFERENCIAS],
)
def test_modelo_representativo_nao_traz_o_processo_antigo(caminho, sigla):
    if caminho is None:
        pytest.skip("referências locais de PCA ausentes")
    origem = Document(caminho)
    desenhos = sum(
        len(list(section.header._element.iter(qn("w:drawing"))))
        for section in origem.sections
    )
    numeros, lugares, iniciais = _venenos(caminho)
    gerado, perfil, data = _preencher(Path(caminho).read_bytes(), sigla, EMENTA_LONGA)
    minuta = Document(BytesIO(gerado))
    texto = _pacote_textual(gerado)
    for numero in numeros:
        assert numero not in re.sub(r"\s+", "", texto)
    for lugar in lugares:
        assert lugar.casefold() not in texto.casefold()
    for inicial in iniciais:
        assert inicial not in texto
    assert "00000/00" in texto
    assert "Teste Alfa" in texto
    assert "2099" in texto
    identificacao_origem = "\n".join(item.text for item in origem.paragraphs[:15])
    if re.search(r"(?m)^(?:respons|gestor)\b", identificacao_origem, re.IGNORECASE):
        assert "Teste Beta" in texto
    assert perfil["nome"].upper() in texto.upper()
    assert sum(1 for item in minuta.paragraphs if _e_data(item.text)) == 1
    ementas = [
        item for item in minuta.paragraphs
        if item.text.strip().casefold().startswith("ementa")
    ]
    celulas = [
        cell.text for table in minuta.tables for row in table.rows for cell in row.cells
        if "ementa" in cell.text.casefold()
    ]
    assert len(ementas) + len(celulas) == 1
    relato = _paragrafo(minuta, "Relato objetivo")
    assert _sombra(relato) is None
    assert sum(
        len(list(section.header._element.iter(qn("w:drawing"))))
        for section in minuta.sections
    ) == desenhos
    assert minuta.sections[0].top_margin == origem.sections[0].top_margin
    assert minuta.sections[0].left_margin == origem.sections[0].left_margin
    assert minuta.sections[0].right_margin == origem.sections[0].right_margin
    if sigla == "BLTC":
        assert len(minuta.tables) >= 2
        assert celulas and EMENTA_LONGA[:40] in celulas[0]
        assert any("PARECER" in table.cell(0, 0).text.upper() for table in minuta.tables)
        assert "PAGE" in minuta.sections[0].footer._element.xml
    if sigla != "BLTC":
        assert ementas and EMENTA_LONGA[:40] in ementas[0].text
    if sigla == "LAF":
        ementa = ementas[0]
        tabela = next(table for table in minuta.tables if "PARECER" in table.cell(0, 0).text.upper())
        filhos = [child for child in minuta.element.body if child.tag != qn("w:sectPr")]
        assert filhos.index(ementa._p) < filhos.index(tabela._tbl)
        assert ementa.runs[0].italic is True
        assert _recuo(ementa, "w:left") == pytest.approx(6.35, abs=0.08)
    if sigla == "IBMF":
        assert "cabedelo" not in texto.casefold()
        assert "2022" not in texto
        textos_paragrafos = [item.text.strip() for item in minuta.paragraphs]
        origem_idx = next(
            index for index, item in enumerate(textos_paragrafos)
            if item.upper().startswith("ORIGEM")
        )
        ementa_idx = next(
            index for index, item in enumerate(textos_paragrafos)
            if item.casefold().startswith("ementa")
        )
        intervalo = textos_paragrafos[origem_idx + 1:ementa_idx]
        assert intervalo and all(not item for item in intervalo)
        assert len(intervalo) <= 2
        faixa_idx = next(
            index for index, item in enumerate(minuta.paragraphs)
            if _eh_faixa_paragrafo(item._p)
        )
        assert ementa_idx < faixa_idx
    if sigla == "MASN":
        assert "2099Jurisdicionado" not in texto
        assert "AlfaRespons" not in texto
    if sigla in {"MASN", "SBBQ", "MTTF"}:
        faixa = next(item for item in minuta.paragraphs if _eh_faixa_paragrafo(item._p))
        assert _sombra(faixa) is not None
        assert _sombra(relato) is None
    if sigla == "SBBQ":
        assert _sombra(_paragrafo(minuta, TITULO)) is not None
        nome = _paragrafo(minuta, perfil["nome"])
        assert nome.runs[0].bold is not True
        assert _fonte(nome.runs[0]) != "Arial Narrow"
    if sigla == "ESPO":
        data_item = _paragrafo(minuta, "João Pessoa,")
        assert data_item.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert data_item.runs[0].italic is not True
        assert data_item.runs[0].font.size.pt == pytest.approx(12)
        assert _fonte(data_item.runs[0]) == "Arial"
        assert _recuo(data_item, "w:firstLine") == pytest.approx(3, abs=0.08)
    if not PERFIS[sigla].get("preservar_geometria"):
        spec = PERFIS[sigla]["corpo"]
        assert _fonte(relato.runs[0]) == spec["font"]
        assert relato.runs[0].font.size.pt == pytest.approx(spec["size"])
        assert _recuo(relato, "w:firstLine") == pytest.approx(spec["first_cm"], abs=0.08)
        assert relato.runs[0].bold is not True
        assert relato.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    repetido = complementar_minuta_pca(
        gerado,
        sigla=sigla,
        structure=_estrutura(EMENTA_LONGA),
        origem_modelo=True,
        data_assinatura=data,
        identificacao=_alvo_atual(),
    )
    outra = Document(BytesIO(repetido))
    assert [item.text for item in outra.paragraphs] == [item.text for item in minuta.paragraphs]
    assert len(outra.tables) == len(minuta.tables)
    assert sum(1 for item in outra.paragraphs if _e_data(item.text)) == 1
    reaberto = BytesIO()
    minuta.save(reaberto)
    assert Document(BytesIO(reaberto.getvalue())).sections
    if sigla == "LAF":
        curta, _perfil_curto, _data_curta = _preencher(Path(caminho).read_bytes(), sigla, EMENTA_CURTA)
        curta_doc = Document(BytesIO(curta))
        ementa_curta = next(
            item for item in curta_doc.paragraphs
            if item.text.strip().casefold().startswith("ementa")
        )
        assert EMENTA_CURTA in ementa_curta.text
        assert "02529" not in _pacote_textual(curta)
