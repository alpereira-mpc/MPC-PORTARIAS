from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import sqlite3
import pytest
from docx import Document
from database.oficios import OficiosStore
from database.store import Store
from document_generator.oficios import generate
from services.oficios import attention, validate_upload, safe_name
from tests.test_postgresql import pg_url, pg_store


def sample(service, member=None, year=2026):
    series = service.series()
    member = member or next(s["membro_id"] for s in series if s["sigla"] == "PROGE")
    return dict(
        direcao="ENVIADO",
        membro_id=member,
        data=f"{year}-09-10",
        assunto="Solicitação de informações",
        destinatario="DESTINATÁRIO TESTE",
        cargo="Promotor de Justiça",
        instituicao="Instituição de teste",
        tratamento="A Sua Excelência o Senhor",
        vocativo="Excelentíssimo Senhor Promotor de Justiça,",
        corpo="Primeiro parágrafo.\n\nSegundo parágrafo.",
        fechamento="Atenciosamente,",
    )


def files(record, series, number):
    return [
        ("docx", "teste.docx", generate(record, series, number)),
        ("pdf", "teste.pdf", b"%PDF-test"),
    ]


def ready(store):
    service = OficiosStore(store)
    service.confirm_sequence("PROGE", 2026, 8, True)
    service.confirm_sequence("BTLC", 2026, 2, True)
    return service


def test_number_lifecycle(store):
    s = ready(store)
    r = sample(s)
    identifier = s.save(r)
    assert s.sequence("PROGE", 2026)["proximo"] == 8
    series = next(x for x in s.series() if x["sigla"] == "PROGE")
    assert (
        b"PK"
        == generate({**r, "signatario": "Teste", "cargo_base": "Procuradora"}, series)[
            :2
        ]
    )
    assert s.sequence("PROGE", 2026)["proximo"] == 8
    result = s.finalize(identifier, files)
    assert result["numero"] == 8 and result["status"] == "Gerado"
    assert s.finalize(identifier, files)["numero"] == 8
    assert s.sequence("PROGE", 2026)["proximo"] == 9
    s.update_status(identifier, "Cancelado", "Teste cancelamento")
    assert s.sequence("PROGE", 2026)["proximo"] == 9
    with pytest.raises(ValueError):
        s.delete_draft(identifier, True)
    with pytest.raises(ValueError):
        s.update_status(identifier, "Enviado", sent="2026-09-10")
    assert s.get(identifier)["numero"] == 8


def test_independent_series_years(store):
    s = ready(store)
    for code, year, expected in [
        ("PROGE", 2026, 8),
        ("BTLC", 2026, 2),
        ("PROGE", 2027, 1),
    ]:
        member = next(x["membro_id"] for x in s.series() if x["sigla"] == code)
        if year == 2027:
            s.confirm_sequence(code, year, 1, True)
        assert s.finalize(s.save(sample(s, member, year)), files)["numero"] == expected


def test_generator_failure_rolls_back(store):
    s = ready(store)
    identifier = s.save(sample(s))

    def fail(*args):
        raise RuntimeError("conversion failure")

    with pytest.raises(RuntimeError):
        s.finalize(identifier, fail)
    assert s.get(identifier)["status"] == "Rascunho"
    assert s.sequence("PROGE", 2026)["proximo"] == 8
    assert not s.files(identifier)


@pytest.mark.parametrize("same", [False, True])
def test_concurrent_sqlite(store, same):
    concurrent(store, same)


def concurrent(store, same):
    s = ready(store)
    ids = [s.save(sample(s)) for _ in range(4)]
    if same:
        ids = [ids[0]] * 4
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda x: s.finalize(x, files)["numero"], ids))
    assert sorted(results) == ([8] * 4 if same else [8, 9, 10, 11])
    assert s.sequence("PROGE", 2026)["proximo"] == (9 if same else 12)


@pytest.mark.parametrize("same", [False, True])
def test_concurrent_postgresql(pg_store, same):
    concurrent(pg_store, same)


def test_postgresql_control(pg_store):
    control(pg_store)


def test_sqlite_control(store):
    control(store)


def control(store):
    s = ready(store)
    member = s.series()[0]["membro_id"]
    d = Document()
    d.add_paragraph("Original recebido")
    out = BytesIO()
    d.save(out)
    identifier = s.save(
        dict(
            direcao="RECEBIDO",
            numero_externo="12/2026",
            remetente="Pessoa",
            instituicao="Órgão",
            data="2026-09-01",
            data_recebimento="2026-09-10",
            membros=[member],
            assunto="Assunto recebido",
            prazo="2026-09-11",
        ),
        uploads=[("original.docx", out.getvalue())],
    )
    metadata = s.files(identifier)
    assert s.download(metadata[0]["id"]) == out.getvalue()
    assert set(metadata[0]) == {"id", "nome", "tipo", "tamanho", "incluida"}
    assert "conteudo" not in str(s.list()) and "payload" not in s.list()[0]
    s.update_status(
        identifier, "Em análise", "Encaminhado ao gabinete.", due="2026-09-12"
    )
    assert s.get(identifier)["prazo"] == "2026-09-12"
    assert s.movements(identifier)[-1]["anterior"] == "Recebido"
    draft = s.save({**sample(s), "responde_a": identifier})
    s.finalize(draft, files)
    assert s.list(related=identifier)[0]["id"] == draft
    assert s.get(identifier)["status"] == "Em análise"
    assert (
        s.list(direction="RECEBIDO", member=member, search="Pessoa")[0]["id"]
        == identifier
    )
    assert not s.list(direction="RECEBIDO", search="inexistente")
    assert s.overview(2026)["recebidos"] == 1


def test_baseline_safeguards(store):
    s = OficiosStore(store)
    assert s.sequence("PROGE", 2026) == {"proximo": 8, "confirmada": False}
    identifier = s.save(sample(s))
    with pytest.raises(ValueError):
        s.finalize(identifier, files)
    for number, confirmed in [(8, False), (7, True), (-1, True)]:
        with pytest.raises(ValueError):
            s.confirm_sequence("PROGE", 2026, number, confirmed)
    s.confirm_sequence("PROGE", 2026, 10, True)
    with pytest.raises(ValueError):
        s.confirm_sequence("PROGE", 2026, 9, True)
    assert s.finalize(identifier, files)["numero"] == 10
    with pytest.raises(ValueError):
        s.confirm_sequence("PROGE", 2026, 10, True)


def test_edit_delete_reopen(store):
    s = ready(store)
    r = sample(s)
    identifier = s.save(r)
    s.save({**r, "assunto": "Alterado"}, identifier)
    with pytest.raises(ValueError):
        s.delete_draft(identifier)
    reopened = OficiosStore(Store(store.path))
    assert reopened.get(identifier)["assunto"] == "Alterado"
    reopened.delete_draft(identifier, True)
    assert not reopened.list()
    assert reopened.sequence("PROGE", 2026)["proximo"] == 8


@pytest.mark.parametrize(
    "series,digits,prefix",
    [("PROGE", 3, "Ofício MPC/PB - PROGE n."), ("BTLC", 2, "Ofício BTLC-MPC-PB nº")],
)
def test_documents(store, series, digits, prefix):
    s = ready(store)
    config = next(x for x in s.series() if x["sigla"] == series)
    person = next(
        x for x in store.catalog("procuradores") if x["id"] == config["membro_id"]
    )
    r = {
        **sample(s, person["id"]),
        "signatario": person["nome"],
        "cargo_base": person["cargo_base"],
    }
    doc = Document(BytesIO(generate(r, config, 8)))
    text = "\n".join(p.text for p in doc.paragraphs)
    for expected in [
        prefix,
        str(8).zfill(digits) + "/2026",
        "10 de setembro de 2026.",
        r["destinatario"],
        r["assunto"],
        r["vocativo"],
        r["signatario"].upper(),
    ]:
        assert expected in text
    assert any(p.text == "Primeiro parágrafo." for p in doc.paragraphs)
    assert any(p.text == "Segundo parágrafo." for p in doc.paragraphs)
    assert "BLTC" not in text
    assert doc.sections[0].header._element.xml


@pytest.mark.parametrize(
    "name,content",
    [
        ("file.exe", b"a"),
        ("a.pdf", b"bad"),
        ("a.docx", b"PKbad"),
        pytest.param("a.pdf", b"%PDF-" + b"x" * (10 * 1024 * 1024), id="oversized"),
    ],
)
def test_invalid_upload(name, content):
    with pytest.raises(ValueError):
        validate_upload(name, content)


@pytest.mark.parametrize(
    "offset,expected", [(-1, "vencido"), (0, "próximo"), (7, "próximo"), (8, "")]
)
def test_deadlines(offset, expected):
    r = {
        "direcao": "RECEBIDO",
        "status": "Recebido",
        "prazo": (date.today() + timedelta(days=offset)).isoformat(),
    }
    assert expected in attention(r)
    r["status"] = "Concluído"
    assert not attention(r)


def test_series_unconfirmed(store):
    s = ready(store)
    laf = next(x for x in s.series() if x["sigla"] == "LAF")
    identifier = s.save(sample(s, laf["membro_id"]))
    s.confirm_sequence("LAF", 2026, 1, True)
    with pytest.raises(ValueError):
        s.finalize(identifier, files)
    s.configure_series(
        laf["membro_id"], "LAF", "BTLC", "Ofício LAF nº {numero}/{ano}", 2, True
    )
    assert s.finalize(identifier, files)["numero"] == 1
    with pytest.raises(ValueError):
        s.configure_series(
            laf["membro_id"], "LAF", "PROGE", "Ofício LAF {numero}/{ano}", 3, True
        )


def test_uniqueness_database(store):
    s = ready(store)
    a = s.save(sample(s))
    b = s.save(sample(s))
    s.finalize(a, files)
    with pytest.raises(sqlite3.IntegrityError):
        with store.connection() as c:
            c.execute("UPDATE oficios SET numero=8,status='Gerado' WHERE id=?", (b,))


def test_home_and_oficios_navigation(store, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from database.store import ROOT

    monkeypatch.setattr("database.store.Store", lambda: store)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    app.button(key="open_oficios").click().run()
    assert not app.exception and not app.error
    for page in [
        "Novo Ofício",
        "Enviados",
        "Recebidos",
        "Acompanhamento",
        "Visão Geral",
    ]:
        app.radio(key="oficio_page").set_value(page).run()
        assert not app.exception and not app.error


@pytest.mark.parametrize(
    "scenario",
    [
        test_number_lifecycle,
        test_independent_series_years,
        test_generator_failure_rolls_back,
        test_baseline_safeguards,
        test_series_unconfirmed,
    ],
)
def test_postgresql_numbering_contract(pg_store, scenario):
    scenario(pg_store)


def test_filters_and_attention_pagination(store):
    s = ready(store)
    member = s.series()[0]["membro_id"]
    for i in range(4):
        s.save(
            dict(
                direcao="RECEBIDO",
                numero_externo=str(i),
                remetente="Origem",
                instituicao="Órgão",
                data="2026-09-10",
                data_recebimento="2026-09-10",
                membros=[member],
                assunto=f"Assunto {i}",
                status="Recebido" if i % 2 else "Em análise",
            )
        )
    results = s.list(direction="RECEBIDO", attention_only=True, limit=1)
    assert len(results) == 1 and results[0]["status"] == "Em análise"
    assert s.list(attention_only=True, offset=1, limit=1)[0]["id"] != results[0]["id"]
    assert not s.list(attention_only=True, offset=2, limit=1)
    assert (
        len(s.list(subject="Assunto", start="2026-09-10", end="2026-09-10", year=2026))
        == 4
    )
    assert not s.list(start="2026-09-11")


def test_metadata_queries_exclude_binary(store, monkeypatch):
    from contextlib import contextmanager

    s = ready(store)
    identifier = s.save(sample(s))
    s.finalize(identifier, files)
    statements = []
    original = store.connection

    @contextmanager
    def traced(**kwargs):
        with original(**kwargs) as c:
            c.set_trace_callback(statements.append)
            yield c

    monkeypatch.setattr(store, "connection", traced)
    s.list()
    s.overview(2026)
    s.files(identifier)
    s.get(identifier)
    assert all("conteudo" not in sql.lower() for sql in statements)
    s.download(s.files(identifier)[0]["id"])
    assert any("select conteudo" in sql.lower() for sql in statements)


def test_pdf_upload_and_fresh_process(store):
    from pypdf import PdfWriter
    import subprocess, sys

    s = ready(store)
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    content = output.getvalue()
    member = s.series()[0]["membro_id"]
    identifier = s.save(
        dict(
            direcao="RECEBIDO",
            numero_externo="12/2026",
            remetente="Pessoa",
            instituicao="Órgão",
            data="2026-09-10",
            data_recebimento="2026-09-10",
            membros=[member],
            assunto="Original",
        ),
        uploads=[("original.pdf", content)],
    )
    file = s.files(identifier)[0]
    code = "from database.store import Store; from database.oficios import OficiosStore; import sys,hashlib; print(hashlib.sha256(OficiosStore(Store(sys.argv[1])).download(sys.argv[2])).hexdigest())"
    import hashlib

    result = subprocess.run(
        [sys.executable, "-c", code, str(store.path), file["id"]],
        capture_output=True,
        text=True,
        check=True,
    )
    assert result.stdout.strip() == hashlib.sha256(content).hexdigest()


def test_invalid_relation_and_status(store):
    s = ready(store)
    draft = s.save(sample(s))
    with pytest.raises(ValueError):
        s.save({**sample(s), "responde_a": draft})
    with pytest.raises(ValueError):
        s.update_status(draft, "Enviado", sent="2026-09-10")
    s.finalize(draft, files)
    with pytest.raises(ValueError):
        s.update_status(draft, "Cancelado", "")
    with pytest.raises(ValueError):
        s.update_status(draft, "Enviado")
    s.update_status(draft, "Enviado", sent="2026-09-10")
    assert s.get(draft)["data_envio"] == "2026-09-10"
    s.update_status(draft, "Cancelado", "Sem efeito")
    assert (
        s.get(draft)["cancelada"]
        and s.movements(draft)[-1]["observacao"] == "Sem efeito"
    )


def test_migration_preserves_existing_data(store):
    before = store.catalog("procuradores")
    sequence = store.next_number(2026)
    s = ready(store)
    identifier = s.save(sample(s))
    OficiosStore(store)
    assert store.catalog("procuradores") == before
    assert store.next_number(2026) == sequence
    assert s.get(identifier)["status"] == "Rascunho"


def test_real_pdf_generation_when_enabled(store):
    import os

    if os.environ.get("MPC_TEST_OFFICE") != "1":
        pytest.skip(
            "Conversor local real habilitado apenas na validação visual explícita"
        )
    from pypdf import PdfReader

    s = ready(store)
    for code in ("PROGE", "BTLC"):
        member = next(x["membro_id"] for x in s.series() if x["sigla"] == code)
        identifier = s.save(sample(s, member))
        result = s.finalize(identifier)
        metadata = s.files(identifier)
        assert len(metadata) == 2
        pdf = next(x for x in metadata if x["tipo"] == "application/pdf")
        text = "\n".join(
            p.extract_text() for p in PdfReader(BytesIO(s.download(pdf["id"]))).pages
        )
        assert code in text and "2026" in text


def test_editor_save_reopen_finalize_ui(store, monkeypatch):
    from streamlit.testing.v1 import AppTest
    from database.store import ROOT

    s = ready(store)
    monkeypatch.setattr("database.store.Store", lambda: store)
    monkeypatch.setattr("document_generator.oficios.official_documents", files)
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    app.button(key="open_oficios").click().run()
    app.radio(key="oficio_page").set_value("Novo Ofício").run()
    for label, value in [
        ("Assunto", "Assunto UI"),
        ("Nome do destinatário", "Pessoa UI"),
    ]:
        next(x for x in app.text_input if x.label == label).set_value(value)
    next(x for x in app.text_area if x.label.startswith("Corpo do ofício")).set_value(
        "Texto UI.\n\nSegundo parágrafo."
    )
    next(x for x in app.button if x.label == "Salvar rascunho").click().run()
    assert not app.exception and not app.error
    identifier = s.list()[0]["id"]
    assert s.get(identifier)["numero"] is None
    app.radio(key="oficio_page").set_value("Enviados").run()
    app.button(key="open_oficio_" + identifier).click().run()
    app.button(key="edit_" + identifier).click().run()
    assert app.radio(key="oficio_page").value == "Novo Ofício"
    next(x for x in app.button if x.label == "Finalizar e gerar ofício").click().run()
    assert not app.exception and not app.error
    assert s.get(identifier)["numero"] == 8
