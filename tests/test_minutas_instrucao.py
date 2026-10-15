"""Instrução processual com número variável de PDFs, sem cadastro no acervo."""

from io import BytesIO

from reportlab.pdfgen import canvas
import inspect
import json
import logging

import pytest

from services.minutas_docx import ELVIRA_TEMPLATE, MinutasDocxService
from services.minutas_generation import MinutasGenerationError, MinutasGenerationService
from services.minutas_instrucao import (
    MinutasInstrucaoError,
    conferencia,
    consolidar,
    fontes,
    receber_pdf,
    sugerir_metadados_instrucao,
    versao,
)


def _pdf(*pages):
    output = BytesIO()
    document = canvas.Canvas(output)
    for text in pages:
        text_object = document.beginText(72, 760)
        for line in str(text).splitlines() or [""]:
            text_object.textLine(line)
        document.drawText(text_object)
        document.showPage()
    document.save()
    return output.getvalue()


def _doc(name, text, tipo, **extra):
    item = receber_pdf(name, _pdf(text), tipo)
    item.update(extra)
    return item


def _library():
    class Repository:
        def create(self, *_args, **_kwargs):
            raise AssertionError("a instrução não pode cadastrar documento")

        def add_file(self, *_args, **_kwargs):
            raise AssertionError("a instrução não pode cadastrar documento")

    class Library:
        repository = Repository()
        storage = object()

        def _authorize(self):
            return None

    return Library()


def _analysis(document):
    return {
        "resumo_processual": "Instrucao recebida para conferencia.",
        "sem_irregularidades_identificadas": False,
        "pontos_atencao_gerais": ["Conferir os documentos enviados."],
        "irregularidades": [{
            "id": "IRR-1",
            "apontamento": "Pagamento sem cobertura documental",
            "fatos_documentados": [{
                "descricao": "Pagamento sem cobertura",
                "fontes": [{
                    "fonte_id": document["fonte_id"],
                    "pagina": 1,
                    "trecho_suporte": (
                        "Pagamento sem cobertura documental no instituto "
                        "de previdencia municipal"
                    ),
                }],
            }],
            "normas_mencionadas": [],
            "argumentos_defesa": [],
            "pontos_controvertidos": ["Ha posicoes distintas na instrucao."],
            "hipoteses_juridicas": ["Conferir a cobertura documental."],
            "avaliacao_auditoria": {
                "descricao": "Apontamento registrado no documento.",
                "tratamento_defesa": "NAO_IDENTIFICADO",
                "fontes": [],
            },
        }],
    }


def _structure():
    return {
        "identificacao_processual": "Processo TC 03333/26",
        "ementa": "Ementa de teste da instrucao.",
        "relatorio": "Relatorio da instrucao processual.",
        "relatorio_paragrafos": ["Relatorio da instrucao processual."],
        "fundamentacao_geral": "Fundamentacao geral de teste.",
        "fundamentacao_geral_paragrafos": ["Fundamentacao geral de teste."],
        "topicos_irregularidades": [{
            "titulo": "Pagamento sem cobertura",
            "texto": "O apontamento permanece em conferencia.",
            "paragrafos": ["O apontamento permanece em conferencia."],
            "irregularidades_ids": ["IRR-1"],
            "avaliacoes_irregularidades": [{
                "irregularidade_id": "IRR-1",
                "tratamento_ministerial_defesa": "NAO_APLICAVEL",
                "fundamentos_divergencia": [],
                "estado_saneamento": "NAO_ALEGADO",
                "fontes_saneamento": [],
            }],
        }],
        "analise_sancionatoria": {
            "conclusao": "REVISAO_HUMANA_NECESSARIA",
            "irregularidades_ids": ["IRR-1"],
            "fundamentos_legais": [],
            "responsaveis": [],
            "dano_ao_erario": "REVISAO_HUMANA_NECESSARIA",
            "texto": "O cabimento de multa depende de revisão humana.",
        },
        "conclusao_pedidos": "Pede-se a conferencia.",
        "pedidos": ["Conferir a instrucao antes da assinatura."],
        "encerramento": "E o parecer, salvo melhor juizo.",
    }


def test_somente_relatorio_inicial_nao_inventa_etapa():
    document = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )
    review = conferencia([document])
    assert "apenas" in review["advertencias"][0]
    assert "Relatório Inicial" in review["advertencias"][0]
    assert "defesa alegou" not in " ".join(review["advertencias"])
    assert fontes([document])[0]["papel"] == "auditoria"


def test_somente_analise_de_defesa():
    document = _doc(
        "defesa.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "analise_defesa",
    )
    review = conferencia([document])
    assert "Relatório de Análise de Defesa" in review["advertencias"][0]
    assert fontes([document])[0]["papel"] == "auditoria"


def test_inicial_e_analise_de_defesa():
    documents = [
        _doc("inicial.pdf", "Relatorio inicial sem apontamento numerado.", "relatorio_inicial"),
        _doc("defesa.pdf", "Relatorio de analise da defesa.", "analise_defesa"),
    ]
    review = conferencia(documents)
    assert not any("apenas" in item for item in review["advertencias"])
    assert [item["tipo"] for item in fontes(documents)] == [
        "Relatório Inicial de Auditoria",
        "Relatório de Análise de Defesa",
    ]


def test_inicial_cota_defesa_e_complementacao():
    documents = [
        _doc("inicial.pdf", "Relatorio inicial.", "relatorio_inicial"),
        _doc("cota.pdf", "Cota ministerial de diligencia.", "cota"),
        _doc("defesa.pdf", "Analise da defesa.", "analise_defesa"),
        _doc("complemento.pdf", "Relatorio de complementacao.", "complementacao"),
    ]
    rows = fontes(documents)
    assert [row["papel"] for row in rows] == [
        "auditoria", "complementar", "auditoria", "auditoria",
    ]
    assert [row["ordem_processual"] for row in rows] == [1, 2, 3, 4]
    assert len({row["fonte_id"] for row in rows}) == 4
    assert all(row["sha256"] for row in rows)


def test_varios_relatorios_complementares_do_mesmo_tipo():
    documents = [
        _doc(f"complemento-{index}.pdf", f"Complementacao {index}.", "complementacao")
        for index in range(1, 4)
    ]
    rows = fontes(documents)
    assert len(rows) == 3
    assert {row["tipo"] for row in rows} == {"Relatório de Complementação"}
    assert [row["arquivo"] for row in rows] == [
        "complemento-1.pdf", "complemento-2.pdf", "complemento-3.pdf",
    ]


def test_ordem_informada_prevalece_sobre_a_data():
    documents = [
        _doc("recente.pdf", "Peca recente.", "despacho", data="2026-03-01"),
        _doc("antiga.pdf", "Peca antiga.", "informacao", data="2026-01-15"),
    ]
    review = conferencia(documents)
    assert any("não segue as datas" in item for item in review["advertencias"])
    assert [row["arquivo"] for row in fontes(documents)] == ["recente.pdf", "antiga.pdf"]


def test_saneamento_posterior_nao_prevalece_sozinho():
    documents = [
        _doc(
            "inicial.pdf",
            "Irregularidade 1 pagamento sem cobertura permanece",
            "relatorio_inicial",
        ),
        _doc(
            "complemento.pdf",
            "Irregularidade 1 pagamento foi sanada",
            "complementacao",
        ),
    ]
    review = conferencia(documents)
    assert review["evolucao"]
    assert "não" in review["evolucao"][0]["leitura"] or "Nenhum" in review["evolucao"][0]["leitura"]
    assert len(review["evolucao"][0]["posicoes"]) == 2


def test_divergencia_entre_relatorios_fica_registrada():
    documents = [
        _doc(
            "inicial.pdf",
            "Irregularidade 2 pagamento sem cobertura",
            "relatorio_inicial",
        ),
        _doc(
            "defesa.pdf",
            "Irregularidade 2 nao se confirma o pagamento",
            "analise_defesa",
        ),
    ]
    review = conferencia(documents)
    assert review["divergencias"]
    assert "sem escolha automática" in review["divergencias"][0]["leitura"]


def test_inclusao_ou_remocao_invalida_a_versao_anterior():
    first = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )
    second = _doc("cota.pdf", "Cota ministerial.", "cota")
    seen = {}

    def analyze(context):
        seen["ids"] = [item["fonte_id"] for item in context["documentos_processo_alvo"]]
        return {"analise": _analysis(first)}

    service = MinutasGenerationService(_library(), analyze=analyze, generate=lambda _context: {"estrutura": _structure()})
    analyzed = service.analyze_instruction([first], processo="03333/26", exercicio=2025)
    assert seen["ids"] == [first["fonte_id"]]
    with pytest.raises(MinutasGenerationError, match="alterada"):
        service.generate_instruction(
            [first, second], analyzed, processo="03333/26", exercicio=2025,
        )
    assert versao([first]) != versao([first, second])
    assert versao([first, second]) != versao([first])


def test_volume_acima_do_contexto_preserva_apontamento_e_lista_omissao():
    finding = "Irregularidade 3 pagamento sem cobertura permanece"
    pages = [{"texto": finding, "pagina": 1}]
    pages.extend(
        {"texto": "registro administrativo sem apontamento " * 30, "pagina": number}
        for number in range(2, 16)
    )
    source = {
        "fonte_id": "volume",
        "arquivo": "volume.pdf",
        "papel": "auditoria",
        "tipo": "Relatório Inicial de Auditoria",
        "ordem_processual": 1,
        "conteudo": pages,
    }
    kept, omissions = consolidar([source], budget=2500)
    assert any(finding in unit["texto"] for unit in kept[0]["conteudo"])
    assert 1 not in omissions[0]["paginas_nao_enviadas_integralmente"]
    assert omissions[0]["paginas_nao_enviadas_integralmente"]


def test_volume_impossivel_nao_trunca_em_silencio():
    source = {
        "fonte_id": "volume",
        "arquivo": "volume.pdf",
        "papel": "auditoria",
        "tipo": "Relatório Inicial de Auditoria",
        "ordem_processual": 1,
        "conteudo": [{
            "texto": "Irregularidade 9 " + ("cobertura " * 400),
            "pagina": 1,
        }],
    }
    with pytest.raises(MinutasInstrucaoError, match="excede o contexto"):
        consolidar([source], budget=120)


@pytest.mark.parametrize("quantidade", [1, 2, 4])
def test_analise_geracao_e_docx_aceitam_quantidade_variavel(quantidade):
    if not ELVIRA_TEMPLATE.is_file():
        pytest.skip("modelo Word indisponível")
    base = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )
    extras = [
        _doc("cota.pdf", "Cota ministerial de diligencia.", "cota"),
        _doc("defesa.pdf", "Analise da manifestacao apresentada.", "analise_defesa"),
        _doc("outro.pdf", "Informacao tecnica complementar.", "informacao"),
    ]
    documents = [base, *extras[: quantidade - 1]]
    captured = {}

    def analyze(context):
        captured["ids"] = [item["fonte_id"] for item in context["documentos_processo_alvo"]]
        captured["alvo"] = context["processo_alvo"]
        captured["instrucao"] = context["processo_alvo"]["instrucao_local"]
        return {"analise": _analysis(base)}

    def generate(context):
        captured["geracao"] = len(context["documentos_processo_alvo"])
        captured["dispositivo"] = context["orientacao_dispositivo"]
        captured["geracao_alvo"] = context["processo_alvo"]
        return {"estrutura": _structure()}

    hits = []

    def search(query, **kwargs):
        hits.append((query, kwargs.get("exclude_process_key")))
        return {"resultados": [
            {"papel": "parecer", "referencia_id": "parecer-1", "processo_chave": "0165326",
             "numero_processo": "01653/26", "trecho": "entendimento historico"},
            {"papel": "complementar", "referencia_id": "relatorio-local", "processo_chave": "0999926",
             "trecho": "nao e precedente"},
            {"papel": "parecer", "referencia_id": "mesmo", "processo_chave": "0333326",
             "trecho": "o proprio processo"},
        ]}

    service = MinutasGenerationService(
        _library(), analyze=analyze, generate=generate, search=search,
    )
    analyzed = service.analyze_instruction(
        documents, processo="03333/26", exercicio=2025, origem="Instituto",
    )
    generated = service.generate_instruction(
        documents, analyzed, processo="03333/26", exercicio=2025, origem="Instituto",
    )
    word = MinutasDocxService(_library(), template_path=ELVIRA_TEMPLATE).export_instruction(
        analyzed, generated, processo="03333/26", exercicio=2025, origem="Instituto",
    )
    assert captured["ids"] == [item["fonte_id"] for item in documents]
    assert captured["instrucao"] is True
    assert captured["alvo"]["especialidade_id"] == "elvira.pca.rpps"
    assert captured["alvo"]["subtipo"] == "Instituto/Autarquia"
    assert captured["geracao_alvo"]["especialidade_id"] == "elvira.pca.rpps"
    assert captured["dispositivo"] == "DEMAIS_PCA_JULGAMENTO_DIRETO"
    assert captured["geracao"] == quantidade
    assert analyzed["fontes_alvo"] == quantidade
    assert hits and hits[0][1] == "0333326"
    assert [item["referencia_id"] for item in analyzed["referencias"]["referencias"]] == ["parecer-1"]
    assert word.content.startswith(b"PK")
    assert word.name == "MINUTA_PARECER_TC_03333_26.docx"


def _irregularidade(identifier, fonte_id, *, trecho):
    return {
        "id": identifier,
        "apontamento": f"Apontamento {identifier} de cobertura documental",
        "fatos_documentados": [{
            "descricao": f"Fato {identifier}",
            "fontes": [{
                "fonte_id": fonte_id,
                "pagina": 1,
                "trecho_suporte": trecho,
            }],
        }],
        "normas_mencionadas": [],
        "argumentos_defesa": [],
        "pontos_controvertidos": [f"Controversia {identifier}."],
        "hipoteses_juridicas": ["Conferir a cobertura documental."],
        "avaliacao_auditoria": {
            "descricao": f"Conclusao indicada {identifier}.",
            "tratamento_defesa": "NAO_IDENTIFICADO",
            "fontes": [{
                "fonte_id": fonte_id,
                "pagina": 1,
                "trecho_suporte": "citacao inexistente na pagina",
            }],
        },
    }


def _estrutura_varias(identifiers):
    return {
        "identificacao_processual": "Processo TC 03333/26",
        "ementa": "Ementa de teste da instrucao.",
        "relatorio": "Relatorio da instrucao processual.",
        "relatorio_paragrafos": ["Relatorio da instrucao processual."],
        "fundamentacao_geral": "Fundamentacao geral de teste.",
        "fundamentacao_geral_paragrafos": ["Fundamentacao geral de teste."],
        "topicos_irregularidades": [{
            "titulo": "Cobertura documental",
            "texto": "Os apontamentos permanecem em conferencia.",
            "paragrafos": ["Os apontamentos permanecem em conferencia."],
            "irregularidades_ids": list(identifiers),
            "avaliacoes_irregularidades": [{
                "irregularidade_id": identifier,
                "tratamento_ministerial_defesa": "NAO_APLICAVEL",
                "fundamentos_divergencia": [],
                "estado_saneamento": "NAO_ALEGADO",
                "fontes_saneamento": [],
            } for identifier in identifiers],
        }],
        "analise_sancionatoria": {
            "conclusao": "REVISAO_HUMANA_NECESSARIA",
            "irregularidades_ids": list(identifiers),
            "fundamentos_legais": [],
            "responsaveis": [],
            "dano_ao_erario": "REVISAO_HUMANA_NECESSARIA",
            "texto": "O cabimento de multa depende de revisão humana.",
        },
        "conclusao_pedidos": "Pede-se a conferencia.",
        "pedidos": ["Pela irregularidade das contas, em conferencia humana."],
        "encerramento": "E o parecer, salvo melhor juizo.",
    }


def test_geracao_com_advertencias_e_precedentes_usa_o_contrato_historico():
    base = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )
    identifiers = ["IRR-1", "IRR-2", "IRR-3", "IRR-4"]
    literal = (
        "Pagamento sem cobertura documental no instituto de previdencia municipal"
    )
    captured = {}

    def analyze(_context):
        return {
            "analise": {
                "resumo_processual": "Instrucao com quatro apontamentos.",
                "sem_irregularidades_identificadas": False,
                "pontos_atencao_gerais": ["Conferir as citacoes."],
                "irregularidades": [
                    _irregularidade(
                        identifier, base["fonte_id"],
                        trecho=literal if identifier == "IRR-1" else "trecho ausente da pagina",
                    )
                    for identifier in identifiers
                ],
            },
            "modelo": "modelo-analise-real",
            "duracao_ms": 1800,
            "tokens_entrada": 1200,
            "tokens_saida": 400,
            "tokens_total": 1600,
        }

    def generate(context):
        captured["historicas"] = context["referencias_historicas"]
        captured["recuperadas"] = context["referencias_juridicas_recuperadas"]
        captured["fatos"] = context["documentos_processo_alvo"]
        captured["orientacao"] = context["orientacao_conferencia_documental"]
        return {
            "estrutura": _estrutura_varias(identifiers),
            "modelo": "modelo-redacao-real",
            "duracao_ms": 2400,
            "tokens_total": 900,
        }

    def search(_query, **kwargs):
        captured["excluido"] = kwargs.get("exclude_process_key")
        return {"resultados": [
            {"papel": "parecer", "referencia_id": "parecer-outro", "processo_chave": "0165326",
             "numero_processo": "01653/26", "trecho": "Art. 40 da Constituicao orienta a contribuicao."},
            {"papel": "parecer", "referencia_id": "parecer-alvo", "processo_chave": "0333326",
             "numero_processo": "03333/26", "trecho": "Segredo do proprio processo."},
            {"papel": "auditoria", "referencia_id": "relatorio", "processo_chave": "0165326",
             "numero_processo": "01653/26", "trecho": "Fato historico que nao pode migrar."},
        ]}

    service = MinutasGenerationService(
        _library(), analyze=analyze, generate=generate, search=search,
    )
    analyzed = service.analyze_instruction(documents := [base], processo="03333/26", exercicio=2025)
    assert analyzed["modelo"] == "modelo-analise-real"
    assert analyzed["tokens_total"] == 1600
    assert analyzed["duracao_ms"] == 1800
    assert len(analyzed["pendencias_conferencia"]) >= 5
    assert len({item["irregularidade"] for item in analyzed["pendencias_conferencia"]}) == 4
    generated = service.generate_instruction(
        documents, analyzed, processo="03333/26", exercicio=2025,
    )
    assert captured["excluido"] == "0333326"
    assert isinstance(captured["historicas"], list)
    assert captured["historicas"][0]["parecer"]["papel"] == "parecer"
    assert "Art. 40" in captured["historicas"][0]["parecer"]["conteudo"][0]["texto"]
    assert "Segredo do proprio processo" not in json.dumps(captured["historicas"])
    assert "Fato historico que nao pode migrar" not in json.dumps(captured)
    assert captured["fatos"][0]["fonte_id"] == base["fonte_id"]
    assert "não impedem a minuta" in captured["orientacao"]
    assert generated["advertencias_documentais"]
    assert generated["referencias_historicas"] == 1
    assert generated["modelo"] == "modelo-redacao-real"
    assert generated["texto"]
    assert "estado" not in captured["historicas"][0]


def test_resposta_invalida_nao_conclui_a_minuta():
    base = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )

    def analyze(_context):
        return {"analise": _analysis(base)}

    def generate(_context):
        return {"estrutura": {"ementa": "Incompleta."}}

    service = MinutasGenerationService(
        _library(), analyze=analyze, generate=generate,
        search=lambda *_args, **_kwargs: {"resultados": []},
    )
    analyzed = service.analyze_instruction([base], processo="03333/26", exercicio=2025)
    with pytest.raises(MinutasGenerationError, match="partes obrigatórias"):
        service.generate_instruction([base], analyzed, processo="03333/26", exercicio=2025)


def test_ausencia_de_precedente_nao_inventa_fundamento():
    base = _doc(
        "inicial.pdf",
        "Pagamento sem cobertura documental no instituto de previdencia municipal",
        "relatorio_inicial",
    )
    captured = {}

    def generate(context):
        captured["historicas"] = context["referencias_historicas"]
        captured["recuperadas"] = context["referencias_juridicas_recuperadas"]
        return {"estrutura": _structure(), "modelo": "modelo-redacao", "duracao_ms": 10, "tokens_total": 1}

    service = MinutasGenerationService(
        _library(), analyze=lambda _context: {"analise": _analysis(base)}, generate=generate,
        search=lambda *_args, **_kwargs: {"resultados": []},
    )
    analyzed = service.analyze_instruction([base], processo="03333/26")
    generated = service.generate_instruction([base], analyzed, processo="03333/26")
    assert captured["historicas"] == []
    assert captured["recuperadas"] == []
    assert generated["referencias_historicas"] == 0
    assert generated["referencias_juridicas_enviadas"] == 0


def test_erro_de_geracao_e_instrucao_chega_ao_usuario_e_o_inesperado_nao_vaza(monkeypatch, caplog):
    import streamlit as st
    from services.minutas_generation import MinutasGenerationError
    from services.minutas_ia_ui import _error
    from services.minutas_instrucao import MinutasInstrucaoError

    shown = []
    monkeypatch.setattr(st, "error", shown.append)
    _error(MinutasGenerationError("As fontes foram alteradas. Analise novamente."))
    _error(MinutasInstrucaoError("Inclua ao menos um PDF de instrução."))
    assert "fontes foram alteradas" in shown[0]
    assert "PDF de instrução" in shown[1]
    secret = "sb_secret_nao_logar"
    try:
        raise RuntimeError(secret)
    except RuntimeError as exc:
        with caplog.at_level(logging.ERROR):
            _error(exc, etapa="geracao_instrucao")
    assert "Tente novamente" in shown[2]
    assert secret not in shown[2]
    assert "Diagnóstico" in shown[2]
    assert "RuntimeError" in caplog.text
    assert "geracao_instrucao" in caplog.text
    assert secret not in caplog.text


def test_exportacao_da_instrucao_recebe_a_biblioteca():
    from services import minutas_ia_ui
    source = inspect.getsource(minutas_ia_ui._instruction_panel)
    assert "service.library" not in source
    assert "MinutasDocxService(service)" in source


def _header_pdf(processo, origem, exercicio, extra="Relatório inicial."):
    return _pdf(
        f"Processo TC: {processo}\n"
        f"Origem: {origem}\n"
        f"Exercício das contas: {exercicio}\n"
        f"{extra}"
    )


def test_sugerir_metadados_instrucao_consensus_conflict_and_absence():
    same = _doc(
        "a.pdf",
        "Processo TC: 01653/26\nOrigem: Instituto Alfa\nExercício das contas: 2026\nTexto.",
        "relatorio_inicial",
    )
    other = _doc(
        "b.pdf",
        "Processo TC: 02720/25\nOrigem: Instituto Beta\nExercício das contas: 2025\nTexto.",
        "relatorio_inicial",
    )
    empty = _doc("c.pdf", "Sem identificação cadastral neste teor.", "relatorio_inicial")
    agreed = sugerir_metadados_instrucao([same, empty])
    assert agreed["processo"] == "01653/26"
    assert agreed["exercicio"] == 2026
    assert agreed["origem"] == "Instituto Alfa"
    conflict = sugerir_metadados_instrucao([same, other])
    assert conflict["processo"] == conflict["origem"] == ""
    assert conflict["exercicio"] is None
    assert any("processos distintos" in item for item in conflict["advertencias"])
    absent = sugerir_metadados_instrucao([empty])
    assert absent["processo"] == absent["origem"] == ""
    assert absent["exercicio"] is None


def test_generator_tab_uses_instruction_flow_without_historical_mode(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from services import minutas_ia_ui

    calls = {"analyze": 0, "generate": 0, "export": 0}

    class Generator:
        def __init__(self, _library):
            pass

        def analyze_instruction(self, documents, **_kwargs):
            calls["analyze"] += 1
            return {
                "analise": {
                    "resumo_processual": "Análise da instrução.",
                    "irregularidades": [],
                    "pontos_atencao_gerais": [],
                },
                "modelo": "modelo-analise",
                "duracao_ms": 1,
                "tokens_total": 1,
                "fontes_versao": "v1",
            }

        def generate_instruction(self, documents, analyzed, **_kwargs):
            calls["generate"] += 1
            assert analyzed["fontes_versao"] == "v1"
            return {
                "estrutura": {"ementa": "Ementa."},
                "texto": "Minuta da instrução.",
                "fontes_versao": "v1",
            }

    class Word:
        def __init__(self, _library):
            pass

        def export_instruction(self, analyzed, generated, **_kwargs):
            calls["export"] += 1
            return type(
                "Doc", (),
                {"content": b"PK", "name": "MINUTA.docx", "mime": "application/vnd.word"},
            )()

    monkeypatch.setattr(minutas_ia_ui, "_service", lambda *_a, **_k: object())
    monkeypatch.setattr(minutas_ia_ui, "MinutasGenerationService", Generator)
    monkeypatch.setattr(minutas_ia_ui, "MinutasDocxService", Word)

    def page():
        from services.minutas_ia_ui import _generator_tab

        _generator_tab(None, {})

    app = AppTest.from_function(page, default_timeout=30).run()
    assert not app.exception
    assert not any(getattr(item, "key", None) == "minutas_generation_mode" for item in app.radio)
    assert "Processo do acervo" not in " ".join(str(item.value) for item in app.radio)
    assert any("Gerar Minuta de Parecer" in str(item.value) for item in app.subheader)
    assert any(item.key == "minutas_instrucao_incluir" for item in app.button)
    assert any(item.key == "minutas_instrucao_analisar" for item in app.button)
    app.file_uploader(key="minutas_instrucao_upload").set_value([
        ("relatorio.pdf", _header_pdf("01653/26", "Instituto Alfa", "2026"), "application/pdf"),
    ]).run()
    app.button(key="minutas_instrucao_incluir").click().run()
    assert app.text_input(key="minutas_instrucao_processo").value == "01653/26"
    assert app.text_input(key="minutas_instrucao_origem_processo").value == "Instituto Alfa"
    assert int(app.number_input(key="minutas_instrucao_exercicio").value) == 2026
    assert len(app.session_state["minutas_instrucao_documentos"]) == 1
    app.button(key="minutas_instrucao_incluir").click().run()
    assert len(app.session_state["minutas_instrucao_documentos"]) == 1
    hashes = set(app.session_state["minutas_instrucao_meta_hashes"])
    app.run()
    assert set(app.session_state["minutas_instrucao_meta_hashes"]) == hashes
    assert calls["analyze"] == 0
    app.button(key="minutas_instrucao_analisar").click().run()
    assert calls["analyze"] == 1
    app.button(key="minutas_instrucao_gerar").click().run()
    assert calls["generate"] == 1
    assert calls["export"] >= 1
    assert any(item.label == "Baixar DOCX" for item in app.get("download_button"))
    exported = calls["export"]
    app.run()
    assert calls["analyze"] == 1
    assert calls["generate"] == 1
    assert calls["export"] >= exported


def test_incluir_preserva_edicao_manual_conflito_e_falha_de_extracao(monkeypatch):
    from streamlit.testing.v1 import AppTest
    from services import minutas_ia_ui
    from services import minutas_instrucao as instrucao

    monkeypatch.setattr(minutas_ia_ui, "_service", lambda *_a, **_k: object())
    monkeypatch.setattr(
        minutas_ia_ui, "MinutasGenerationService", lambda _library: type("G", (), {})()
    )

    def page():
        from services.minutas_ia_ui import _generator_tab

        _generator_tab(None, {})

    app = AppTest.from_function(page, default_timeout=30).run()
    assert not app.exception
    app.file_uploader(key="minutas_instrucao_upload").set_value([
        ("a.pdf", _header_pdf("01653/26", "Instituto Alfa", "2026"), "application/pdf"),
    ]).run()
    app.button(key="minutas_instrucao_incluir").click().run()
    app.text_input(key="minutas_instrucao_processo").set_value("09999/99").run()
    app.file_uploader(key="minutas_instrucao_upload").set_value([
        ("b.pdf", _header_pdf("02720/25", "Instituto Beta", "2025"), "application/pdf"),
    ]).run()
    app.button(key="minutas_instrucao_incluir").click().run()
    assert app.text_input(key="minutas_instrucao_processo").value == "09999/99"
    assert len(app.session_state["minutas_instrucao_documentos"]) == 2
    assert any("processos distintos" in str(item.value) for item in app.caption)

    def boom(_documents):
        raise RuntimeError("falha sintética")

    monkeypatch.setattr(instrucao, "sugerir_metadados_instrucao", boom)
    app.file_uploader(key="minutas_instrucao_upload").set_value([
        ("c.pdf", _header_pdf("01499/26", "Instituto Gama", "2024"), "application/pdf"),
    ]).run()
    app.button(key="minutas_instrucao_incluir").click().run()
    assert len(app.session_state["minutas_instrucao_documentos"]) == 3
    assert app.text_input(key="minutas_instrucao_processo").value == "09999/99"
    assert any("Não foi possível extrair automaticamente" in item.value for item in app.warning)


def _profiled_library(specialty, *, motor="", estado="A iniciar"):
    from services.minutas_specialties import scope

    profile = dict(scope(specialty))
    profile["id"] = specialty
    profile["motor"] = motor
    profile["estado"] = estado
    profile["config_versao"] = "1"
    searches = []

    class Repository:
        specialty_id = specialty

        def profile(self):
            return profile

        def create(self, *_args, **_kwargs):
            raise AssertionError("a instrução não pode cadastrar documento")

        def add_file(self, *_args, **_kwargs):
            raise AssertionError("a instrução não pode cadastrar documento")

        def search_opinions(self, query, procuradora_key, extractor_version, **kwargs):
            searches.append(kwargs.get("specialty_id"))
            return []

    class Library:
        specialty_id = specialty
        repository = Repository()
        storage = object()

        def _authorize(self):
            return None

    library = Library()
    library.searches = searches
    return library


def test_camaras_desabilitada_e_outras_especialidades_nao_geram():
    from services.minutas_specialties import CAMARAS, SECRETARIAS_MUNICIPAIS, SHARED_INSTRUCTION_MOTOR
    document = _doc(
        "inicial.pdf",
        "Relatório inicial da Câmara Municipal sem irregularidades.",
        "relatorio_inicial",
    )
    blocked = [
        _profiled_library(CAMARAS),
        _profiled_library(CAMARAS, motor=SHARED_INSTRUCTION_MOTOR, estado="A iniciar"),
        _profiled_library(SECRETARIAS_MUNICIPAIS),
        _profiled_library(SECRETARIAS_MUNICIPAIS, motor=SHARED_INSTRUCTION_MOTOR, estado="A iniciar"),
        _profiled_library(
            "elvira.pca.governo", motor=SHARED_INSTRUCTION_MOTOR, estado="Em testes",
        ),
        _profiled_library(
            "elvira.tema.licitacoes", motor=SHARED_INSTRUCTION_MOTOR, estado="Em testes",
        ),
    ]
    for library in blocked:
        service = MinutasGenerationService(
            library, analyze=lambda _context: {"analise": _analysis(document)},
            generate=lambda _context: {"estrutura": _structure()},
        )
        with pytest.raises(MinutasGenerationError, match="motor habilitado"):
            service.analyze_instruction(
                [document], processo="02108/26", exercicio=2025,
                origem="Câmara Municipal de Exemplo",
            )


def test_camaras_habilitada_usa_identidade_propria_sem_misturar_rpps():
    from services.minutas_search import MinutasSearchService
    from services.minutas_specialties import CAMARAS, RPPS, SHARED_INSTRUCTION_MOTOR

    document = _doc(
        "inicial.pdf",
        "Relatório inicial da Câmara Municipal sem irregularidades.",
        "relatorio_inicial",
    )
    library = _profiled_library(CAMARAS, motor=SHARED_INSTRUCTION_MOTOR, estado="Em testes")
    captured = {}
    searches = []

    def search(query, **kwargs):
        searches.append(kwargs)
        return {"resultados": [
            {"papel": "parecer", "referencia_id": "camara-1", "processo_chave": "0999900",
             "numero_processo": "09999/00", "trecho": "parecer de câmara"},
            {"papel": "parecer", "referencia_id": "alvo", "processo_chave": "0210826",
             "trecho": "o proprio processo"},
        ]}

    def analyze(context):
        captured["analise"] = context["processo_alvo"]
        return {"analise": _analysis(document)}

    def generate(context):
        captured["geracao"] = context["processo_alvo"]
        captured["dispositivo"] = context["orientacao_dispositivo"]
        return {"estrutura": _structure()}

    service = MinutasGenerationService(
        library, analyze=analyze, generate=generate, search=search,
    )
    analyzed = service.analyze_instruction(
        [document], processo="02108/26", exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    )
    generated = service.generate_instruction(
        [document], analyzed, processo="02108/26", exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    )
    assert captured["analise"]["especialidade_id"] == CAMARAS
    assert captured["analise"]["subtipo"] == "Câmara Municipal"
    assert captured["analise"]["classe"] == "PCA"
    assert captured["geracao"]["especialidade_id"] == CAMARAS
    assert captured["dispositivo"] == "DEMAIS_PCA_JULGAMENTO_DIRETO"
    assert captured["dispositivo"] != "PREFEITURA_CONDICIONAL_DUAS_COMPETENCIAS"
    assert searches and searches[0].get("exclude_process_key") == "0210826"
    assert [item["referencia_id"] for item in analyzed["referencias"]["referencias"]] == ["camara-1"]

    rpps_library = _profiled_library(RPPS, motor=SHARED_INSTRUCTION_MOTOR)
    MinutasSearchService(library).search("regularidade das contas")
    MinutasSearchService(rpps_library).search("regularidade das contas")
    assert library.searches == [CAMARAS]
    assert rpps_library.searches == [RPPS]
    if not ELVIRA_TEMPLATE.is_file():
        pytest.skip("modelo Word indisponível")
    word = MinutasDocxService(library, template_path=ELVIRA_TEMPLATE).export_instruction(
        analyzed, generated, processo="02108/26", exercicio=2025,
        origem="Câmara Municipal de Exemplo",
    )
    from docx import Document
    body = "\n".join(paragraph.text for paragraph in Document(BytesIO(word.content)).paragraphs)
    assert "Câmara Municipal de Exemplo" in body
    assert "elvira.pca.rpps" not in body
    assert "Institutos Previdenciários" not in body
    assert "RPPS" not in body


@pytest.mark.parametrize("estado", ("Em testes", "Funcional com ajustes", "Validada"))
def test_secretarias_habilitada_usa_identidade_propria_e_julgamento_direto(estado):
    from services.minutas_search import MinutasSearchService
    from services.minutas_specialties import SECRETARIAS_MUNICIPAIS, SHARED_INSTRUCTION_MOTOR

    document = _doc(
        "inicial.pdf",
        "Relatório inicial da Secretaria Municipal de Educação sem irregularidades.",
        "relatorio_inicial",
    )
    library = _profiled_library(
        SECRETARIAS_MUNICIPAIS, motor=SHARED_INSTRUCTION_MOTOR, estado=estado,
    )
    captured = {}

    def search(_query, **kwargs):
        captured["pesquisa"] = kwargs
        return {"resultados": [
            {"papel": "parecer", "referencia_id": "secretaria-1", "processo_chave": "0999900",
             "numero_processo": "09999/00", "trecho": "parecer de secretaria"},
            {"papel": "parecer", "referencia_id": "alvo", "processo_chave": "0210826",
             "trecho": "o proprio processo"},
        ]}

    def analyze(context):
        captured["analise"] = context["processo_alvo"]
        return {"analise": _analysis(document)}

    def generate(context):
        captured["geracao"] = context["processo_alvo"]
        captured["dispositivo"] = context["orientacao_dispositivo"]
        return {"estrutura": _structure()}

    service = MinutasGenerationService(library, analyze=analyze, generate=generate, search=search)
    analyzed = service.analyze_instruction(
        [document], processo="02108/26", exercicio=2025,
        origem="Secretaria Municipal de Educação de Exemplo",
    )
    generated = service.generate_instruction(
        [document], analyzed, processo="02108/26", exercicio=2025,
        origem="Secretaria Municipal de Educação de Exemplo",
    )
    assert captured["analise"]["especialidade_id"] == SECRETARIAS_MUNICIPAIS
    assert captured["analise"]["classe"] == "PCA"
    assert captured["analise"]["subtipo"] == "Secretaria Municipal ou Estadual"
    assert captured["geracao"]["especialidade_id"] == SECRETARIAS_MUNICIPAIS
    assert captured["dispositivo"] == "DEMAIS_PCA_JULGAMENTO_DIRETO"
    assert captured["pesquisa"]["exclude_process_key"] == "0210826"
    assert [item["referencia_id"] for item in analyzed["referencias"]["referencias"]] == ["secretaria-1"]
    MinutasSearchService(library).search("regularidade das contas")
    assert library.searches == [SECRETARIAS_MUNICIPAIS]
    if not ELVIRA_TEMPLATE.is_file():
        pytest.skip("modelo Word indisponível")
    word = MinutasDocxService(library, template_path=ELVIRA_TEMPLATE).export_instruction(
        analyzed, generated, processo="02108/26", exercicio=2025,
        origem="Secretaria Municipal de Educação de Exemplo",
    )
    from docx import Document
    body = "\n".join(paragraph.text for paragraph in Document(BytesIO(word.content)).paragraphs)
    assert "Secretaria Municipal de Educação de Exemplo" in body
    assert "Câmara Municipal" not in body
    assert "Institutos Previdenciários" not in body


@pytest.mark.parametrize(
    ("texts", "requests", "has_precedent", "expected_primary"),
    (
        (
            ("Contas de governo: a instrução não contém elementos de gestão.",),
            ("Proponha-se a emissão de parecer contrário sobre as contas de governo.",),
            False,
                ("parecer prévio contrário",),
        ),
        (
            ("Contas de gestão: a instrução não contém elementos de governo.",),
            ("Proponha-se a irregularidade das contas submetidas a julgamento.",),
            True,
            ("irregularidade das contas",),
        ),
        (
            (
                "Contas de governo: proposta de parecer prévio consta da instrução.",
                "Contas de gestão: há objeto distinto submetido a julgamento direto.",
            ),
            (
                "Proponha-se a emissão de parecer contrário sobre as contas de governo.",
                "Proponha-se a irregularidade das contas submetidas a julgamento.",
                "Aplique-se multa individualizada ao responsável pela irregularidade.",
            ),
            True,
                ("parecer prévio contrário", "irregularidade das contas"),
        ),
    ),
)
def test_prefeituras_habilitada_gera_parecer_unico_condicional_e_isolado(
    texts, requests, has_precedent, expected_primary,
):
    from services.minutas_search import MinutasSearchService
    from services.minutas_specialties import PREFEITURAS, SHARED_INSTRUCTION_MOTOR

    documents = [
        _doc(
            f"instrucao-{index}.pdf", text,
            "relatorio_inicial" if index == 1 else "complementacao",
        )
        for index, text in enumerate(texts, start=1)
    ]
    library = _profiled_library(
        PREFEITURAS, motor=SHARED_INSTRUCTION_MOTOR, estado="Em testes",
    )
    captured = {}

    def search(_query, **kwargs):
        captured["pesquisa"] = kwargs
        rows = []
        if has_precedent:
            rows.append({
                "papel": "parecer", "referencia_id": "prefeitura-1",
                "processo_chave": "0999900", "numero_processo": "09999/00",
                "trecho": "precedente de Prefeitura pertinente",
            })
        rows.append({
            "papel": "parecer", "referencia_id": "alvo",
            "processo_chave": "0210826", "trecho": "o próprio processo",
        })
        return {"resultados": rows}

    def analyze(context):
        captured["analise"] = context["processo_alvo"]
        return {"analise": _analysis(documents[0])}

    def generate(context):
        captured["geracao"] = context
        structure = _structure()
        structure["pedidos"] = list(requests)
        return {"estrutura": structure}

    service = MinutasGenerationService(
        library, analyze=analyze, generate=generate, search=search,
    )
    analyzed = service.analyze_instruction(
        documents, processo="02108/26", exercicio=2025,
        origem="Prefeitura Municipal de Exemplo",
    )
    generated = service.generate_instruction(
        documents, analyzed, processo="02108/26", exercicio=2025,
        origem="Prefeitura Municipal de Exemplo",
    )

    assert captured["analise"]["especialidade_id"] == PREFEITURAS
    assert captured["analise"]["subtipo"] == "Prefeitura"
    assert captured["geracao"]["orientacao_dispositivo"] == (
        "PREFEITURA_CONDICIONAL_DUAS_COMPETENCIAS"
    )
    assert captured["pesquisa"]["exclude_process_key"] == "0210826"
    assert [item["referencia_id"] for item in analyzed["referencias"]["referencias"]] == (
        ["prefeitura-1"] if has_precedent else []
    )
    assert analyzed["referencias"]["estado"] == (
        "RECUPERADO" if has_precedent else "SEM_RESULTADO"
    )
    assert all(
        expected.casefold() in item.casefold()
        for item, expected in zip(generated["estrutura"]["pedidos"], expected_primary)
    )
    assert len(generated["estrutura"]["pedidos"]) == len(requests)
    MinutasSearchService(library).search("regularidade das contas")
    assert library.searches == [PREFEITURAS]
    if not ELVIRA_TEMPLATE.is_file():
        pytest.skip("modelo Word indisponível")
    word = MinutasDocxService(library, template_path=ELVIRA_TEMPLATE).export_instruction(
        analyzed, generated, processo="02108/26", exercicio=2025,
        origem="Prefeitura Municipal de Exemplo",
    )
    from docx import Document
    body = "\n".join(paragraph.text for paragraph in Document(BytesIO(word.content)).paragraphs)
    assert "Prefeitura Municipal de Exemplo" in body
    assert "Institutos Previdenciários" not in body


def test_desabilitar_camaras_nao_afeta_rpps():
    from services.minutas_specialties import CAMARAS, SHARED_INSTRUCTION_MOTOR

    document = _doc(
        "inicial.pdf",
        "Relatório inicial da Câmara Municipal sem irregularidades.",
        "relatorio_inicial",
    )
    disabled = _profiled_library(CAMARAS, motor="", estado="A iniciar")
    with pytest.raises(MinutasGenerationError, match="motor habilitado"):
        MinutasGenerationService(disabled).analyze_instruction(
            [document], processo="02108/26", origem="Câmara Municipal",
        )
    rpps = MinutasGenerationService(
        _library(), analyze=lambda _context: {"analise": _analysis(document)},
        generate=lambda _context: {"estrutura": _structure()},
        search=lambda *_a, **_k: {"resultados": []},
    )
    analyzed = rpps.analyze_instruction([document], processo="03333/26", exercicio=2025)
    generated = rpps.generate_instruction([document], analyzed, processo="03333/26", exercicio=2025)
    assert analyzed["processo_alvo_id"].startswith("instrucao-")
    assert generated["estrutura"]
