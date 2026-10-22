"""Resolução e aplicação do DOCX visual cadastrado por Procurador e especialidade.

O arquivo vigente é a base do documento exportado. Quando o modelo traz regiões
de apresentação, a minuta nova substitui somente o texto dessas regiões.
O gerador não recria a página com estilos de outro gabinete nem reaproveita o
texto processual do modelo.
"""

import re
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from services.minutas_docx import (
    ELVIRA_TEMPLATE,
    MinutasDocxError,
    _ALINEA,
    _blocks,
    _closing,
    _conclusion,
    _is_report_transition,
    _numbering_definition,
    _apply_numbering,
    _paragraph_metrics,
    _signature_date,
)


MENSAGEM_MODELO_AUSENTE = (
    "Não foi possível gerar o DOCX porque não há um modelo visual válido "
    "cadastrado para este Procurador e esta especialidade. Cadastre ou "
    "regularize o modelo correspondente."
)
MENSAGEM_MODELO_INVALIDO = (
    "Não foi possível gerar o DOCX porque o modelo visual cadastrado para "
    "este Procurador e esta especialidade é inválido ou incompatível. "
    "Regularize o modelo correspondente."
)
_MAX_BYTES = 10 * 1024 * 1024
_MAX_ENTRIES = 200
_MAX_UNCOMPRESSED = 40 * 1024 * 1024
_BLOQUEADOS = (
    "vbaproject",
    "oleobject",
    "activex",
    "embeddings",
    "attachedtemplate",
)


def _legado_institucional(library):
    """Só o piloto que já possui arquivo institucional próprio pode seguir sem cadastro."""
    from services.minutas_procuradores import legado_espo

    procurador = getattr(library, "procurador", None) or legado_espo()
    return procurador.sigla == "ESPO" and ELVIRA_TEMPLATE.is_file()


def _servico(library):
    from services.minutas_modelos import MinutasModeloVisualService

    repository = getattr(library, "repository", None)
    store = getattr(repository, "store", None) if repository is not None else None
    storage = getattr(library, "storage", None)
    specialty = getattr(library, "specialty_id", None)
    if store is None or storage is None or not specialty:
        return None
    return MinutasModeloVisualService(library)


def impressao_modelo(library):
    """Identidade do modelo vigente, sem baixar o arquivo nem registrar o texto."""
    service = _servico(library)
    if service is None or not service.enabled():
        return ""
    row = service.active(getattr(library, "specialty_id", None))
    if not row:
        return "ausente"
    return f"{row['id']}:{row['versao']}:{row['sha256']}"


def validar_pacote(content):
    """Recusa pacote malformado, macro ou conteúdo ativo. Não altera o original."""
    if not isinstance(content, (bytes, bytearray)) or not content:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    blob = bytes(content)
    if len(blob) > _MAX_BYTES or not blob.startswith(b"PK"):
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    try:
        package = ZipFile(BytesIO(blob))
    except Exception as exc:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO) from exc
    with package:
        return _validar_zip(package, blob)


def _validar_zip(package, blob):
    total = 0
    names = package.namelist()
    if len(names) > _MAX_ENTRIES:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    for info in package.infolist():
        name = str(info.filename or "").replace("\\", "/")
        if not name or name.startswith("/") or ".." in name.split("/"):
            raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
        lowered = name.casefold()
        if any(token in lowered for token in _BLOQUEADOS) or lowered.endswith(".docm"):
            raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
        total += int(info.file_size or 0)
        if total > _MAX_UNCOMPRESSED:
            raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    if "[Content_Types].xml" not in names or not any(
        name.startswith("word/") for name in names
    ):
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    return blob


def consultar_modelo(library):
    """Metadados e bytes do modelo ativo deste Procurador e desta especialidade."""
    service = _servico(library)
    if service is None or not service.enabled():
        return None
    specialty = getattr(library, "specialty_id", None)
    row = service.active(specialty)
    if not row:
        return {
            "estado": "ausente",
            "procuradora_chave": service._chave(),
            "especialidade_id": str(specialty),
        }
    try:
        content = library.storage.download(row["path"])
        blob = validar_pacote(content)
    except MinutasDocxError:
        return {"estado": "invalido", "id": str(row["id"]), "versao": int(row["versao"])}
    except Exception:
        return {"estado": "invalido", "id": str(row["id"]), "versao": int(row["versao"])}
    if sha256(blob).hexdigest() != str(row["sha256"]):
        return {"estado": "invalido", "id": str(row["id"]), "versao": int(row["versao"])}
    return {
        "estado": "valido",
        "id": str(row["id"]),
        "versao": int(row["versao"]),
        "sha256": str(row["sha256"]),
        "procuradora_chave": str(row["procuradora_chave"]),
        "especialidade_id": str(row["especialidade_id"]),
        "content": blob,
    }


def garantir_modelo_exportavel(library):
    """Impede a chamada ao modelo de IA quando o cadastro consultável não tem DOCX válido.

    Sem tabela ou sem armazenamento não há o que resolver aqui; a exportação é que
    recusa outro Procurador sem modelo próprio.
    """
    found = consultar_modelo(library)
    if found is None or (found["estado"] == "ausente" and _legado_institucional(library)):
        return None
    if found["estado"] == "valido":
        return found
    if found["estado"] == "invalido":
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    raise MinutasDocxError(MENSAGEM_MODELO_AUSENTE)


def resolver_para_exportacao(library):
    """Modelo vigente. O arquivo institucional legado não substitui outro Procurador."""
    found = consultar_modelo(library)
    if found is not None and found["estado"] == "valido":
        return found
    if found is not None and found["estado"] == "invalido":
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    if _legado_institucional(library):
        return None
    raise MinutasDocxError(MENSAGEM_MODELO_AUSENTE)


def _alinhamento(paragraph):
    try:
        return paragraph.alignment
    except ValueError:
        return None


def _snapshot(paragraph):
    if paragraph is None:
        return None
    properties = paragraph._p.find(qn("w:pPr"))
    run_properties = None
    for run in paragraph.runs:
        if str(run.text or "").strip():
            run_properties = run._r.find(qn("w:rPr"))
            break
    return {
        "pPr": deepcopy(properties) if properties is not None else None,
        "rPr": deepcopy(run_properties) if run_properties is not None else None,
    }


_MESES = (
    "janeiro", "fevereiro", "março", "marco", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
)


def _eh_data(text):
    """Linha curta de data. Menção a um mês no corpo não serve de protótipo."""
    plain = " ".join(str(text or "").split())
    return bool(plain) and len(plain) <= 140 and _DATA_CURTA.search(plain) is not None


def _prototipos(document):
    samples = [
        (paragraph, _paragraph_metrics(paragraph))
        for paragraph in document.paragraphs
        if str(paragraph.text or "").strip() and not _eh_faixa_paragrafo(paragraph._p)
    ]
    dated = next((item for item, _metrics in samples if _eh_data(item.text)), None)
    if dated is not None:
        samples = [(item, metrics) for item, metrics in samples if item is not dated]
    body = next((item for item, metrics in samples if metrics.get("first_cm")), None)
    if body is None:
        body = next(
            (item for item, metrics in samples if not metrics["center"]),
            samples[0][0] if samples else None,
        )
    indented = [
        (item, metrics)
        for item, metrics in samples
        if (metrics.get("left_cm") or 0) >= 5
    ]
    ementa = max(indented, key=lambda item: item[1]["left_cm"])[0] if indented else None
    if ementa is None:
        ementa = next(
            (
                item for item, _metrics in samples
                if _alinhamento(item) == WD_ALIGN_PARAGRAPH.RIGHT
            ),
            None,
        )
    heading = next(
        (
            item
            for item, metrics in samples
            if metrics["bold"] and metrics["italic"] and not metrics["center"]
        ),
        None,
    )
    centered = [(item, metrics) for item, metrics in samples if metrics["center"]]
    named = next((item for item, metrics in reversed(centered) if metrics["bold"]), None)
    role = next((item for item, metrics in reversed(centered) if not metrics["bold"]), None)
    if named is None:
        named = next(
            (
                item
                for item, metrics in reversed(samples)
                if metrics["bold"] and len(str(item.text or "").strip()) <= 80
            ),
            None,
        )
    ordered = [item for item, _metrics in samples]
    if role is None and named in ordered:
        index = ordered.index(named)
        if index + 1 < len(ordered):
            role = ordered[index + 1]
    return {
        "corpo": _snapshot(body),
        "ementa": _snapshot(ementa or body),
        "titulo": _snapshot(heading or body),
        "assinatura": _snapshot(named or body),
        "cargo": _snapshot(role or named or body),
        "data": _snapshot(dated),
    }


def _limpar_corpo(document):
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _sem_rastreio(snapshot):
    if not snapshot or snapshot["pPr"] is None:
        return snapshot
    properties = deepcopy(snapshot["pPr"])
    for tag in ("w:numPr", "w:sectPr", "w:rPr"):
        node = properties.find(qn(tag))
        if node is not None:
            properties.remove(node)
    return {"pPr": properties, "rPr": snapshot["rPr"]}


def _escrever(document, snapshot, text):
    paragraph = document.add_paragraph()
    chosen = _sem_rastreio(snapshot)
    if chosen and chosen["pPr"] is not None:
        current = paragraph._p.find(qn("w:pPr"))
        if current is not None:
            paragraph._p.remove(current)
        paragraph._p.insert(0, deepcopy(chosen["pPr"]))
    lines = str(text or "").splitlines() or [""]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        if chosen and chosen["rPr"] is not None:
            current = run._r.find(qn("w:rPr"))
            if current is not None:
                run._r.remove(current)
            run._r.insert(0, deepcopy(chosen["rPr"]))
    return paragraph


def _sanitizar(document):
    properties = document.core_properties
    properties.title = ""
    properties.subject = ""
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.author = ""
    properties.last_modified_by = ""
    part = document.part
    for relationship in list(part.rels.values()):
        kind = str(relationship.reltype or "").casefold()
        if any(token in kind for token in ("comment", "customxml", "oleobject", "embedding", "activex")):
            part.drop_rel(relationship.rId)


def _blocos(structure):
    pieces = []
    pieces.append(("ementa", f"Ementa: {structure['ementa']}"))
    for block in _blocks(structure["relatorio"], structure.get("relatorio_paragrafos")):
        _transition, text = _is_report_transition(block)
        pieces.append(("corpo", text))
    if str(structure.get("admissibilidade") or "").strip():
        pieces.append(("titulo", "Admissibilidade"))
        for block in _blocks(structure["admissibilidade"]):
            pieces.append(("corpo", block))
    for block in _blocks(
        structure["fundamentacao_geral"], structure.get("fundamentacao_geral_paragrafos")
    ):
        pieces.append(("corpo", block))
    for topic in structure["topicos_irregularidades"]:
        pieces.append(("titulo", topic["titulo"]))
        for block in _blocks(topic["texto"], topic.get("paragrafos")):
            pieces.append(("corpo", block))
    responsibility = structure.get("responsabilizacao_administrativa")
    if isinstance(responsibility, dict) and str(responsibility.get("texto") or "").strip():
        pieces.append(("titulo", str(responsibility.get("titulo") or "Responsabilização administrativa")))
        for block in _blocks(responsibility["texto"]):
            pieces.append(("corpo", block))
    sanction = structure.get("analise_sancionatoria")
    if isinstance(sanction, dict) and str(sanction.get("texto") or "").strip():
        pieces.append(("titulo", "Análise sancionatória"))
        for block in _blocks(sanction["texto"]):
            pieces.append(("corpo", block))
    conclusion = _conclusion(structure)
    if conclusion:
        for block in _blocks(conclusion):
            pieces.append(("corpo", block))
    return pieces


_DATA_CURTA = re.compile(
    r"\d{1,2}\s+de\s+(?:janeiro|fevereiro|março|marco|abril|maio|junho|julho|"
    r"agosto|setembro|outubro|novembro|dezembro|"
    r"jan\.|fev\.|mar\.|abr\.|mai\.|jun\.|jul\.|ago\.|set\.|out\.|nov\.|dez\.)"
    r"\s+de\s+\d{4}",
    re.IGNORECASE,
)
_LINHA_PARECER = re.compile(r"^parecer\s*n[.ºo°]?\s*$", re.IGNORECASE)
_ALINEA_MANUAL = re.compile(r"^[a-z]\)\s+\S", re.IGNORECASE)
_PREFIXO_ALINEA = re.compile(r"^[a-z]\)\s+", re.IGNORECASE)


def _filhos(document):
    return [child for child in document.element.body if child.tag != qn("w:sectPr")]


def _texto(elemento):
    return "".join(node.text or "" for node in elemento.iter(qn("w:t"))).strip()


def _texto_run(run):
    return "".join(node.text or "" for node in run.findall(qn("w:t")))


def _tem_midia(elemento):
    for tag in ("w:drawing", "w:pict", "w:object"):
        if next(elemento.iter(qn(tag)), None) is not None:
            return True
    return False


def _propriedade(paragraph, tag):
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        return None
    return properties.find(qn(tag))


def _jc(paragraph):
    align = _propriedade(paragraph, "w:jc")
    if align is None:
        return ""
    return str(align.get(qn("w:val")) or "")


def _estilo(paragraph):
    style = _propriedade(paragraph, "w:pStyle")
    if style is None:
        return ""
    return str(style.get(qn("w:val")) or "")


def _tem_num(paragraph):
    return _propriedade(paragraph, "w:numPr") is not None


def _marca(run, tag):
    properties = run.find(qn("w:rPr"))
    if properties is None:
        return False
    node = properties.find(qn(tag))
    if node is None:
        return False
    return str(node.get(qn("w:val")) or "true").casefold() not in {"0", "false", "off"}


def _runs_texto(paragraph):
    return [run for run in paragraph.iter(qn("w:r")) if _texto_run(run).strip()]


def _sombreado(table):
    for shade in table.iter(qn("w:shd")):
        fill = str(shade.get(qn("w:fill")) or "").upper()
        if fill and fill not in {"AUTO", "FFFFFF", "NONE"}:
            return fill
    return ""


def _linhas_tabela(table):
    return table.findall(qn("w:tr"))


def _celulas(table):
    rows = _linhas_tabela(table)
    if not rows:
        return []
    return rows[0].findall(qn("w:tc"))


def _eh_ementa(table):
    rows = _linhas_tabela(table)
    cells = _celulas(table)
    if len(rows) > 2 or len(cells) < 2:
        return False
    texts = [_texto(cell) for cell in cells]
    if any("ementa" in text.casefold() for text in texts):
        return True
    if len(rows) == 1 and len(cells) == 2:
        sizes = [len(text) for text in texts]
        return min(sizes) <= 20 and max(sizes) >= 20
    return False


def _eh_faixa(table):
    if _eh_ementa(table):
        return False
    text = _texto(table)
    if len(_linhas_tabela(table)) > 2 or not 3 <= len(text) <= 80:
        return False
    if _sombreado(table):
        return True
    return bool(re.match(r"(parecer|voto|decis[aã]o)\b", text, re.IGNORECASE))


def _paragrafo_sombreado(elemento):
    properties = elemento.find(qn("w:pPr"))
    if properties is None:
        return False
    shade = properties.find(qn("w:shd"))
    if shade is None:
        return False
    fill = str(shade.get(qn("w:fill")) or "").upper()
    return bool(fill) and fill not in {"AUTO", "FFFFFF", "NONE"}


def _eh_faixa_paragrafo(child):
    """Faixa PARECER desenhada como parágrafo sombreado ou com letras espaçadas."""
    if child.tag != qn("w:p") or _tem_midia(child):
        return False
    plain = _texto(child).strip()
    if not plain or len(plain) > 40:
        return False
    compacto = re.sub(r"[^0-9A-Za-zÀ-ÿ]+", "", plain).casefold()
    if not compacto.startswith("parecer"):
        return False
    if re.search(r"p\s+a\s+r\s+e\s+c\s+e\s+r", plain, re.IGNORECASE):
        return True
    if compacto == "parecer" or _paragrafo_sombreado(child):
        return True
    return False


def _titulo_sombreado(item):
    return bool(
        _paragrafo_sombreado(item)
        and _runs_texto(item)
        and all(_marca(run, "w:b") for run in _runs_texto(item))
    )


def _paragrafo_data(paragraph):
    text = _texto(paragraph).replace("\u00a0", " ")
    return bool(text) and len(text) <= 140 and not _tem_num(paragraph) and bool(
        _DATA_CURTA.search(text)
    )


def _mapear_regioes(document):
    """Localiza as regiões. Sem tabela, uma faixa PARECER em parágrafo também delimita."""
    children = _filhos(document)
    if not children:
        return None
    estruturais = []
    limite = max(20, len(children) // 3)
    for index, child in enumerate(children):
        if index > limite or child.tag != qn("w:tbl"):
            continue
        if _eh_ementa(child):
            estruturais.append((index, "ementa", child))
        elif _eh_faixa(child):
            estruturais.append((index, "faixa", child))
    if not estruturais:
        for index, child in enumerate(children):
            if index > limite:
                break
            if _eh_faixa_paragrafo(child):
                estruturais.append((index, "faixa", child))
                break
    if not estruturais:
        return None
    data = next(
        (child for child in reversed(children) if child.tag == qn("w:p") and _paragrafo_data(child)),
        None,
    )
    if data is None:
        return None
    primeiro = min(item[0] for item in estruturais)
    ultimo = max(item[0] for item in estruturais)
    data_index = children.index(data)
    if data_index <= ultimo:
        return None
    posteriores = [
        child for child in children[data_index + 1:]
        if child.tag == qn("w:p") and _texto(child)
    ]
    if not posteriores:
        return None
    return {
        "filhos": children,
        "ementa": next((item[2] for item in estruturais if item[1] == "ementa"), None),
        "faixa": next((item[2] for item in estruturais if item[1] == "faixa"), None),
        "data": data,
        "identificacao": [
            child for child in children[:primeiro]
            if child.tag == qn("w:p") and _texto(child)
        ],
        "nome": posteriores[0],
        "cargo": posteriores[1] if len(posteriores) > 1 else None,
        "primeiro": primeiro,
        "ultimo": ultimo,
        "zona": children[ultimo + 1:data_index],
        "ancora": children[primeiro],
    }


def _vazio_estrutural(elemento):
    return elemento.tag == qn("w:p") and not _texto(elemento) and not _tem_midia(elemento)


def _miolo(zona):
    inicio = 0
    while inicio < len(zona) and _vazio_estrutural(zona[inicio]):
        inicio += 1
    fim = len(zona)
    while fim > inicio and _vazio_estrutural(zona[fim - 1]):
        fim -= 1
    return zona[inicio:fim]


def _grupo_dispositivo(meio):
    dispositivo = []
    modo = None
    for elemento in reversed(meio):
        if elemento.tag != qn("w:p"):
            break
        texto = _texto(elemento)
        if not texto:
            continue
        if _tem_midia(elemento):
            if texto:
                break
            continue
        if _tem_num(elemento):
            if modo not in {None, "num"}:
                break
            modo = "num"
            dispositivo.append(elemento)
            continue
        if _ALINEA_MANUAL.match(texto):
            if modo not in {None, "letra"}:
                break
            modo = "letra"
            dispositivo.append(elemento)
            continue
        break
    dispositivo.reverse()
    return dispositivo, modo


def _intro_dispositivo(meio, dispositivo):
    if not dispositivo:
        return None
    indice = meio.index(dispositivo[0])
    for elemento in reversed(meio[:indice]):
        if elemento.tag != qn("w:p") or not _texto(elemento):
            continue
        if _tem_num(elemento) or _tem_midia(elemento):
            return None
        return elemento
    return None


def _prototipos_regiao(meio, dispositivo, intro):
    excluidos = set(dispositivo)
    if intro is not None:
        excluidos.add(intro)
    narrativa = [
        item for item in meio
        if item.tag == qn("w:p")
        and item not in excluidos
        and _texto(item)
        and not _tem_midia(item)
        and not _tem_num(item)
    ]
    relatorio = next(
        (
            item for item in narrativa
            if len(_texto(item)) >= 80 and not _titulo_sombreado(item)
        ),
        narrativa[0] if narrativa else None,
    )
    estilo_relatorio = _estilo(relatorio) if relatorio is not None else ""
    contagem = {}
    for item in narrativa:
        if item is relatorio or len(_texto(item)) < 80:
            continue
        estilo = _estilo(item)
        if estilo and estilo != estilo_relatorio:
            contagem[estilo] = contagem.get(estilo, 0) + 1
    fundamento = None
    if contagem:
        escolhido = max(contagem, key=contagem.get)
        fundamento = next(
            item for item in narrativa
            if _estilo(item) == escolhido and item is not relatorio
        )
    candidatos_titulo = [
        item for item in narrativa
        if len(_texto(item)) <= 180
        and _runs_texto(item)
        and all(_marca(run, "w:b") for run in _runs_texto(item))
    ]
    titulo = next((item for item in candidatos_titulo if _paragrafo_sombreado(item)), None)
    if titulo is None:
        titulo = next(
            (
                item for item in narrativa
                if _paragrafo_sombreado(item) and 8 < len(_texto(item)) <= 180
            ),
            None,
        )
    if titulo is None and candidatos_titulo:
        titulo = candidatos_titulo[0]
    citacao = next(
        (
            item for item in narrativa
            if len(_texto(item)) >= 40
            and _runs_texto(item)
            and all(_marca(run, "w:i") for run in _runs_texto(item))
        ),
        None,
    )
    return relatorio, _ou(fundamento, relatorio), titulo, citacao


def _definir_negrito(properties, ativo):
    if properties is None:
        return
    for tag in ("w:b", "w:bCs"):
        node = properties.find(qn(tag))
        if ativo and node is None:
            properties.append(OxmlElement(tag))
        elif not ativo and node is not None:
            properties.remove(node)


def _propriedades_run(paragraph, preferir_negrito):
    escolhida = None
    for run in _runs_texto(paragraph):
        properties = run.find(qn("w:rPr"))
        negrito = _marca(run, "w:b")
        if preferir_negrito is True and negrito:
            return deepcopy(properties) if properties is not None else None
        if preferir_negrito is False and not negrito:
            return deepcopy(properties) if properties is not None else None
        if escolhida is None:
            escolhida = properties
    if escolhida is None:
        return None
    copia = deepcopy(escolhida)
    if preferir_negrito is True:
        _definir_negrito(copia, True)
    elif preferir_negrito is False:
        _definir_negrito(copia, False)
    return copia


def _montar_run(texto, properties, *, negrito):
    run = OxmlElement("w:r")
    props = deepcopy(properties) if properties is not None else None
    if negrito is True:
        if props is None:
            props = OxmlElement("w:rPr")
        _definir_negrito(props, True)
    elif negrito is False and props is not None:
        _definir_negrito(props, False)
    if props is not None:
        run.append(props)
    linhas = str(texto or "").splitlines() or [""]
    for index, linha in enumerate(linhas):
        if index:
            run.append(OxmlElement("w:br"))
        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = linha
        run.append(node)
    return run


def _substituir_paragrafo(paragraph, texto, *, negrito=None):
    properties = paragraph.find(qn("w:pPr"))
    run_properties = _propriedades_run(paragraph, negrito)
    for child in list(paragraph):
        if child is not properties:
            paragraph.remove(child)
    paragraph.append(_montar_run(texto, run_properties, negrito=negrito))


def _clonar_paragrafo(origem, texto, *, negrito, manter_numeracao):
    clone = deepcopy(origem)
    properties = clone.find(qn("w:pPr"))
    if properties is not None and not manter_numeracao:
        number = properties.find(qn("w:numPr"))
        if number is not None:
            properties.remove(number)
    for child in list(clone):
        if child is not properties:
            clone.remove(child)
    clone.append(_montar_run(
        texto, _propriedades_run(origem, negrito), negrito=negrito,
    ))
    return clone


def _definir_numeracao(paragraph, number_id):
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph.insert(0, properties)
    current = properties.find(qn("w:numPr"))
    if current is not None:
        properties.remove(current)
    number = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), str(number_id))
    number.extend((level, identifier))
    properties.append(number)


def _manter_com_proximo(paragraph):
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph.insert(0, properties)
    if properties.find(qn("w:keepNext")) is None:
        properties.append(OxmlElement("w:keepNext"))


def _unir_bloco(inicio, fim):
    atual = inicio
    while atual is not None and atual is not fim:
        if atual.tag == qn("w:p"):
            _manter_com_proximo(atual)
        atual = atual.getnext()


def _remover(elementos):
    for elemento in elementos:
        pai = elemento.getparent()
        if pai is not None:
            pai.remove(elemento)


def _remover_apos(elemento):
    atual = elemento.getnext()
    while atual is not None:
        proximo = atual.getnext()
        if atual.tag != qn("w:sectPr"):
            pai = atual.getparent()
            if pai is not None:
                pai.remove(atual)
        atual = proximo


def _linhas_cabecalho(header_lines, tem_faixa):
    lines = []
    for line in header_lines or ():
        text = str(line or "").strip()
        if not text:
            continue
        if tem_faixa and _LINHA_PARECER.match(text):
            continue
        lines.append(text)
    return lines


def _pecas(structure):
    pieces = []
    for block in _blocks(structure.get("relatorio"), structure.get("relatorio_paragrafos")):
        _transicao, text = _is_report_transition(block)
        pieces.append(("relatorio", text))
    if str(structure.get("admissibilidade") or "").strip():
        pieces.append(("titulo", "Admissibilidade"))
        for block in _blocks(structure.get("admissibilidade")):
            pieces.append(("fundamentacao", block))
    for block in _blocks(
        structure.get("fundamentacao_geral"),
        structure.get("fundamentacao_geral_paragrafos"),
    ):
        pieces.append(("fundamentacao", block))
    for topic in structure.get("topicos_irregularidades") or ():
        if not isinstance(topic, dict):
            continue
        title = str(topic.get("titulo") or "").strip()
        if title:
            pieces.append(("titulo", title))
        for block in _blocks(topic.get("texto"), topic.get("paragrafos")):
            pieces.append(("fundamentacao", block))
    responsibility = structure.get("responsabilizacao_administrativa")
    if isinstance(responsibility, dict) and str(responsibility.get("texto") or "").strip():
        pieces.append((
            "titulo",
            str(responsibility.get("titulo") or "Responsabilização administrativa"),
        ))
        for block in _blocks(responsibility.get("texto")):
            pieces.append(("fundamentacao", block))
    sanction = structure.get("analise_sancionatoria")
    if isinstance(sanction, dict) and str(sanction.get("texto") or "").strip():
        pieces.append(("titulo", "Análise sancionatória"))
        for block in _blocks(sanction.get("texto")):
            pieces.append(("fundamentacao", block))
    conclusion = _conclusion(structure)
    if conclusion:
        for block in _blocks(conclusion):
            pieces.append(("conclusao", block))
    closing = _closing(structure.get("encerramento"))
    if closing:
        pieces.append(("encerramento", closing))
    return pieces


def _inserir_pecas(ancora, pecas, copias):
    base = _ou(copias["relatorio"], copias["fundamento"])
    for role, text in pecas:
        if not str(text or "").strip():
            continue
        if role == "titulo" and copias["titulo"] is not None:
            origem, negrito = copias["titulo"], None
        elif role == "titulo":
            origem, negrito = _ou(copias["fundamento"], base), True
        elif role == "relatorio":
            origem, negrito = _ou(copias["relatorio"], base), False
        elif role == "conclusao":
            origem, negrito = _ou(copias["intro"], copias["fundamento"], base), False
        elif str(text).startswith(('"', "“", "‘", "'")) and copias["citacao"] is not None:
            origem, negrito = copias["citacao"], False
        elif role == "encerramento":
            origem, negrito = _ou(copias["relatorio"], base), False
        else:
            origem, negrito = _ou(copias["fundamento"], base), False
        ancora.addprevious(_clonar_paragrafo(
            origem, text, negrito=negrito, manter_numeracao=False,
        ))


def _inserir_pedidos(document, ancora, proto, modo, pedidos, fallback):
    numero = None
    for pedido in pedidos or ():
        for line in str(pedido).splitlines():
            text = line.strip()
            if not text:
                continue
            if modo in {"num", "letra"}:
                text = _PREFIXO_ALINEA.sub("", text, count=1).strip()
                text = re.sub(r"^\d{1,2}[.)]\s+", "", text).strip()
            if not text:
                continue
            origem = proto if proto is not None else fallback
            paragraph = _clonar_paragrafo(
                origem,
                text,
                negrito=False,
                manter_numeracao=modo == "num" and proto is not None,
            )
            if modo == "letra":
                if numero is None:
                    numero = _numbering_definition(
                        document,
                        kind="lowerLetter",
                        marker="%1)",
                        left_twips=720,
                        hanging_twips=360,
                    )
                _definir_numeracao(paragraph, numero)
            ancora.addprevious(paragraph)


def _celula_ementa(table):
    cells = _celulas(table)
    labeled = [cell for cell in cells if "ementa" in _texto(cell).casefold()]
    return max(labeled or cells, key=lambda cell: len(_texto(cell)))


def _preencher_ementa(table, texto):
    cell = _celula_ementa(table)
    paragraphs = cell.findall(qn("w:p"))
    if not paragraphs:
        return
    paragraph = paragraphs[0]
    for extra in paragraphs[1:]:
        cell.remove(extra)
    runs = paragraph.findall(qn("w:r"))
    label_index = next(
        (
            index for index, run in enumerate(runs)
            if _texto_run(run).strip()
            and "ementa" in _texto_run(run).casefold()
            and len(_texto_run(run).strip()) <= 24
        ),
        None,
    )
    plain = re.sub(
        r"^ementa\s*:\s*", "", str(texto or "").strip(), flags=re.IGNORECASE,
    )
    if label_index is None:
        _substituir_paragrafo(paragraph, plain)
    else:
        keep_until = label_index
        if label_index + 1 < len(runs) and not _texto_run(runs[label_index + 1]).strip():
            keep_until = label_index + 1
        sample = next(
            (
                run.find(qn("w:rPr"))
                for run in runs[keep_until + 1:]
                if _texto_run(run).strip()
            ),
            None,
        )
        for run in runs[keep_until + 1:]:
            paragraph.remove(run)
        paragraph.append(_montar_run(
            plain,
            deepcopy(sample) if sample is not None else None,
            negrito=False,
        ))
    for row in _linhas_tabela(table)[1:]:
        for extra_cell in row.findall(qn("w:tc")):
            for extra in extra_cell.findall(qn("w:p")):
                _substituir_paragrafo(extra, "")


def _neutralizar_faixa(table):
    text = _texto(table)
    if not re.search(r"\d", text):
        return
    cleaned = re.sub(r"(?:\s*[/._-]\s*)?\d[\d\s./_-]*$", "", text).strip() or text
    paragraphs = list(table.iter(qn("w:p")))
    if not paragraphs:
        return
    _substituir_paragrafo(paragraphs[0], cleaned)
    for extra in paragraphs[1:]:
        parent = extra.getparent()
        if parent is not None:
            parent.remove(extra)


def _processo_identificacao(target):
    from services.minutas_generation import _draft_process_display

    return _draft_process_display(target or {})


def _valor_cabecalho(header_lines, prefixo):
    for line in header_lines or ():
        text = str(line or "").strip()
        if text.casefold().startswith(prefixo):
            return text.split(":", 1)[1].strip() if ":" in text else text
    return ""


_PROXIMO_ROTULO = (
    r"processo|natureza|exerc[ií]cio|unidades?|jurisdicionad\w*|origem|entidade|"
    r"respons[aá]vel|gestor"
)


def _trocar_rotulo(texto, rotulo, valor):
    if not valor:
        return texto
    novo = re.sub(
        rf"({rotulo}\s*:\s*)(.*?)(?=(?:{_PROXIMO_ROTULO})\s*:|$)",
        lambda match: match.group(1) + valor,
        texto,
        count=1,
        flags=re.IGNORECASE,
    )
    return re.sub(
        rf"({re.escape(valor)})((?:{_PROXIMO_ROTULO})\s*:)",
        r"\1 \2",
        novo,
        count=1,
        flags=re.IGNORECASE,
    )


def _atualizar_linha_identificacao(texto, target, header_lines, tem_exercicio):
    """Troca só o valor de um rótulo de identificação. Ementa e citação ficam intactas."""
    plain = " ".join(str(texto or "").split())
    folded = plain.casefold()
    if (
        not plain
        or len(plain) > 400
        or folded.startswith("ementa")
        or not re.match(
            r"(?:processo|natureza|exerc[ií]cio|unidades?(?:\s+jurisdicionad[oa]s?)?|"
            r"jurisdicionado|origem|entidade|respons[aá]vel|gestor)\b",
            folded,
        )
    ):
        return None
    target = target or {}
    if re.match(r"processo\b", folded):
        process = _processo_identificacao(target)
        if not process:
            return None
        updated = re.sub(r"\d{2,}\s*/\s*\d{2,}", process, plain, count=1)
        return updated if updated != plain else None
    novo = plain
    if re.search(r"jurisdicion|unidade|origem|entidade", folded):
        origin = str(target.get("origem") or "").strip() or _valor_cabecalho(header_lines, "origem")
        novo = _trocar_rotulo(
            novo,
            r"(?:unidades?(?:\s+jurisdicionad[oa]s?)?|jurisdicionado|origem|entidade)",
            origin,
        )
    if re.search(r"exerc[ií]cio", folded):
        novo = _trocar_rotulo(novo, r"exerc[ií]cio", str(target.get("exercicio") or "").strip())
    if re.search(r"respons[aá]vel|gestor", folded):
        responsible = str(target.get("responsavel") or "").strip() or _valor_cabecalho(
            header_lines, "responsável"
        )
        novo = _trocar_rotulo(novo, r"(?:respons[aá]vel|gestor)", responsible)
    if re.search(r"natureza\s*:", folded):
        value = _valor_cabecalho(header_lines, "natureza")
        if tem_exercicio or re.search(r"exerc[ií]cio\s*:", folded):
            value = re.sub(
                r"\s*[-–]\s*exerc[ií]cio\b.*$", "", value, flags=re.IGNORECASE,
            ).strip()
        novo = _trocar_rotulo(novo, r"natureza", value)
    return novo if novo != plain else None


def _preencher_identificacao(regioes, header_lines, target, ementa=""):
    linhas = _linhas_cabecalho(header_lines, regioes["faixa"] is not None)
    identificacao = regioes["identificacao"]
    textos = [_texto(paragraph) for paragraph in identificacao]
    tem_exercicio = any(
        re.search(r"exerc[ií]cio", texto, re.IGNORECASE) and "natureza" not in texto.casefold()
        for texto in textos
    )
    atualizadas = [
        _atualizar_linha_identificacao(texto, target, linhas, tem_exercicio)
        for texto in textos
    ]
    if identificacao and atualizadas and all(item is not None for item in atualizadas):
        for paragraph, texto in zip(identificacao, atualizadas):
            _substituir_paragrafo(paragraph, texto)
        return
    removeu_ementa = False
    for paragraph, novo in zip(identificacao, atualizadas):
        if novo:
            _substituir_paragrafo(paragraph, novo)
            continue
        bruto = _texto(paragraph).strip()
        if bruto.casefold().startswith("ementa"):
            _remover([paragraph])
            removeu_ementa = True
            continue
        if len(bruto) > 180 and not re.match(
            rf"(?:{_PROXIMO_ROTULO}|parecer)\b", bruto, re.IGNORECASE,
        ):
            texto = re.sub(r"\s+", " ", str(ementa or "")).strip()
            if texto:
                _substituir_paragrafo(paragraph, "Ementa: " + texto)
            else:
                _remover([paragraph])
            removeu_ementa = True
    if any(atualizadas) or removeu_ementa:
        return
    if not linhas:
        return
    for index, linha in enumerate(linhas):
        if index < len(identificacao):
            _substituir_paragrafo(identificacao[index], linha)
            continue
        origem = identificacao[-1] if identificacao else None
        if origem is None:
            novo = OxmlElement("w:p")
            novo.append(_montar_run(linha, None, negrito=True))
        else:
            novo = _clonar_paragrafo(origem, linha, negrito=None, manter_numeracao=False)
        regioes["ancora"].addprevious(novo)
    _remover(identificacao[len(linhas):])


def _copiar(elemento):
    return deepcopy(elemento) if elemento is not None else None


def _ou(*elementos):
    for elemento in elementos:
        if elemento is not None:
            return elemento
    return None


def _aplicar_regioes(document, regioes, target, structure, profile, header_lines):
    meio = _miolo(regioes["zona"])
    dispositivo, modo = _grupo_dispositivo(meio)
    intro = _intro_dispositivo(meio, dispositivo)
    relatorio, fundamento, titulo, citacao = _prototipos_regiao(meio, dispositivo, intro)
    base = _ou(
        relatorio,
        fundamento,
        intro,
        regioes["identificacao"][0] if regioes["identificacao"] else None,
    )
    if base is None:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    copias = {
        "relatorio": _copiar(_ou(relatorio, base)),
        "fundamento": _copiar(_ou(fundamento, relatorio, base)),
        "titulo": _copiar(titulo),
        "citacao": _copiar(citacao),
        "intro": _copiar(intro),
        "dispositivo": _copiar(dispositivo[0]) if dispositivo else None,
    }
    entre = []
    for child in regioes["filhos"][regioes["primeiro"] + 1:regioes["ultimo"]]:
        if child is regioes["ementa"] or child is regioes["faixa"]:
            continue
        if _vazio_estrutural(child):
            continue
        entre.append(child)
    _remover(entre)
    _remover(meio)
    ancora = regioes["data"]
    pecas = _pecas(structure)
    _inserir_pecas(
        ancora, [item for item in pecas if item[0] != "encerramento"], copias,
    )
    _inserir_pedidos(
        document,
        ancora,
        copias["dispositivo"],
        modo,
        structure.get("pedidos") or (),
        copias["relatorio"],
    )
    _inserir_pecas(
        ancora, [item for item in pecas if item[0] == "encerramento"], copias,
    )
    _preencher_identificacao(
        regioes, header_lines, target, structure.get("ementa") or "",
    )
    if regioes["ementa"] is not None:
        _preencher_ementa(regioes["ementa"], structure.get("ementa") or "")
    if regioes["faixa"] is not None:
        _neutralizar_faixa(regioes["faixa"])
    _substituir_paragrafo(regioes["data"], _signature_date(target, profile))
    _substituir_paragrafo(regioes["nome"], str(profile["nome"]))
    cargo = regioes["cargo"]
    if cargo is None:
        cargo = _clonar_paragrafo(
            regioes["nome"], str(profile["cargo"]), negrito=False, manter_numeracao=False,
        )
        regioes["nome"].addnext(cargo)
    else:
        _substituir_paragrafo(cargo, str(profile["cargo"]))
    _remover_apos(cargo)
    _unir_bloco(regioes["data"], cargo)


def _aplicar_paragrafos(document, target, structure, profile, header_lines):
    """Modelos sem tabela estrutural continuam reconstruídos por protótipo de parágrafo."""
    prototypes = _prototipos(document)
    if prototypes["corpo"] is None:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    faixas = [deepcopy(child) for child in _filhos(document) if _eh_faixa_paragrafo(child)]
    _limpar_corpo(document)
    for line in header_lines:
        _escrever(document, prototypes["corpo"], line)
    for role, text in _blocos(structure):
        snapshot = prototypes["titulo"] if role == "titulo" else prototypes[role]
        _escrever(document, snapshot, text)
    requests = [str(item) for item in structure.get("pedidos") or () if str(item).strip()]
    if requests:
        number_id = _numbering_definition(
            document, kind="decimal", marker="%1.", left_twips=720, hanging_twips=360,
        )
        alpha_id = None
        for request in requests:
            lines = [line.strip() for line in request.splitlines() if line.strip()]
            if not lines:
                continue
            paragraph = _escrever(document, prototypes["corpo"], lines[0])
            _apply_numbering(paragraph, number_id)
            alineas = []
            for line in lines[1:]:
                match = _ALINEA.match(line)
                if match:
                    alineas.append(match.group(1).strip())
                    continue
                _escrever(document, prototypes["corpo"], line)
            if alineas:
                if alpha_id is None:
                    alpha_id = _numbering_definition(
                        document,
                        kind="lowerLetter",
                        marker="%1)",
                        left_twips=1440,
                        hanging_twips=360,
                    )
                for item in alineas:
                    subparagraph = _escrever(document, prototypes["corpo"], item)
                    _apply_numbering(subparagraph, alpha_id)
    closing = _closing(structure.get("encerramento"))
    if closing:
        _escrever(document, prototypes["corpo"], closing)
    if prototypes.get("data"):
        _escrever(document, prototypes["data"], _signature_date(target, profile))
    _escrever(document, prototypes["assinatura"], profile["nome"])
    _escrever(document, prototypes["cargo"], profile["cargo"])
    if faixas:
        ancora = next(
            (item for item in document.paragraphs if item.text.strip().casefold().startswith("ementa")),
            None,
        )
        if ancora is None and document.paragraphs:
            ancora = document.paragraphs[0]
        if ancora is not None:
            for faixa in reversed(faixas):
                ancora._p.addnext(faixa)


def renderizar(content, target, structure, profile, header_lines):
    """Copia o modelo em memória, descarta o texto de exemplo e grava a minuta nova."""
    from docx import Document

    blob = validar_pacote(content)
    try:
        document = Document(BytesIO(blob))
    except Exception as exc:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO) from exc
    if not document.sections:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    regioes = _mapear_regioes(document)
    _sanitizar(document)
    if regioes is not None:
        _aplicar_regioes(document, regioes, target, structure, profile, header_lines)
    elif document.tables:
        raise MinutasDocxError(MENSAGEM_MODELO_INVALIDO)
    else:
        _aplicar_paragrafos(document, target, structure, profile, header_lines)
    _descartar_nos_vazios(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def metadados_publicos(modelo):
    if not modelo:
        return None
    return {
        "id": modelo.get("id"),
        "versao": modelo.get("versao"),
        "sha256": modelo.get("sha256"),
        "procuradora_chave": modelo.get("procuradora_chave"),
        "especialidade_id": modelo.get("especialidade_id"),
    }


def _descartar_nos_vazios(document):
    """Remove marcadores de comentário que tenham sobrado fora do texto substituído."""
    body = document.element.body
    for node in list(body.iter()):
        if node.tag in {
            qn("w:commentRangeStart"),
            qn("w:commentRangeEnd"),
            qn("w:commentReference"),
        }:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)
