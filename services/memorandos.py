"""Rules for substitution memoranda, independent from Streamlit and SQL."""

from datetime import date
import hashlib
import json
import re
import unicodedata

TIPO_SUBSTITUICAO = "SUBSTITUICAO"
STATUS_RASCUNHO = "RASCUNHO"
STATUS_FINALIZADO = "FINALIZADO"
STATUS_CANCELADO = "CANCELADO"
NATUREZAS = ("Cargo comissionado", "Função de confiança")
MOTIVOS = (
    "Férias regulamentares",
    "Licença para tratamento de saúde",
    "Licença especial",
    "Outro",
)


def normalize(value):
    return "".join(
        c for c in unicodedata.normalize("NFKD", (value or "").casefold())
        if not unicodedata.combining(c)
    ).strip()


def normalized_registration(value):
    text = re.sub(r"[^0-9A-Za-z]", "", str(value or "").strip())
    return "" if not text or set(text) == {"0"} else text.upper()


def fingerprint(record):
    return hashlib.sha256(json.dumps(record, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def substitution_status(record, today=None):
    if record["status"] == STATUS_CANCELADO:
        return "CANCELADA"
    today = today or date.today()
    start, end = date.fromisoformat(record["data_inicio"]), date.fromisoformat(record["data_fim"])
    return "AGENDADA" if start > today else "ENCERRADA" if end < today else "EM ANDAMENTO"


def cabinet_text(procurador, genero):
    if normalize(procurador["nome"]) == normalize("Elvira Samara Pereira de Oliveira"):
        return "lotada na Procuradoria-Geral" if genero == "Feminino" else "lotado na Procuradoria-Geral"
    article = "da" if procurador.get("genero") == "feminino" else "do"
    role = "Procuradora" if procurador.get("genero") == "feminino" else "Procurador"
    return f"{'lotada' if genero == 'Feminino' else 'lotado'} no gabinete {article} {role} {procurador['nome']}"


def display_sector(value):
    return {"1CAM": "Secretaria da 1ª Câmara", "2CAM": "Secretaria da 2ª Câmara"}.get((value or "").strip().upper(), value or "")


def validate(record):
    if record.get("tipo") != TIPO_SUBSTITUICAO:
        raise ValueError("Tipo de memorando inválido.")
    if record.get("natureza_funcao") not in NATUREZAS:
        raise ValueError("Informe a natureza da função.")
    if record.get("motivo") not in MOTIVOS or (record.get("motivo") == "Outro" and not record.get("motivo_texto", "").strip()):
        raise ValueError("Informe o motivo do afastamento.")
    if date.fromisoformat(record["data_fim"]) < date.fromisoformat(record["data_inicio"]):
        raise ValueError("A data final deve ser igual ou posterior à inicial.")
    steps = record.get("etapas") or []
    if not steps:
        raise ValueError("Inclua ao menos uma substituição.")
    seen = set()
    previous = None
    for step in steps:
        for side in ("substituido", "substituto"):
            person = step.get(side) or {}
            if not person.get("nome", "").strip() or person.get("genero") not in ("Masculino", "Feminino"):
                raise ValueError("Nome e gênero são obrigatórios em cada etapa.")
        left, right = step["substituido"], step["substituto"]
        left_key = left.get("servidor_id") or normalize(left["nome"])
        right_key = right.get("servidor_id") or normalize(right["nome"])
        if left_key == right_key or right_key in seen:
            raise ValueError("A cadeia contém pessoa repetida, ciclo ou autossubstituição.")
        if previous is not None and left_key != previous:
            raise ValueError("A cadeia de substituição está quebrada.")
        seen.add(left_key); seen.add(right_key); previous = right_key


def safe_filename(record):
    name = record["etapas"][0]["substituido"]["nome"]
    name = re.sub(r"[^\w.-]+", "_", unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()).strip("_")
    return f"Memorando_Substituicao_{name}_{record['data_inicio']}.pdf"


def short_date(value):
    parsed = date.fromisoformat(value)
    return f"{parsed.day}/{parsed.month}/{parsed.year}"


def period_text(start, end):
    return f"no período de {short_date(start)} a {short_date(end)}"


def role_article(gender):
    return "a servidora" if gender == "Feminino" else "o servidor"


def participle(gender):
    return "lotada" if gender == "Feminino" else "lotado"


def read_server_xlsx(content):
    """Read only the four supported columns; callers never persist upload bytes."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("A importação XLSX exige a dependência openpyxl.") from exc
    from io import BytesIO
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = book.active
    aliases = {"nome": ("nome", "servidor"), "matricula": ("matricula", "matrcula"), "cargo": ("cargo", "funcao"), "setor": ("setor", "lotacao")}
    source = sheet.iter_rows(values_only=True)
    header = None
    for candidate in source:
        normalized = [normalize(str(value or "")) for value in candidate]
        if any(value in aliases["nome"] for value in normalized) and any(value in aliases["matricula"] for value in normalized):
            header = normalized
            break
    if header is None:
        raise ValueError("Não foi encontrada uma linha de cabeçalhos na planilha.")
    indexes = {}
    for target, choices in aliases.items():
        indexes[target] = next((i for i, value in enumerate(header) if value in choices), None)
    if any(value is None for value in indexes.values()):
        raise ValueError("A planilha deve conter os cabeçalhos nome, matrícula, cargo e setor/lotação.")
    rows = []
    for values in source:
        rows.append({key: values[indexes[key]] if indexes[key] < len(values) else "" for key in indexes})
    return rows


def open_service(store):
    from database.memorandos import MemorandosStore
    return MemorandosStore(store)
