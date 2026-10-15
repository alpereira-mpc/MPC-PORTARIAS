"""Exportação DOCX institucional das minutas experimentais já geradas."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PROJECT_ROOT = Path(__file__).resolve().parents[1]
ELVIRA_TEMPLATE = (
    PROJECT_ROOT / "modelos" / "modelo_visual_elvira_bananeiras_convertido.docx"
)


@dataclass(frozen=True)
class MinutaDocx:
    name: str
    mime: str
    content: bytes


class MinutasDocxError(ValueError):
    pass


_PROFILES = {
    "Elvira Samara": {
        "nome": "ELVIRA SAMARA PEREIRA DE OLIVEIRA",
        "cargo": (
            "Procuradora-Geral do Ministério Público de Contas do Estado da Paraíba"
        ),
        "template": ELVIRA_TEMPLATE,
    }
}
_ALINEA = re.compile(r"^\s*[a-z]\s*[).:-]\s*(.+)$", re.IGNORECASE)
LOGO_FROM_TOP_CM = 0.81
LOGO_AFTER_PT = 24
IDENTIFICATION_AFTER_PT = 48
EMENTA_AFTER_PT = 72
CLOSING_BEFORE_PT = 24
PEDIDO_LEFT_TWIPS = 2121
PEDIDO_HANGING_TWIPS = 360
ALINEA_LEFT_TWIPS = 2841
ALINEA_HANGING_TWIPS = 360


def _plain(value):
    decomposed = unicodedata.normalize("NFKD", str(value or "").casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _canonical_process_key(value):
    normalized = re.sub(r"\s+", " ", _plain(value)).strip()
    normalized = re.sub(r"^processo\s*", "", normalized)
    normalized = re.sub(r"^tc\s*[-:]?\s*", "", normalized)
    if not re.fullmatch(r"\d{5}\s*/?\s*\d{2}", normalized):
        return ""
    return re.sub(r"\D", "", normalized)


def _normalized_support(value):
    return re.sub(r"\s+", " ", _plain(value)).strip()


def _set_run_font(run, size, *, bold=None, italic=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial")
    fonts.set(qn("w:cs"), "Arial")


def _configure_style(
    document,
    name,
    *,
    size=12,
    bold=False,
    italic=False,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    left_cm=0,
    right_cm=0,
    first_cm=0,
    before_pt=0,
    after_pt=6,
    keep_next=False,
):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial")
    fonts.set(qn("w:cs"), "Arial")
    fmt = style.paragraph_format
    fmt.alignment = alignment
    fmt.left_indent = Cm(left_cm)
    fmt.right_indent = Cm(right_cm)
    fmt.first_line_indent = Cm(first_cm)
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    fmt.line_spacing = 1
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_next
    return style


def _apply_style_overrides(document, overrides):
    for name, values in (overrides or {}).items():
        if name not in document.styles or not isinstance(values, dict):
            continue
        style = document.styles[name]
        fmt = style.paragraph_format
        if values.get("size"):
            style.font.size = Pt(values["size"])
        if "bold" in values:
            style.font.bold = values["bold"]
        if "italic" in values:
            style.font.italic = values["italic"]
        if values.get("left_cm") is not None:
            fmt.left_indent = Cm(values["left_cm"])
        if values.get("first_cm") is not None:
            fmt.first_line_indent = Cm(values["first_cm"])
        if values.get("right_cm") is not None:
            fmt.right_indent = Cm(values["right_cm"])
        if values.get("before_pt") is not None:
            fmt.space_before = Pt(values["before_pt"])
        if values.get("after_pt") is not None:
            fmt.space_after = Pt(values["after_pt"])


def _configure_styles(document, overrides=None):
    _configure_style(
        document,
        "Normal",
        size=12,
        first_cm=3,
        after_pt=6,
    )
    _configure_style(
        document,
        "Minuta Identificação",
        size=12,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        after_pt=0,
    )
    _configure_style(
        document,
        "Minuta Ementa",
        size=11,
        left_cm=6.752,
        after_pt=EMENTA_AFTER_PT,
    )
    _configure_style(document, "Minuta Corpo", size=12, first_cm=3, after_pt=12)
    _configure_style(
        document,
        "Minuta Transição",
        size=12,
        bold=True,
        first_cm=3,
        after_pt=12,
        keep_next=True,
    )
    _configure_style(
        document,
        "Minuta Irregularidade",
        size=12,
        bold=True,
        italic=True,
        left_cm=3,
        after_pt=6,
        keep_next=True,
    )
    _configure_style(
        document,
        "Minuta Transcrição",
        size=11,
        left_cm=3,
        right_cm=0.13,
        after_pt=6,
    )
    _configure_style(
        document,
        "Minuta Pedido",
        size=12,
        left_cm=PEDIDO_LEFT_TWIPS / 567,
        first_cm=-(PEDIDO_HANGING_TWIPS / 567),
        after_pt=6,
    )
    _configure_style(
        document,
        "Minuta Alínea",
        size=12,
        left_cm=ALINEA_LEFT_TWIPS / 567,
        right_cm=0.13,
        first_cm=-(ALINEA_HANGING_TWIPS / 567),
        after_pt=6,
    )
    _configure_style(
        document,
        "Minuta Encerramento",
        size=12,
        first_cm=3,
        before_pt=CLOSING_BEFORE_PT,
        after_pt=6,
    )
    _configure_style(
        document,
        "Minuta Procuradora",
        size=14,
        bold=True,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after_pt=0,
    )
    _configure_style(
        document,
        "Minuta Cargo",
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after_pt=0,
    )
    _configure_style(
        document,
        "Minuta Revisão Título",
        size=12,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before_pt=0,
        after_pt=12,
        keep_next=True,
    )
    _configure_style(
        document,
        "Minuta Revisão Item",
        size=10,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        left_cm=0.75,
        first_cm=-0.5,
        after_pt=6,
    )
    _apply_style_overrides(document, overrides)


def _metric(value, low, high):
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if low <= number <= high:
        return number
    return None


def _length_cm(value):
    return None if value is None else float(value.cm)


def _length_pt(value):
    return None if value is None else float(value.pt)


def _paragraph_metrics(paragraph):
    fmt = paragraph.paragraph_format
    style_fmt = paragraph.style.paragraph_format if paragraph.style is not None else None
    run = next((item for item in paragraph.runs if str(item.text or "").strip()), None)
    size = None
    bold = False
    italic = False
    if run is not None and run.font.size is not None:
        size = float(run.font.size.pt)
    elif paragraph.style is not None and paragraph.style.font.size is not None:
        size = float(paragraph.style.font.size.pt)
    if run is not None:
        bold = bool(run.font.bold)
        italic = bool(run.font.italic)
    left = _length_cm(fmt.left_indent)
    if left is None and style_fmt is not None:
        left = _length_cm(style_fmt.left_indent)
    first = _length_cm(fmt.first_line_indent)
    if first is None and style_fmt is not None:
        first = _length_cm(style_fmt.first_line_indent)
    after = _length_pt(fmt.space_after)
    if after is None and style_fmt is not None:
        after = _length_pt(style_fmt.space_after)
    before = _length_pt(fmt.space_before)
    if before is None and style_fmt is not None:
        before = _length_pt(style_fmt.space_before)
    return {
        "center": paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER,
        "left_cm": _metric(left, 0, 12),
        "first_cm": _metric(first, -3, 8),
        "before_pt": _metric(before, 0, 72),
        "after_pt": _metric(after, 0, 72),
        "size": _metric(size, 8, 18),
        "bold": bold,
        "italic": italic,
    }


def read_visual_model(content):
    """Lê só medidas, estilos e a logomarca. O texto do modelo não entra no resultado."""
    try:
        template = Document(BytesIO(content))
    except Exception as exc:
        raise MinutasDocxError("O modelo visual da especialidade é inválido.") from exc
    if not template.sections:
        raise MinutasDocxError("O modelo visual da especialidade não tem página configurada.")
    section = template.sections[0]
    spec = {
        "page_width_cm": _metric(_length_cm(section.page_width), 14, 30),
        "page_height_cm": _metric(_length_cm(section.page_height), 14, 45),
        "top_cm": _metric(_length_cm(section.top_margin), 0.5, 8),
        "bottom_cm": _metric(_length_cm(section.bottom_margin), 0.5, 8),
        "left_cm": _metric(_length_cm(section.left_margin), 0.5, 8),
        "right_cm": _metric(_length_cm(section.right_margin), 0.5, 8),
        "header_cm": _metric(_length_cm(section.header_distance), 0, 5),
        "footer_cm": _metric(_length_cm(section.footer_distance), 0, 5),
        "styles": {},
    }
    for relationship in section.header.part.rels.values():
        if relationship.reltype.endswith("/image"):
            spec["logo"] = relationship.target_part.blob
            break
    extents = section.header._element.xpath(".//wp:extent")
    if extents:
        spec["logo_width_in"] = _metric(int(extents[0].get("cx") or 0) / 914400, 1, 4)
    samples = [_paragraph_metrics(paragraph) for paragraph in template.paragraphs if paragraph.text.strip()]
    body = next((item for item in samples if item.get("first_cm")), None)
    ementa = max((item for item in samples if (item.get("left_cm") or 0) >= 5), default=None, key=lambda item: item["left_cm"])
    heading = next((item for item in samples if item["bold"] and item["italic"] and not item["center"]), None)
    signature = [item for item in samples if item["center"]]
    styles = {}
    if body:
        styles["Minuta Corpo"] = {key: body[key] for key in ("size", "first_cm", "after_pt") if body.get(key) is not None}
    if ementa:
        styles["Minuta Ementa"] = {
            key: ementa[key] for key in ("size", "left_cm") if ementa.get(key) is not None
        }
    if heading:
        styles["Minuta Irregularidade"] = {
            key: heading[key] for key in ("size", "left_cm", "after_pt") if heading.get(key) is not None
        }
        styles["Minuta Irregularidade"]["bold"] = True
        styles["Minuta Irregularidade"]["italic"] = True
    spec["styles"] = styles
    named = next((item for item in reversed(signature) if item["bold"]), None)
    plain = next((item for item in reversed(signature) if not item["bold"]), None)
    if named:
        spec["signature_name"] = {
            key: named[key] for key in ("size", "before_pt") if named.get(key) is not None
        }
    if plain and plain.get("size"):
        spec["signature_role"] = {"size": plain["size"]}
    return spec


def _template_logo(template_path):
    if not template_path.is_file():
        raise MinutasDocxError(
            "O modelo visual institucional da Procuradora não está disponível."
        )
    try:
        template = Document(template_path)
    except Exception as exc:
        raise MinutasDocxError("O modelo visual institucional é inválido.") from exc
    for section in template.sections:
        for relationship in section.header.part.rels.values():
            if relationship.reltype.endswith("/image"):
                return relationship.target_part.blob
    raise MinutasDocxError("A logomarca institucional não foi localizada no modelo.")


def _next_numbering_id(numbering, tag):
    values = []
    for element in numbering.findall(qn(f"w:{tag}")):
        value = element.get(qn(f"w:{tag}Id"))
        if value is not None:
            values.append(int(value))
    return max(values, default=-1) + 1


def _numbering_definition(document, *, kind, marker, left_twips, hanging_twips):
    numbering = document.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "abstractNum")
    number_id = _next_numbering_id(numbering, "num")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), kind)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), marker)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    paragraph_properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), str(left_twips))
    indentation.set(qn("w:hanging"), str(hanging_twips))
    paragraph_properties.append(indentation)
    level.extend((start, number_format, level_text, suffix, paragraph_properties))
    abstract.append(level)
    first_number = numbering.find(qn("w:num"))
    if first_number is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_number), abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return number_id


def _apply_numbering(paragraph, number_id):
    properties = paragraph._p.get_or_add_pPr()
    number_properties = properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        properties.append(number_properties)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), str(number_id))
    number_properties.extend((level, identifier))


def _add_text(paragraph, text, *, size=12, bold=None, italic=None):
    lines = str(text).splitlines() or [""]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        _set_run_font(run, size, bold=bold, italic=italic)


def _blocks(value, structured=None):
    if isinstance(structured, list):
        return [
            block.strip()
            for item in structured
            for block in re.split(r"\n\s*\n", str(item or "").strip())
            if block.strip()
        ]
    return [
        block.strip()
        for block in re.split(r"\n\s*\n", str(value or "").strip())
        if block.strip()
    ]


def _is_report_transition(block):
    from services.minutas_generation import REPORT_TRANSITION, _REPORT_TRANSITION_RE

    text = re.sub(r"\s+", " ", str(block or "")).strip()
    if not text:
        return False, text
    if text == REPORT_TRANSITION:
        return True, REPORT_TRANSITION
    if _REPORT_TRANSITION_RE.fullmatch(text):
        return True, REPORT_TRANSITION
    return False, text


def _add_body(document, value, style="Minuta Corpo", *, paragraphs=None, size=12):
    for block in _blocks(value, paragraphs):
        is_transition, block = _is_report_transition(block)
        if is_transition:
            selected_style = "Minuta Transição"
        elif block.startswith(('"', "“", "‘", "'")):
            selected_style = "Minuta Transcrição"
        else:
            selected_style = style
        paragraph = document.add_paragraph(style=selected_style)
        _add_text(
            paragraph,
            block,
            size=11 if selected_style == "Minuta Transcrição" else size,
            bold=True if is_transition else None,
        )


def _conclusion(structure):
    request_keys = {
        _normalized_support(request) for request in structure.get("pedidos", ())
    }
    kept = []
    for line in str(structure.get("conclusao_pedidos") or "").splitlines():
        unnumbered = re.sub(r"^\s*(?:\d+[.)]|[-*])\s*", "", line)
        if _normalized_support(unnumbered) not in request_keys:
            kept.append(line)
    return "\n".join(kept).strip()


def _closing(value):
    if re.search(r"\bé\s+o\s+parecer\.", str(value or ""), re.IGNORECASE):
        return "É o Parecer."
    return str(value or "").strip()


def _add_requests(document, requests):
    decimal_id = _numbering_definition(
        document,
        kind="decimal",
        marker="%1.",
        left_twips=PEDIDO_LEFT_TWIPS,
        hanging_twips=PEDIDO_HANGING_TWIPS,
    )
    for request in requests:
        lines = [line.strip() for line in str(request).splitlines() if line.strip()]
        if not lines:
            continue
        paragraph = document.add_paragraph(style="Minuta Pedido")
        _apply_numbering(paragraph, decimal_id)
        _add_text(paragraph, lines[0])
        alineas = []
        for line in lines[1:]:
            match = _ALINEA.match(line)
            if not match:
                paragraph.add_run().add_break()
                _add_text(paragraph, line)
                continue
            alineas.append(match.group(1).strip())
        if alineas:
            alpha_id = _numbering_definition(
                document,
                kind="lowerLetter",
                marker="%1)",
                left_twips=ALINEA_LEFT_TWIPS,
                hanging_twips=ALINEA_HANGING_TWIPS,
            )
            for item in alineas:
                subparagraph = document.add_paragraph(style="Minuta Alínea")
                _apply_numbering(subparagraph, alpha_id)
                _add_text(subparagraph, item)


def _filename(target):
    key = _canonical_process_key(target.get("numero_processo")) or _canonical_process_key(
        target.get("processo_chave")
    )
    if len(key) == 7:
        return f"MINUTA_PARECER_TC_{key[:5]}_{key[5:]}.docx"
    return "MINUTA_PARECER.docx"


def _style_size(visual, name, default):
    values = ((visual or {}).get("styles") or {}).get(name) or {}
    return values.get("size") or default


def _build_docx(target, structure, profile, header_lines, visual=None):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm((visual or {}).get("page_width_cm") or 21.001)
    section.page_height = Cm((visual or {}).get("page_height_cm") or 29.7)
    section.top_margin = Cm((visual or {}).get("top_cm") or 6.35)
    section.bottom_margin = Cm((visual or {}).get("bottom_cm") or 2)
    section.left_margin = Cm((visual or {}).get("left_cm") or 2.75)
    section.right_margin = Cm((visual or {}).get("right_cm") or 1.75)
    header_cm = (
        (visual or {}).get("header_cm")
        if (visual or {}).get("header_cm") is not None
        else 0.5
    )
    section.header_distance = Cm(header_cm)
    section.footer_distance = Cm((visual or {}).get("footer_cm") or 0)
    _configure_styles(document, (visual or {}).get("styles"))

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.right_indent = Cm(0)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    logo_before_pt = max(0, (LOGO_FROM_TOP_CM - float(header_cm)) * 72 / 2.54)
    paragraph.paragraph_format.space_before = Pt(logo_before_pt)
    paragraph.paragraph_format.space_after = Pt(LOGO_AFTER_PT)
    run = paragraph.add_run()
    logo_bytes = (visual or {}).get("logo") or _template_logo(Path(profile["template"]))
    logo = run.add_picture(
        BytesIO(logo_bytes), width=Inches((visual or {}).get("logo_width_in") or 2.62)
    )
    logo._inline.docPr.set("title", "MPC-PB")
    logo._inline.docPr.set(
        "descr", "Logomarca do Ministério Público de Contas da Paraíba"
    )
    extent = paragraph._p.xpath(".//wp:extent")
    if extent:
        logo_height_cm = int(extent[0].get("cy") or 0) / 914400 * 2.54
        needed_top = (
            float(header_cm)
            + logo_before_pt * 2.54 / 72
            + logo_height_cm
            + LOGO_AFTER_PT * 2.54 / 72
        )
        configured_top = float((visual or {}).get("top_cm") or 6.35)
        section.top_margin = Cm(max(configured_top, needed_top))

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_run = footer_paragraph.add_run(
        "MINUTA EXPERIMENTAL — SUJEITA À REVISÃO HUMANA"
    )
    _set_run_font(footer_run, 9, italic=True)

    for index, line in enumerate(header_lines):
        item = document.add_paragraph(style="Minuta Identificação")
        _add_text(item, line, bold=True)
        if index == len(header_lines) - 1:
            item.paragraph_format.space_after = Pt(IDENTIFICATION_AFTER_PT)

    body_size = _style_size(visual, "Minuta Corpo", 12)
    ementa_size = _style_size(visual, "Minuta Ementa", 11)
    heading_size = _style_size(visual, "Minuta Irregularidade", 12)
    ementa = document.add_paragraph(style="Minuta Ementa")
    _add_text(ementa, f"Ementa: {structure['ementa']}", size=ementa_size)
    _add_body(
        document,
        structure["relatorio"],
        paragraphs=structure.get("relatorio_paragrafos"),
        size=body_size,
    )
    _add_body(
        document,
        structure["fundamentacao_geral"],
        paragraphs=structure.get("fundamentacao_geral_paragrafos"),
        size=body_size,
    )
    for topic in structure["topicos_irregularidades"]:
        heading = document.add_paragraph(style="Minuta Irregularidade")
        _add_text(heading, topic["titulo"], bold=True, italic=True, size=heading_size)
        _add_body(
            document,
            topic["texto"],
            paragraphs=topic.get("paragrafos"),
            size=body_size,
        )

    conclusion = _conclusion(structure)
    if conclusion:
        _add_body(document, conclusion, size=body_size)
    _add_requests(document, structure["pedidos"])
    _add_body(
        document,
        _closing(structure["encerramento"]),
        style="Minuta Encerramento",
        size=body_size,
    )

    signature = (visual or {}).get("signature_name") or {}
    role = (visual or {}).get("signature_role") or {}
    name = document.add_paragraph(style="Minuta Procuradora")
    name.paragraph_format.space_before = Pt(signature.get("before_pt") or 24)
    _add_text(
        name,
        profile["nome"],
        size=signature.get("size") or 14,
        bold=True,
        italic=True,
    )
    title = document.add_paragraph(style="Minuta Cargo")
    _add_text(title, profile["cargo"], size=role.get("size") or 11)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class MinutasDocxService:
    """Revalida a sessão e exporta em memória, sem nova operação de IA."""

    def __init__(self, library, *, source_summary=None, template_path=None):
        self.library = library
        if source_summary is None:
            from services.minutas_generation import MinutasGenerationService

            source_summary = MinutasGenerationService(library).source_summary
        self._source_summary = source_summary
        self.template_path = Path(template_path) if template_path else None

    def _active_visual(self):
        if self.template_path is not None:
            return None
        from services.minutas_modelos import active_model_bytes

        content = active_model_bytes(self.library)
        if not content:
            return None
        return read_visual_model(content)

    def export(self, target_reference_id, analyzed, generated):
        from services.minutas_generation import (
            MinutasGenerationService,
            _canonical_process_key as canonical_process_key,
            _draft_header_lines,
            _target_with_header_metadata,
        )

        self.library._authorize()
        from services.minutas_specialties import scope, RPPS
        if scope(getattr(self.library, "specialty_id", None))["id"] != RPPS:
            raise MinutasDocxError("Exportação indisponível nesta especialidade.")
        target = self.library.get(target_reference_id)
        MinutasGenerationService._validate_target(target)
        if not isinstance(analyzed, dict) or not isinstance(generated, dict):
            raise MinutasDocxError("Gere novamente a minuta antes de exportar.")
        summary = self._source_summary(target_reference_id)
        current_version = summary.get("fontes_versao")
        if summary.get("issues"):
            raise MinutasDocxError(
                "Confira ou reprocesse os documentos instrutórios antes de exportar."
            )
        if any(
            item.get("fontes_versao") != current_version
            for item in (analyzed, generated)
        ):
            raise MinutasDocxError(
                "As fontes foram alteradas. Analise e gere novamente a minuta."
            )

        expected = {
            "processo_alvo_id": str(target["id"]),
            "processo_alvo_chave": canonical_process_key(
                target.get("processo_chave") or target.get("numero_processo")
            ),
            "procuradora_chave": str(target.get("procuradora_chave") or ""),
        }
        for name, value in expected.items():
            if str(analyzed.get(name) or "") != value or str(
                generated.get(name) or ""
            ) != value:
                raise MinutasDocxError(
                    "A minuta não pertence ao processo e à Procuradora selecionados."
                )

        profile = dict(_PROFILES.get(target.get("procuradora")) or {})
        if not profile:
            raise MinutasDocxError("Não há modelo Word autorizado para esta Procuradora.")
        if self.template_path:
            profile["template"] = self.template_path
        structure = MinutasGenerationService._validate_structure(
            generated.get("estrutura")
        )
        header_target = _target_with_header_metadata(
            target, generated.get("metadados_cabecalho")
        )
        structure, _normalization_warnings = (
            MinutasGenerationService._normalize_structure(structure, header_target)
        )
        structure = MinutasGenerationService._validate_structure(structure)
        return MinutaDocx(
            name=_filename(target),
            mime=DOCX_MIME,
            content=_build_docx(
                target,
                structure,
                profile,
                _draft_header_lines(header_target),
                self._active_visual(),
            ),
        )

    def export_instruction(self, analyzed, generated, *, processo, exercicio=None, origem=""):
        """DOCX da minuta gerada sobre a instrução local, sem consultar o acervo."""
        from services.minutas_generation import (
            MinutasGenerationService,
            _canonical_process_key as canonical_process_key,
            _draft_header_lines,
            _target_with_header_metadata,
        )
        from services.minutas_library import PROCURADORA

        self.library._authorize()
        key = canonical_process_key(processo)
        if len(key) != 7:
            raise MinutasDocxError("Informe o processo TC no formato 00000/00.")
        target = {
            "id": f"instrucao-{key}",
            "numero_processo": f"{key[:5]}/{key[5:]}",
            "processo_chave": key,
            "exercicio": exercicio,
            "origem": origem or "",
            "procuradora": PROCURADORA,
            "procuradora_chave": PROCURADORA.casefold(),
            "especialidade_id": "elvira.pca.rpps",
            "classificacao_estado": "CLASSIFICADO",
            "subtipo": "Instituto/Autarquia",
        }
        if not isinstance(analyzed, dict) or not isinstance(generated, dict):
            raise MinutasDocxError("Gere novamente a minuta antes de exportar.")
        if analyzed.get("fontes_versao") != generated.get("fontes_versao"):
            raise MinutasDocxError(
                "As fontes foram alteradas. Analise e gere novamente a minuta."
            )
        if (
            str(analyzed.get("processo_alvo_id") or "") != target["id"]
            or str(generated.get("processo_alvo_id") or "") != target["id"]
        ):
            raise MinutasDocxError("A minuta não pertence ao processo informado.")
        profile = dict(_PROFILES.get(PROCURADORA) or {})
        if not profile:
            raise MinutasDocxError("Não há modelo Word autorizado para esta Procuradora.")
        if self.template_path:
            profile["template"] = self.template_path
        structure = MinutasGenerationService._validate_structure(generated.get("estrutura"))
        header_target = _target_with_header_metadata(
            target, generated.get("metadados_cabecalho")
        )
        structure, _warnings = MinutasGenerationService._normalize_structure(
            structure, header_target
        )
        structure = MinutasGenerationService._validate_structure(structure)
        return MinutaDocx(
            name=_filename(target),
            mime=DOCX_MIME,
            content=_build_docx(
                target,
                structure,
                profile,
                _draft_header_lines(header_target),
                self._active_visual(),
            ),
        )
