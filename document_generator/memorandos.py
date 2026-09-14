"""Controlled institutional DOCX template for final substitution memoranda."""
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips
from services.memorandos import participle, role_article, short_date

ROOT = Path(__file__).resolve().parents[1]
PORTARIAS_LOGO = ROOT / "assets" / "logo.jpeg"
PORTARIAS_LOGO_WIDTH = Pt(163.5)
PORTARIAS_LOGO_HEIGHT = Pt(105.75)


def _add_portarias_logo(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.3
    # A4 21 cm with 3 cm / 2 cm margins: content center is 0.5 cm right of the
    # physical page center. A -1 cm left indent recenters the mark without
    # changing section margins or logo size.
    paragraph.paragraph_format.left_indent = Cm(-1)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.add_run().add_picture(
        str(PORTARIAS_LOGO),
        width=PORTARIAS_LOGO_WIDTH,
        height=PORTARIAS_LOGO_HEIGHT,
    )
    spacer = doc.add_paragraph()
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacer.paragraph_format.line_spacing = 1.3
    spacer.paragraph_format.left_indent = Twips(540)


def _font(run, *, size=13, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return run


def _paragraph(doc, text="", *, alignment=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, first=None, bold=False, size=13, physical_center=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.3
    # A4 21 cm with 3 cm / 2 cm margins: content center is 0.5 cm right of the
    # physical page center. A -1 cm left indent recenters CENTER text without
    # changing section margins.
    paragraph.paragraph_format.left_indent = Cm(-1) if physical_center else Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(first) if first else Pt(0)
    if text:
        _font(paragraph.add_run(text), size=size, bold=bold)
    return paragraph


def _runs(paragraph, chunks, *, size=13):
    for text, bold in chunks:
        _font(paragraph.add_run(text), size=size, bold=bold)
    return paragraph


def _subject(record):
    return "Substituição de servidor ocupante de " + record["natureza_funcao"].lower()


def _motive(record):
    return record["motivo_texto"].strip() if record["motivo"] == "Outro" else record["motivo"].lower()


def _person_label(person):
    return f"{person['nome']}, matrícula nº {person.get('matricula', '')}"


def _first_body_chunks(record, step):
    left, right = step["substituido"], step["substituto"]
    ending = "a" if left["genero"] == "Feminino" else "o"
    return [
        (f"Ao cumprimentá-lo, e considerando que {role_article(left['genero'])} ", False),
        (_person_label(left), True),
        (
            f", ocupante de {record['natureza_funcao'].lower()} de {left.get('cargo', '')}, "
            f"{record['gabinete_snapshot']}, encontra-se afastad{ending} de suas atividades laborais "
            f"no período de ",
            False,
        ),
        (short_date(record["data_inicio"]), True),
        (" a ", False),
        (short_date(record["data_fim"]), True),
        (
            f", em decorrência de {_motive(record)}, indico {role_article(right['genero'])} ",
            False,
        ),
        (_person_label(right), True),
        (
            f", {right.get('cargo', '')}, {participle(right['genero'])} em {right.get('lotacao', '')}, "
            f"para substituir {role_article(left['genero'])} antes mencionad{ending} durante o respectivo período.",
            False,
        ),
    ]


def _cascade_body_chunks(step):
    left, right = step["substituido"], step["substituto"]
    return [
        (f"Em decorrência da substituição acima, indico {role_article(right['genero'])} ", False),
        (_person_label(right), True),
        (
            f", {right.get('cargo', '')}, {participle(right['genero'])} em {right.get('lotacao', '')}, "
            f"para substituir, no mesmo período, {role_article(left['genero'])} ",
            False,
        ),
        (_person_label(left), True),
        (f", {left.get('cargo', '')}.", False),
    ]


def generate(record):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.25), Cm(2.0)
    section.left_margin, section.right_margin = Cm(3.0), Cm(2.0)
    section.header_distance = Cm(1.25)
    _add_portarias_logo(doc)
    _paragraph(doc, "MEMORANDO", alignment=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=18, bold=True, size=14, physical_center=True)
    _paragraph(doc, "Ao Excelentíssimo Senhor Presidente do Tribunal de Contas do Estado da Paraíba", after=12)
    subject = _paragraph(doc, "", after=14)
    _runs(subject, [("Assunto:", False), (" " + _subject(record), True)])
    _paragraph(doc, "", after=12)
    _paragraph(doc, "Senhor Presidente,", after=12)
    _paragraph(doc, "", after=12)
    for index, step in enumerate(record["etapas"]):
        body = _paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after=12, first=1.25)
        _runs(body, _first_body_chunks(record, step) if index == 0 else _cascade_body_chunks(step))
    _paragraph(doc, "", after=24)
    _paragraph(doc, "Com os meus melhores cumprimentos,", after=24)
    _paragraph(doc, "", after=24)
    _paragraph(doc, record["signatario_nome"].upper(), alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    _paragraph(
        doc,
        record["signatario_cargo"] + " do Ministério Público de Contas da Paraíba",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=13,
    )
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def preview_documents(record):
    from document_generator.pdf import convert

    docx = generate(record)
    pdf_bytes, _ = convert(docx)
    return docx, pdf_bytes


def pdf(record):
    _, content = preview_documents(record)
    return content
