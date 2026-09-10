"""Document generation from sanitized institutional packages, using existing PDF conversion."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from services.oficios import GABINETES, long_date, safe_name

ROOT = Path(__file__).resolve().parents[1]


def template_path(series):
    """A future cabinet package replaces its fallback without numbering changes."""
    folder = ROOT / "templates" / "oficios"
    code = series["sigla"]
    own = folder / f"{code}.docx" if code in GABINETES else None
    if own is not None and own.is_file():
        return own
    model = series["modelo"]
    if model not in ("PROGE", "BTLC"):
        raise ValueError("Confirme o modelo institucional desta série.")
    return folder / f"{model}.docx"


def _apply_run_props(run, props):
    if props is not None:
        run._r.insert(0, deepcopy(props))


def _heading_line(target, doc, heading, dated, model, props):
    """Keep number on the left and the full date on the right of one line."""
    target.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    usable = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    stops = target.paragraph_format.tab_stops
    for stop in list(stops):
        stops.remove_tab_stop(stop.position)
    stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    pPr = target._p.get_or_add_pPr()
    for node in pPr.findall(qn("w:wordWrap")):
        pPr.remove(node)
    wrap = OxmlElement("w:wordWrap")
    wrap.set(qn("w:val"), "off")
    pPr.append(wrap)
    target.clear()
    if model == "BTLC" and heading.startswith("Ofício "):
        first = target.add_run("Ofício ")
        first.font.name = "Calibri"
        first.font.size = Pt(12)
        middle = target.add_run(heading.removeprefix("Ofício "))
        middle.font.name = "Calibri"
        middle.font.size = Pt(13)
        last = target.add_run("\t" + dated)
        last.font.name = "Calibri"
        last.font.size = Pt(12)
        return
    left = target.add_run(heading)
    _apply_run_props(left, props)
    right = target.add_run("\t" + dated)
    _apply_run_props(right, props)


def generate(record, series, number=None):
    model = series["modelo"]
    if model not in ("PROGE", "BTLC"):
        raise ValueError("Confirme o modelo institucional desta série.")
    doc = Document(template_path(series))
    number_text = (
        str(number).zfill(series["digitos"]) if number is not None else "RASCUNHO"
    )
    heading = series["cabecalho"].format(numero=number_text, ano=record["data"][:4])
    title = record.get("titulo_assinatura") or (
        "Procuradora-Geral do Ministério Público de Contas da Paraíba"
        if series["sigla"] == "PROGE"
        else record["cargo_base"] + " do Ministério Público de Contas da Paraíba"
    )
    values = {
        "heading": heading,
        "treatment": record.get("tratamento", ""),
        "name": record.get("destinatario", ""),
        "role": record.get("cargo", ""),
        "unit": record.get("unidade", ""),
        "institution": record.get("instituicao", ""),
        "subject": "Assunto: "
        + record["assunto"]
        + ("\nReferência: " + record["referencia"] if record.get("referencia") else ""),
        "vocative": record.get("vocativo", ""),
        "body": record.get("corpo", ""),
        "closing": record.get("fechamento", "Atenciosamente,"),
        "signature": record["signatario"].upper(),
        "title": title,
    }
    dated = long_date(record["data"])
    for p in list(doc.paragraphs):
        if not p.text:
            continue
        key = p.text.strip("{}")
        chunks = (
            re.split(r"\r?\n+", values[key])
            if key in ("body", "subject")
            else [values[key]]
        )
        if not chunks[0] and key in (
            "treatment",
            "name",
            "role",
            "unit",
            "institution",
        ):
            p._p.getparent().remove(p._p)
            continue
        template = deepcopy(p._p)
        previous = p
        for index, text in enumerate(chunks):
            target = p if index == 0 else Paragraph(deepcopy(template), p._parent)
            if index:
                previous._p.addnext(target._p)
            props = (
                deepcopy(target.runs[0]._r.rPr)
                if target.runs and target.runs[0]._r.rPr is not None
                else None
            )
            target.clear()
            if key == "heading":
                _heading_line(target, doc, heading, dated, model, props)
            elif key == "subject" and index == 0:
                label = target.add_run("Assunto: ")
                _apply_run_props(label, props)
                value = target.add_run(record["assunto"].strip())
                _apply_run_props(value, props)
                value.bold = True
            else:
                run = target.add_run(text.strip())
                _apply_run_props(run, props)
            if key in ("signature", "closing"):
                target.paragraph_format.keep_with_next = True
            previous = target
    # Empty paragraphs after the signature can spill onto an otherwise blank page.
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        tail = doc.paragraphs[-1]._p
        if tail.xpath(".//w:drawing | .//w:sectPr"):
            break
        tail.getparent().remove(tail)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def filename(record, series, number, ext):
    return (
        safe_name(
            f"Oficio_{series['sigla']}_{str(number).zfill(series['digitos'])}_{record['data'][:4]}_{record.get('destinatario') or record['assunto']}"
        )
        + "."
        + ext
    )


def official_documents(record, series, number):
    from document_generator.pdf import convert

    docx = generate(record, series, number)
    try:
        pdf, _ = convert(docx)
    except RuntimeError as exc:
        raise RuntimeError(
            "Conversão PDF indisponível. O rascunho foi mantido e nenhum número foi consumido. Verifique Word/LibreOffice e tente novamente."
        ) from exc
    return [
        ("docx", filename(record, series, number, "docx"), docx),
        ("pdf", filename(record, series, number, "pdf"), pdf),
    ]
